"""从已有 .docx 读取格式，导出可复刻的配置与代码片段。

与 ``WordFormatter`` 的生成链路相反：本模块只读，输入一份现成文档，输出

* ``document``：docDefaults、分节页面、页眉页脚与页码、自动编号；
* ``styles``：标题/正文等段落样式的**生效值**，并标注每个字段来自哪一层；
* ``observed``：正文逐块的实测格式，按「样式 + 格式签名」聚类；
* ``python_code``：可直接粘贴执行的 ``WordFormatter`` 调用；
* ``notes``：现有写接口无法还原、需要手工处理的部分。

设计要点
--------
Word 的格式是三层叠加的：``docDefaults`` → 段落样式链（``basedOn``）→ 段落/字符
直接格式。只读 python-docx 的 ``paragraph_format`` 属性会得到一片 ``None``，既分不清
「没设置」和「继承自样式」，也读不到章节编号、主题字体、字符单位缩进。因此这里按字段
逐层取值，并记录来源。

本模块不直接依赖 lxml：主题部件用 ``docx.oxml.parse_xml`` 解析。
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.table import Table

from .word import WordFormatter

SCHEMA_VERSION = 1

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

_ALIGNMENT_MAP = {
    "left": "left",
    "start": "left",
    "center": "center",
    "right": "right",
    "end": "right",
    "both": "justify",
    "justify": "justify",
    "distribute": "distribute",
    "thaiDistribute": "distribute",
}

_KEY_STYLE_NAMES = (
    {"Normal", "Title", "Subtitle"}
    | {f"Heading {level}" for level in range(1, 10)}
    | {f"TOC {level}" for level in range(1, 4)}
)

_FIELD_NAMES = ("PAGE", "NUMPAGES", "SECTIONPAGES", "STYLEREF", "REF", "SEQ", "TOC")
_FIELD_RE = re.compile(r"\b(" + "|".join(_FIELD_NAMES) + r")\b")

# 这些属性必须成组生效：高层一旦写了其中任意一项，低层同组属性就不能再回填。
# 否则会出现「段落直接写固定行距、样式里的倍数又被合并进来」这种自相矛盾的配置。
_ATOMIC_GROUPS = {
    "line_spacing": ("line_spacing.",),
    "numbering": ("numbering.",),
    "first_line": ("indent.first_line_chars", "indent.first_line_pt"),
    "hanging": ("indent.hanging_chars", "indent.hanging_pt"),
    "left": ("indent.left_chars", "indent.left_pt"),
    "right": ("indent.right_chars", "indent.right_pt"),
}
_PREFIX_GROUPS = {"line_spacing.": "line_spacing", "numbering.": "numbering"}
_KEY_GROUPS = {
    key: group
    for group, keys in _ATOMIC_GROUPS.items()
    for key in keys
}


def _group_of(key: str) -> str | None:
    for prefix, group in _PREFIX_GROUPS.items():
        if key.startswith(prefix):
            return group
    return _KEY_GROUPS.get(key)


def _xml(element) -> str:
    """序列化 XML；编号/主题里的普通 lxml 元素没有 ``.xml``。"""
    if element is None:
        return ""
    try:
        return element.xml
    except AttributeError:
        from docx.oxml.xmlchemy import serialize_for_reading

        return serialize_for_reading(element)


# ====================================================================
# 基础工具：单位、属性、主题字体
# ====================================================================


def _int(value) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _float(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _twips_to_pt(value) -> float | None:
    number = _float(value)
    return None if number is None else round(number / 20, 2)


def _twips_to_cm(value) -> float | None:
    number = _float(value)
    return None if number is None else round(number / 567, 2)


def _half_points_to_pt(value) -> float | None:
    number = _float(value)
    return None if number is None else round(number / 2, 2)


def _chars(value) -> float | None:
    number = _float(value)
    return None if number is None else round(number / 100, 2)


def _emu_to_cm(value) -> float | None:
    number = _float(value)
    return None if number is None else round(number / 360000, 2)


def _length_cm(value) -> float | None:
    if value is None:
        return None
    try:
        return round(value.cm, 2)
    except (AttributeError, ValueError):
        return None


def _length_pt(value) -> float | None:
    if value is None:
        return None
    try:
        return round(value.pt, 2)
    except (AttributeError, ValueError):
        return None


def _child_val(parent, tag: str):
    """取子元素 ``w:val``；不存在返回 None。"""
    if parent is None:
        return None
    element = parent.find(qn(tag))
    return None if element is None else element.get(qn("w:val"))


def _on_off(element) -> bool | None:
    """``w:keepNext`` 这类开关：无 val 为真，显式 0/false/off 为假。"""
    if element is None:
        return None
    value = element.get(qn("w:val"))
    if value is None:
        return True
    return value.lower() not in {"0", "false", "off"}


def _attr_on_off(value) -> bool | None:
    """直接解析属性形式的开关（如 ``w:beforeAutospacing="0"``）。"""
    if value is None:
        return None
    return str(value).lower() not in {"0", "false", "off"}


def _theme_fonts(doc) -> dict[str, str]:
    """解析 ``theme1.xml``，把 ``minorEastAsia`` 之类的引用换成真实字体名。"""
    result: dict[str, str] = {}
    for part in doc.part.package.parts:
        if not str(part.partname).endswith("theme/theme1.xml"):
            continue
        try:
            root = parse_xml(part.blob)
        except Exception:  # noqa: BLE001 - 主题损坏时退化为无主题
            continue
        scheme = root.find(f".//{{{_A_NS}}}fontScheme")
        if scheme is None:
            continue
        for kind, prefix in (("majorFont", "major"), ("minorFont", "minor")):
            font = scheme.find(f"{{{_A_NS}}}{kind}")
            if font is None:
                continue
            for child, suffix in (("latin", "Ascii"), ("ea", "EastAsia"), ("cs", "Cs")):
                node = font.find(f"{{{_A_NS}}}{child}")
                if node is None:
                    continue
                typeface = (node.get("typeface") or "").strip()
                if typeface:
                    result[f"{prefix}{suffix}"] = typeface
                    # OOXML 引用名与主题槽位不同：minorHAnsi / minorBidi
                    if suffix == "Ascii":
                        result[f"{prefix}HAnsi"] = typeface
                    elif suffix == "Cs":
                        result[f"{prefix}Bidi"] = typeface
    return result


# ====================================================================
# 字段提取：把 pPr / rPr 摊平成「点路径 -> 值」
# ====================================================================


def _p_fields(pPr) -> dict[str, Any]:
    """提取段落属性。返回扁平键，便于跨层合并。"""
    out: dict[str, Any] = {}
    if pPr is None:
        return out

    jc = pPr.find(qn("w:jc"))
    if jc is not None:
        raw = jc.get(qn("w:val"))
        if raw is not None:
            out["alignment"] = _ALIGNMENT_MAP.get(raw, raw)

    spacing = pPr.find(qn("w:spacing"))
    if spacing is not None:
        before = spacing.get(qn("w:before"))
        if before is not None:
            out["space_before_pt"] = _twips_to_pt(before)
        after = spacing.get(qn("w:after"))
        if after is not None:
            out["space_after_pt"] = _twips_to_pt(after)
        before_lines = spacing.get(qn("w:beforeLines"))
        if before_lines is not None:
            out["space_before_lines"] = _chars(before_lines)
        after_lines = spacing.get(qn("w:afterLines"))
        if after_lines is not None:
            out["space_after_lines"] = _chars(after_lines)
        before_auto = spacing.get(qn("w:beforeAutospacing"))
        if before_auto is not None:
            out["space_before_autospacing"] = _attr_on_off(before_auto)
        after_auto = spacing.get(qn("w:afterAutospacing"))
        if after_auto is not None:
            out["space_after_autospacing"] = _attr_on_off(after_auto)

        line = spacing.get(qn("w:line"))
        if line is not None:
            rule = spacing.get(qn("w:lineRule")) or "auto"
            if rule in ("exact", "atLeast"):
                out["line_spacing.rule"] = "exact" if rule == "exact" else "at_least"
                out["line_spacing.pt"] = _twips_to_pt(line)
            else:
                out["line_spacing.rule"] = "multiple"
                out["line_spacing.value"] = round((_float(line) or 0) / 240, 2)
            out["line_spacing.raw_twips"] = _int(line)

    ind = pPr.find(qn("w:ind"))
    if ind is not None:
        for attr, key, converter in (
            ("w:firstLine", "first_line_pt", _twips_to_pt),
            ("w:firstLineChars", "first_line_chars", _chars),
            ("w:hanging", "hanging_pt", _twips_to_pt),
            ("w:hangingChars", "hanging_chars", _chars),
            ("w:left", "left_pt", _twips_to_pt),
            ("w:leftChars", "left_chars", _chars),
            ("w:right", "right_pt", _twips_to_pt),
            ("w:rightChars", "right_chars", _chars),
        ):
            raw = ind.get(qn(attr))
            if raw is None:
                continue
            out[f"indent.{key}"] = converter(raw)

    for tag, key in (
        ("w:keepNext", "keep_next"),
        ("w:keepLines", "keep_lines"),
        ("w:pageBreakBefore", "page_break_before"),
        ("w:widowControl", "widow_control"),
        ("w:snapToGrid", "snap_to_grid"),
    ):
        element = pPr.find(qn(tag))
        if element is not None:
            out[key] = _on_off(element)

    outline = pPr.find(qn("w:outlineLvl"))
    if outline is not None:
        out["outline_level"] = _int(outline.get(qn("w:val")))

    num_pr = pPr.find(qn("w:numPr"))
    if num_pr is not None:
        num_id = _int(_child_val(num_pr, "w:numId"))
        if num_id is not None:
            out["numbering.num_id"] = num_id
            level = _int(_child_val(num_pr, "w:ilvl"))
            if level is not None:
                out["numbering.level"] = level
    return out


def _r_fields(rPr, theme: dict[str, str]) -> dict[str, Any]:
    """提取字符属性；主题字体在读取时就换成真实字体名。"""
    out: dict[str, Any] = {}
    if rPr is None:
        return out

    fonts = rPr.find(qn("w:rFonts"))
    if fonts is not None:
        for attr, key, theme_attr in (
            ("w:ascii", "font.ascii", "w:asciiTheme"),
            ("w:hAnsi", "font.hansi", "w:hAnsiTheme"),
            ("w:eastAsia", "font.cn", "w:eastAsiaTheme"),
            ("w:cs", "font.cs", "w:cstheme"),
        ):
            literal = fonts.get(qn(attr))
            if literal:
                out[key] = literal
                continue
            reference = fonts.get(qn(theme_attr))
            if reference and reference in theme:
                out[key] = theme[reference]
            elif reference:
                out[f"{key}.theme"] = reference

    size = rPr.find(qn("w:sz"))
    if size is not None:
        out["size_pt"] = _half_points_to_pt(size.get(qn("w:val")))
    else:
        size_cs = rPr.find(qn("w:szCs"))
        if size_cs is not None:
            out["size_pt"] = _half_points_to_pt(size_cs.get(qn("w:val")))

    for tag, key in (
        ("w:b", "bold"),
        ("w:bCs", "bold_cs"),
        ("w:i", "italic"),
        ("w:iCs", "italic_cs"),
        ("w:caps", "caps"),
        ("w:smallCaps", "small_caps"),
    ):
        element = rPr.find(qn(tag))
        if element is not None:
            out[key] = _on_off(element)

    underline = rPr.find(qn("w:u"))
    if underline is not None:
        value = underline.get(qn("w:val")) or "single"
        out["underline"] = None if value == "none" else value

    color = rPr.find(qn("w:color"))
    if color is not None:
        value = color.get(qn("w:val"))
        if value:
            out["color"] = value

    highlight = rPr.find(qn("w:highlight"))
    if highlight is not None:
        out["highlight"] = highlight.get(qn("w:val"))

    shading = rPr.find(qn("w:shd"))
    if shading is not None:
        fill = shading.get(qn("w:fill"))
        if fill and fill not in ("auto", "FFFFFF"):
            out["shading"] = fill

    char_spacing = rPr.find(qn("w:spacing"))
    if char_spacing is not None:
        out["char_spacing_pt"] = _twips_to_pt(char_spacing.get(qn("w:val")))

    vert_align = rPr.find(qn("w:vertAlign"))
    if vert_align is not None:
        out["vert_align"] = vert_align.get(qn("w:val"))
    return out


def _merge_layers(layers, reader) -> tuple[dict[str, Any], dict[str, str]]:
    """按「最具体 → 最宽泛」的顺序合并，第一次出现的字段即生效值。"""
    merged: dict[str, Any] = {}
    sources: dict[str, str] = {}
    claimed_groups: set[str] = set()
    for element, source in layers:
        if element is None:
            continue
        fields = reader(element)
        layer_groups = {
            group for group in (_group_of(key) for key in fields) if group
        }
        for key, value in fields.items():
            if key in merged:
                continue
            group = _group_of(key)
            if group and group in claimed_groups:
                continue
            merged[key] = value
            sources[key] = source
        claimed_groups.update(layer_groups)
    return merged, sources


def _style_ppr(style):
    return None if style is None else style.element.find(qn("w:pPr"))


def _style_rpr(style):
    return None if style is None else style.element.find(qn("w:rPr"))


def _style_outline_level(style) -> int | None:
    return _int(_child_val(_style_ppr(style), "w:outlineLvl"))


# ====================================================================
# 输出结构
# ====================================================================


def _font_dict(fields: dict[str, Any]) -> dict[str, Any]:
    return {
        "cn": fields.get("font.cn"),
        "ascii": fields.get("font.ascii"),
        "hansi": fields.get("font.hansi"),
        "cs": fields.get("font.cs"),
        "size_pt": fields.get("size_pt"),
        "bold": fields.get("bold"),
        "italic": fields.get("italic"),
        "underline": fields.get("underline"),
        "color": fields.get("color"),
        "highlight": fields.get("highlight"),
        "shading": fields.get("shading"),
        "char_spacing_pt": fields.get("char_spacing_pt"),
        "vert_align": fields.get("vert_align"),
    }


def _paragraph_dict(fields: dict[str, Any]) -> dict[str, Any]:
    line_spacing = None
    if "line_spacing.rule" in fields:
        line_spacing = {"rule": fields["line_spacing.rule"]}
        for key in ("value", "pt", "raw_twips"):
            if f"line_spacing.{key}" in fields:
                line_spacing[key] = fields[f"line_spacing.{key}"]

    indent = {
        key.split(".", 1)[1]: value
        for key, value in fields.items()
        if key.startswith("indent.") and value is not None
    }

    numbering = None
    num_id = fields.get("numbering.num_id")
    if num_id not in (None, 0):
        numbering = {"num_id": num_id}
        if "numbering.level" in fields:
            numbering["level"] = fields["numbering.level"]

    paragraph = {
        "alignment": fields.get("alignment"),
        "line_spacing": line_spacing,
        "space_before_pt": fields.get("space_before_pt"),
        "space_after_pt": fields.get("space_after_pt"),
        "indent": indent or None,
        "keep_next": fields.get("keep_next"),
        "page_break_before": fields.get("page_break_before"),
        "snap_to_grid": fields.get("snap_to_grid"),
        "outline_level": fields.get("outline_level"),
        "numbering": numbering,
    }
    for key in ("space_before_lines", "space_after_lines"):
        if key in fields:
            paragraph[key] = fields[key]
    return paragraph


def _sources_for(p_sources: dict[str, str], r_sources: dict[str, str]) -> dict[str, str]:
    sources: dict[str, str] = {}
    for key, source in r_sources.items():
        if key.endswith(".theme"):
            continue
        sources[key if key.startswith("font.") else f"font.{key}"] = source
    for key, source in p_sources.items():
        sources[f"paragraph.{key}"] = source
    return sources


# ====================================================================
# 上下文与解析
# ====================================================================


class _Context:
    def __init__(self, doc, theme, defaults_rpr, defaults_ppr,
                 *, include_sources: bool, include_raw: bool):
        self.doc = doc
        self.theme = theme
        self.defaults_rpr = defaults_rpr
        self.defaults_ppr = defaults_ppr
        self.include_sources = include_sources
        self.include_raw = include_raw
        self.style_usage: dict[str, int] = {}
        self.numbering_usage: dict[int, dict[str, list]] = {}


def _default_fields(doc, theme: dict[str, str]):
    styles_element = doc.styles.element
    rprs = styles_element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")
    pprs = styles_element.xpath("./w:docDefaults/w:pPrDefault/w:pPr")
    rpr = rprs[0] if rprs else None
    ppr = pprs[0] if pprs else None
    return _r_fields(rpr, theme), _p_fields(ppr), rpr, ppr


def _paragraph_fields(paragraph, ctx: _Context):
    layers = [(paragraph._p.pPr, "direct")]
    style = getattr(paragraph, "style", None)
    if style is not None:
        for item in WordFormatter._style_chain(style):
            layers.append((_style_ppr(item), item.name))
    layers.append((ctx.defaults_ppr, "docDefaults"))
    return _merge_layers(layers, _p_fields)


def _run_fields(run, paragraph, ctx: _Context):
    layers = [(run._r.rPr, "direct")]
    try:
        char_style = run.style
    except Exception:  # noqa: BLE001 - 个别残缺文档没有字符样式
        char_style = None
    if char_style is not None:
        for item in WordFormatter._style_chain(char_style):
            layers.append((_style_rpr(item), item.name))
    style = getattr(paragraph, "style", None)
    if style is not None:
        for item in WordFormatter._style_chain(style):
            layers.append((_style_rpr(item), item.name))
    layers.append((ctx.defaults_rpr, "docDefaults"))
    return _merge_layers(layers, lambda element: _r_fields(element, ctx.theme))


def _style_only_run_fields(paragraph, ctx: _Context):
    layers = []
    style = getattr(paragraph, "style", None)
    if style is not None:
        for item in WordFormatter._style_chain(style):
            layers.append((_style_rpr(item), item.name))
    layers.append((ctx.defaults_rpr, "docDefaults"))
    return _merge_layers(layers, lambda element: _r_fields(element, ctx.theme))


def _dominant_run(paragraph):
    for run in paragraph.runs:
        if run.text.strip():
            return run
    return paragraph.runs[0] if paragraph.runs else None


def _describe_paragraph(paragraph, ctx: _Context):
    p_fields, p_sources = _paragraph_fields(paragraph, ctx)
    run = _dominant_run(paragraph)
    if run is not None:
        r_fields, r_sources = _run_fields(run, paragraph, ctx)
    else:
        r_fields, r_sources = _style_only_run_fields(paragraph, ctx)
    sources = _sources_for(p_sources, r_sources) if ctx.include_sources else {}
    return _font_dict(r_fields), _paragraph_dict(p_fields), sources


def _describe_style(style, ctx: _Context) -> dict[str, Any]:
    chain = WordFormatter._style_chain(style)
    p_layers = [(_style_ppr(item), item.name) for item in chain]
    p_layers.append((ctx.defaults_ppr, "docDefaults"))
    r_layers = [(_style_rpr(item), item.name) for item in chain]
    r_layers.append((ctx.defaults_rpr, "docDefaults"))

    p_fields, p_sources = _merge_layers(p_layers, _p_fields)
    r_fields, r_sources = _merge_layers(
        r_layers, lambda element: _r_fields(element, ctx.theme),
    )

    try:
        based_on = style.base_style.name if style.base_style is not None else None
    except Exception:  # noqa: BLE001
        based_on = None
    try:
        next_style = (
            style.next_paragraph_style.name
            if style.next_paragraph_style is not None else None
        )
    except Exception:  # noqa: BLE001
        next_style = None

    entry: dict[str, Any] = {
        "style_id": style.style_id,
        "based_on": based_on,
        "next": next_style,
        "outline_level": p_fields.get("outline_level"),
        "used_count": ctx.style_usage.get(style.name, 0),
        "font": _font_dict(r_fields),
        "paragraph": _paragraph_dict(p_fields),
    }
    if ctx.include_sources:
        entry["resolved_from"] = _sources_for(p_sources, r_sources)
    if ctx.include_raw:
        entry["raw_xml"] = _xml(style.element)
    return entry


def _collect_style_usage(doc) -> dict[str, int]:
    counts: dict[str, int] = {}

    def add(paragraph):
        style = getattr(paragraph, "style", None)
        name = style.name if style is not None else None
        if name:
            counts[name] = counts.get(name, 0) + 1

    for paragraph in doc.paragraphs:
        add(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    add(paragraph)
    return counts


def _describe_styles(doc, ctx: _Context, *, include_all: bool) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH or not style.name:
            continue
        if (
            not include_all
            and style.name not in ctx.style_usage
            and style.name not in _KEY_STYLE_NAMES
            and _style_outline_level(style) is None
        ):
            continue
        result[style.name] = _describe_style(style, ctx)
    return result


# ====================================================================
# 逐块实测
# ====================================================================


def _role(style_name: str | None, paragraph: dict[str, Any]) -> str:
    outline = paragraph.get("outline_level")
    if isinstance(outline, int) and 0 <= outline <= 8:
        return f"heading_{outline + 1}"
    name = (style_name or "").strip().lower()
    if "toc" in name or "目录" in name:
        return "toc"
    if name in {"title", "标题"}:
        return "title"
    if name in {"subtitle", "副标题"}:
        return "subtitle"
    return "body"


def _signature(font: dict[str, Any], paragraph: dict[str, Any]) -> str:
    return json.dumps(
        {"font": font, "paragraph": paragraph},
        ensure_ascii=False, sort_keys=True, default=str,
    )


def _describe_observed(doc, ctx: _Context, *, max_groups: int, max_samples: int):
    groups: dict[str, dict[str, Any]] = {}
    paragraph_index = 0
    for item in doc.iter_inner_content():
        if isinstance(item, Table):
            continue
        index = paragraph_index
        paragraph_index += 1
        if not item.text.strip():
            continue
        font, paragraph, sources = _describe_paragraph(item, ctx)
        style_name = item.style.name if getattr(item, "style", None) is not None else None
        key = _signature(font, paragraph) + f"|{style_name}"
        group = groups.get(key)
        if group is None:
            group = {
                "role": _role(style_name, paragraph),
                "style": style_name,
                "count": 0,
                "samples": [],
                "font": font,
                "paragraph": paragraph,
            }
            if sources:
                group["resolved_from"] = sources
            groups[key] = group
        group["count"] += 1
        if len(group["samples"]) < max_samples:
            group["samples"].append(index)

    ordered = sorted(groups.values(), key=lambda item: item["samples"][0])
    return ordered[:max_groups]


# ====================================================================
# 表格
# ====================================================================


def _cell_shading(cell) -> str | None:
    tc_pr = cell._tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return None
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        return None
    return shading.get(qn("w:fill"))


def _row_sample(row, ctx: _Context):
    for cell in row.cells:
        if not cell.text.strip() or not cell.paragraphs:
            continue
        font, paragraph, _sources = _describe_paragraph(cell.paragraphs[0], ctx)
        return {
            "font": font,
            "paragraph": paragraph,
            "shading": _cell_shading(cell),
        }
    return None


def _table_borders(tbl_pr) -> dict[str, Any] | None:
    if tbl_pr is None:
        return None
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        return None
    result: dict[str, Any] = {}
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            continue
        value = element.get(qn("w:val"))
        if value in (None, "nil", "none"):
            continue
        size = _float(element.get(qn("w:sz")))
        result[edge] = {
            "style": value,
            "size_pt": round(size / 8, 2) if size is not None else None,
            "color": element.get(qn("w:color")) or None,
        }
    return result or None


def _row_grid(trow) -> dict[int, tuple[int, str | None]]:
    """返回 ``{起始列: (跨列数, vMerge 状态)}``。"""
    grid: dict[int, tuple[int, str | None]] = {}
    column = 0
    for cell in trow.findall(qn("w:tc")):
        tc_pr = cell.find(qn("w:tcPr"))
        span = _int(_child_val(tc_pr, "w:gridSpan")) or 1
        vmerge_element = tc_pr.find(qn("w:vMerge")) if tc_pr is not None else None
        vmerge = None
        if vmerge_element is not None:
            vmerge = vmerge_element.get(qn("w:val")) or "continue"
        grid[column] = (span, vmerge)
        column += span
    return grid


def _table_merges(tbl) -> list[dict[str, int]]:
    rows = tbl.findall(qn("w:tr"))
    grids = [_row_grid(row) for row in rows]
    result: list[dict[str, int]] = []
    for row_index, grid in enumerate(grids):
        for column, (span, vmerge) in grid.items():
            if span == 1 and vmerge != "restart":
                continue
            rowspan = 1
            if vmerge == "restart":
                for following in grids[row_index + 1:]:
                    state = following.get(column, (1, None))[1]
                    if state == "continue":
                        rowspan += 1
                    else:
                        break
            result.append({
                "row": row_index,
                "col": column,
                "rowspan": rowspan,
                "colspan": span,
            })
    return result


def _describe_table(table, ctx: _Context, index: int) -> dict[str, Any]:
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    try:
        style_name = table.style.name if table.style is not None else None
    except Exception:  # noqa: BLE001
        style_name = None

    width_cm = None
    width_element = tbl_pr.find(qn("w:tblW")) if tbl_pr is not None else None
    if width_element is not None and width_element.get(qn("w:type")) in (None, "dxa"):
        width_cm = _twips_to_cm(width_element.get(qn("w:w")))

    grid: list[float | None] = []
    grid_element = tbl.find(qn("w:tblGrid"))
    if grid_element is not None:
        for column in grid_element.findall(qn("w:gridCol")):
            grid.append(_twips_to_cm(column.get(qn("w:w"))))

    conditional_styles: list[str] = []
    try:
        style_element = table.style.element if table.style is not None else None
    except Exception:  # noqa: BLE001
        style_element = None
    if style_element is not None:
        conditional_styles = list(dict.fromkeys(
            element.get(qn("w:type"))
            for element in style_element.findall(qn("w:tblStylePr"))
            if element.get(qn("w:type"))
        ))

    rows = table.rows
    entry: dict[str, Any] = {
        "index": index,
        "style": style_name,
        "rows": len(rows),
        "columns": len(table.columns),
        "alignment": _ALIGNMENT_MAP.get(
            _child_val(tbl_pr, "w:jc"), _child_val(tbl_pr, "w:jc"),
        ),
        "width_cm": width_cm,
        "layout": _child_val(tbl_pr, "w:tblLayout"),
        "column_widths_cm": grid,
        "borders": _table_borders(tbl_pr),
        "header_row": _row_sample(rows[0], ctx) if rows else None,
        "body_row": _row_sample(rows[1], ctx) if len(rows) > 1 else None,
        "merges": _table_merges(tbl),
        "conditional_styles": conditional_styles,
    }
    if ctx.include_raw and style_element is not None:
        entry["conditional_raw_xml"] = [
            _xml(element) for element in style_element.findall(qn("w:tblStylePr"))
        ]
    return entry


def _describe_tables(doc, ctx: _Context) -> list[dict[str, Any]]:
    return [
        _describe_table(table, ctx, index)
        for index, table in enumerate(doc.tables)
    ]


# ====================================================================
# 页面、页眉页脚、编号
# ====================================================================


def _describe_page(doc) -> dict[str, Any]:
    sections = []
    for index, section in enumerate(doc.sections):
        sect_pr = section._sectPr
        pg_sz = sect_pr.find(qn("w:pgSz"))
        pg_mar = sect_pr.find(qn("w:pgMar"))
        doc_grid = sect_pr.find(qn("w:docGrid"))
        cols = sect_pr.find(qn("w:cols"))
        pg_num = sect_pr.find(qn("w:pgNumType"))

        margins = {}
        if pg_mar is not None:
            for attr, key in (
                ("w:top", "top"),
                ("w:bottom", "bottom"),
                ("w:left", "left"),
                ("w:right", "right"),
                ("w:gutter", "gutter"),
            ):
                margins[key] = _twips_to_cm(pg_mar.get(qn(attr)))
            header_cm = _twips_to_cm(pg_mar.get(qn("w:header")))
            footer_cm = _twips_to_cm(pg_mar.get(qn("w:footer")))
        else:
            header_cm = footer_cm = None

        orientation = None
        if pg_sz is not None:
            orientation = pg_sz.get(qn("w:orient"))
        if not orientation:
            width = _length_cm(section.page_width)
            height = _length_cm(section.page_height)
            orientation = "landscape" if (
                width is not None and height is not None and width > height
            ) else "portrait"

        entry: dict[str, Any] = {
            "index": index,
            "orientation": orientation,
            "width_cm": _length_cm(section.page_width),
            "height_cm": _length_cm(section.page_height),
            "margins_cm": margins,
            "header_distance_cm": header_cm,
            "footer_distance_cm": footer_cm,
            "different_first_page": sect_pr.find(qn("w:titlePg")) is not None,
            "horizontal_alignment": _ALIGNMENT_MAP.get(
                _child_val(sect_pr, "w:jc"), _child_val(sect_pr, "w:jc"),
            ),
            "vertical_alignment": _child_val(sect_pr, "w:vAlign"),
            "page_numbering": {
                "format": pg_num.get(qn("w:fmt")) if pg_num is not None else None,
                "start": _int(pg_num.get(qn("w:start"))) if pg_num is not None else None,
            },
        }
        if doc_grid is not None:
            entry["doc_grid"] = {
                "type": doc_grid.get(qn("w:type")),
                "line_pitch_pt": _twips_to_pt(doc_grid.get(qn("w:linePitch"))),
                "char_space": _int(doc_grid.get(qn("w:charSpace"))),
            }
        if cols is not None:
            entry["columns"] = {
                "count": _int(cols.get(qn("w:num"))) or 1,
                "spacing_cm": _twips_to_cm(cols.get(qn("w:space"))),
            }
        sections.append(entry)

    even_and_odd = doc.settings.element.find(qn("w:evenAndOddHeaders")) is not None
    return {"even_and_odd_headers": even_and_odd, "sections": sections}


def _paragraph_text(paragraph) -> str:
    return "".join(WordFormatter._run_text(run) for run in paragraph.runs)


def _paragraph_tabs(paragraph) -> list[dict[str, Any]]:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return []
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        return []
    result = []
    for tab in tabs.findall(qn("w:tab")):
        result.append({
            "pos_cm": _twips_to_cm(tab.get(qn("w:pos"))),
            "align": tab.get(qn("w:val")),
        })
    return result


def _paragraph_field_names(paragraph) -> list[str]:
    found: list[str] = []
    for node in paragraph._p.iter():
        if node.tag == qn("w:fldSimple"):
            found.extend(_FIELD_RE.findall(node.get(qn("w:instr")) or ""))
        elif node.tag == qn("w:instrText"):
            found.extend(_FIELD_RE.findall(node.text or ""))
    return list(dict.fromkeys(name.upper() for name in found))


def _paragraph_images(paragraph) -> list[dict[str, Any]]:
    images: list[dict[str, Any]] = []
    p_element = paragraph._p
    for drawing in p_element.findall(f".//{qn('w:drawing')}"):
        extent = drawing.find(f".//{{{_WP_NS}}}extent")
        anchor = drawing.find(f".//{{{_WP_NS}}}anchor")
        item: dict[str, Any] = {
            "type": "drawing",
            "floating": anchor is not None,
        }
        if extent is not None:
            item["width_cm"] = _emu_to_cm(extent.get("cx"))
            item["height_cm"] = _emu_to_cm(extent.get("cy"))
        images.append(item)
    for _pict in p_element.findall(f".//{qn('w:pict')}"):
        images.append({"type": "vml", "floating": None})
    return images


def _paragraph_border(paragraph, edge: str) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return False
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        return False
    element = borders.find(qn(f"w:{edge}"))
    return element is not None and element.get(qn("w:val")) not in (None, "nil", "none")


def _describe_header_part(part, ctx: _Context, *, kind: str) -> dict[str, Any]:
    entry: dict[str, Any] = {
        "linked_to_previous": bool(part.is_linked_to_previous),
        "paragraphs": [],
        "tables": len(getattr(part, "tables", []) or []),
    }
    if entry["linked_to_previous"]:
        return entry

    border_edge = "bottom" if kind == "header" else "top"
    for paragraph in part.paragraphs:
        fields = _paragraph_field_names(paragraph)
        images = _paragraph_images(paragraph)
        text = _paragraph_text(paragraph)
        if not text.strip() and not fields and not images:
            continue
        font, paragraph_format, _sources = _describe_paragraph(paragraph, ctx)
        runs = []
        for run in paragraph.runs:
            run_text = WordFormatter._run_text(run)
            if not run_text.strip():
                continue
            run_fields, _run_sources = _run_fields(run, paragraph, ctx)
            runs.append({**_font_dict(run_fields), "text": run_text})
        item: dict[str, Any] = {
            "text": text,
            "alignment": paragraph_format.get("alignment"),
            "font": font,
            "runs": runs,
            "tabs": _paragraph_tabs(paragraph),
            "fields": fields,
            "images": images,
        }
        item[f"{border_edge}_border"] = _paragraph_border(paragraph, border_edge)
        if ctx.include_raw:
            item["raw_xml"] = _xml(paragraph._p)
        entry["paragraphs"].append(item)
    return entry


def _describe_headers_footers(doc, ctx: _Context) -> list[dict[str, Any]]:
    variants = (
        ("primary", "header", "footer"),
        ("first", "first_page_header", "first_page_footer"),
        ("even", "even_page_header", "even_page_footer"),
    )
    result = []
    for index, section in enumerate(doc.sections):
        entry: dict[str, Any] = {"section": index, "header": {}, "footer": {}}
        for variant, header_attr, footer_attr in variants:
            entry["header"][variant] = _describe_header_part(
                getattr(section, header_attr), ctx, kind="header",
            )
            entry["footer"][variant] = _describe_header_part(
                getattr(section, footer_attr), ctx, kind="footer",
            )
        result.append(entry)
    return result


def _numbering_level(level_element, ctx: _Context) -> dict[str, Any]:
    p_pr = level_element.find(qn("w:pPr"))
    r_pr = level_element.find(qn("w:rPr"))
    p_fields = _p_fields(p_pr)
    r_fields = _r_fields(r_pr, ctx.theme)
    entry: dict[str, Any] = {
        "level": _int(level_element.get(qn("w:ilvl"))),
        "format": _child_val(level_element, "w:numFmt"),
        "text": _child_val(level_element, "w:lvlText"),
        "start": _int(_child_val(level_element, "w:start")),
        "suffix": _child_val(level_element, "w:suff"),
        "justification": _child_val(level_element, "w:lvlJc"),
        "paragraph_style": _child_val(level_element, "w:pStyle"),
        "paragraph": _paragraph_dict(p_fields),
        "font": _font_dict(r_fields),
    }
    if ctx.include_raw:
        entry["raw_xml"] = _xml(level_element)
    return entry


def _describe_numbering(doc, ctx: _Context) -> dict[str, Any]:
    part = getattr(doc.part, "numbering_part", None)
    if part is None:
        return {"definitions": []}

    root = part.element
    abstracts: dict[int, dict[str, Any]] = {}
    for abstract in root.findall(qn("w:abstractNum")):
        abstract_id = _int(abstract.get(qn("w:abstractNumId")))
        if abstract_id is None:
            continue
        abstracts[abstract_id] = {
            "abstract_id": abstract_id,
            "multi_level_type": _child_val(abstract, "w:multiLevelType"),
            "levels": [
                _numbering_level(level, ctx)
                for level in abstract.findall(qn("w:lvl"))
            ],
        }
        if ctx.include_raw:
            abstracts[abstract_id]["raw_xml"] = _xml(abstract)

    definitions = []
    for num in root.findall(qn("w:num")):
        num_id = _int(num.get(qn("w:numId")))
        abstract_id = _int(_child_val(num, "w:abstractNumId"))
        if num_id is None or abstract_id not in abstracts:
            continue
        base = abstracts[abstract_id]
        levels = {level["level"]: dict(level) for level in base["levels"]}

        for override in num.findall(qn("w:lvlOverride")):
            level_no = _int(override.get(qn("w:ilvl")))
            if level_no is None:
                continue
            replacement = override.find(qn("w:lvl"))
            if replacement is not None:
                levels[level_no] = _numbering_level(replacement, ctx)
            start = _int(_child_val(override, "w:startOverride"))
            if start is not None and level_no in levels:
                levels[level_no]["start"] = start

        usage = ctx.numbering_usage.get(num_id, {"paragraphs": [], "styles": []})
        if not usage.get("paragraphs") and not usage.get("styles"):
            continue
        definition: dict[str, Any] = {
            "num_id": num_id,
            "abstract_id": abstract_id,
            "multi_level_type": base.get("multi_level_type"),
            "used_by": {
                "paragraphs": usage.get("paragraphs", []),
                "styles": usage.get("styles", []),
            },
            "levels": [levels[key] for key in sorted(levels)],
        }
        if ctx.include_raw:
            definition["raw_xml"] = base.get("raw_xml")
        definitions.append(definition)
    return {"definitions": definitions}


def _collect_numbering_usage(doc, ctx: _Context) -> None:
    paragraph_index = 0
    for item in doc.iter_inner_content():
        if isinstance(item, Table):
            continue
        index = paragraph_index
        paragraph_index += 1
        if not item.text.strip():
            continue
        fields, _sources = _paragraph_fields(item, ctx)
        num_id = fields.get("numbering.num_id")
        if num_id in (None, 0):
            continue
        usage = ctx.numbering_usage.setdefault(
            num_id, {"paragraphs": [], "styles": []},
        )
        usage["paragraphs"].append(index)

    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH or not style.name:
            continue
        if style.name not in ctx.style_usage:
            continue
        chain = WordFormatter._style_chain(style)
        layers = [(_style_ppr(item), item.name) for item in chain]
        fields, _sources = _merge_layers(layers, _p_fields)
        num_id = fields.get("numbering.num_id")
        if num_id in (None, 0):
            continue
        usage = ctx.numbering_usage.setdefault(
            num_id, {"paragraphs": [], "styles": []},
        )
        if style.name not in usage["styles"]:
            usage["styles"].append(style.name)


# ====================================================================
# 代码片段
# ====================================================================


def _py_string(value) -> str:
    return repr(str(value))


def _py_number(value) -> str:
    number = round(float(value), 2)
    return repr(int(number)) if number.is_integer() else repr(number)


def _color_tuple(value) -> tuple[int, int, int] | None:
    if not value or value == "auto" or len(value) != 6:
        return None
    try:
        return tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))
    except ValueError:
        return None


def _unsupported_style_items(style: dict[str, Any]) -> list[str]:
    font = style.get("font") or {}
    paragraph = style.get("paragraph") or {}
    items = []
    if font.get("italic"):
        items.append("斜体")
    if font.get("underline"):
        items.append("下划线")
    if font.get("highlight"):
        items.append("高亮")
    if font.get("char_spacing_pt"):
        items.append("字符间距")
    if font.get("vert_align"):
        items.append("上下标")
    if font.get("shading"):
        items.append("底纹")
    if paragraph.get("space_before_lines") or paragraph.get("space_after_lines"):
        items.append("段前/段后按行")
    if paragraph.get("numbering"):
        items.append("自动编号")
    return items


def _style_call(name: str, style: dict[str, Any]) -> tuple[str | None, list[str]]:
    font = style.get("font") or {}
    paragraph = style.get("paragraph") or {}
    arguments: list[str] = []
    if style.get("based_on") and style["based_on"] != "Normal":
        arguments.append(f"base_style_name={_py_string(style['based_on'])}")
    if font.get("cn"):
        arguments.append(f"font_name={_py_string(font['cn'])}")
    if font.get("ascii"):
        arguments.append(f"en_font={_py_string(font['ascii'])}")
    if font.get("size_pt") is not None:
        arguments.append(f"font_size={_py_number(font['size_pt'])}")
    if font.get("bold") is not None:
        arguments.append(f"bold={bool(font['bold'])!r}")
    color = _color_tuple(font.get("color"))
    if color is not None:
        arguments.append(f"color=({color[0]}, {color[1]}, {color[2]})")

    line_spacing = paragraph.get("line_spacing") or {}
    if line_spacing.get("rule") in ("exact", "at_least") and line_spacing.get("pt") is not None:
        arguments.append(f"line_spacing_pt={_py_number(line_spacing['pt'])}")
        if line_spacing.get("rule") == "at_least":
            arguments.append('line_spacing_rule="atLeast"')
    elif line_spacing.get("value") is not None:
        arguments.append(f"line_spacing={_py_number(line_spacing['value'])}")

    if paragraph.get("space_before_pt") is not None:
        arguments.append(f"space_before={_py_number(paragraph['space_before_pt'])}")
    if paragraph.get("space_after_pt") is not None:
        arguments.append(f"space_after={_py_number(paragraph['space_after_pt'])}")

    indent = paragraph.get("indent") or {}
    if indent.get("first_line_chars") is not None:
        arguments.append(
            f"first_line_indent_chars={_py_number(indent['first_line_chars'])}",
        )
    elif indent.get("first_line_pt") is not None:
        arguments.append(f"first_line_indent={_py_number(indent['first_line_pt'])}")
    if indent.get("left_pt") is not None:
        arguments.append(f"left_indent={_py_number(indent['left_pt'])}")
    if indent.get("right_pt") is not None:
        arguments.append(f"right_indent={_py_number(indent['right_pt'])}")
    if paragraph.get("alignment"):
        arguments.append(f"alignment={_py_string(paragraph['alignment'])}")
    if paragraph.get("keep_next") is not None:
        arguments.append(f"keep_next={bool(paragraph['keep_next'])!r}")
    if paragraph.get("page_break_before") is not None:
        arguments.append(
            f"page_break_before={bool(paragraph['page_break_before'])!r}",
        )

    if not arguments:
        return None, _unsupported_style_items(style)

    head = f"formatter.set_paragraph_style({_py_string(name)}"
    if len(head) + sum(len(item) + 2 for item in arguments) < 88:
        return f"{head}, " + ", ".join(arguments) + ")", _unsupported_style_items(style)

    lines = [head + ","]
    lines.extend(f"    {item}," for item in arguments)
    lines.append(")")
    return "\n".join(lines), _unsupported_style_items(style)


def _page_calls(profile: dict[str, Any]) -> list[str]:
    calls = []
    page = profile["document"].get("page") or {}
    for section in page.get("sections") or []:
        index = section.get("index", 0)
        arguments = []
        if section.get("orientation"):
            arguments.append(f"orientation={_py_string(section['orientation'])}")
        for key, name in (
            ("width_cm", "page_width"),
            ("height_cm", "page_height"),
        ):
            if section.get(key) is not None:
                arguments.append(f"{name}={_py_number(section[key])}")
        margins = section.get("margins_cm") or {}
        for key in ("top", "bottom", "left", "right", "gutter"):
            if margins.get(key) is not None:
                arguments.append(f"{key}={_py_number(margins[key])}")
        if section.get("horizontal_alignment"):
            arguments.append(
                f"horizontal_alignment={_py_string(section['horizontal_alignment'])}",
            )
        if arguments:
            calls.append(
                f"formatter.set_section_page(doc.sections[{index}], "
                + ", ".join(arguments) + ")",
            )
        if section.get("different_first_page"):
            calls.append(
                f"formatter.set_different_first_page(doc.sections[{index}], True)",
            )
        page_numbering = section.get("page_numbering") or {}
        if page_numbering.get("format") or page_numbering.get("start") is not None:
            number_format = _py_string(page_numbering.get("format") or "decimal")
            start = page_numbering.get("start")
            extra = f", start={_py_number(start)}" if start is not None else ""
            calls.append(
                f"formatter.set_page_number_format(doc.sections[{index}], "
                f"{number_format}{extra})",
            )
    if page.get("even_and_odd_headers"):
        calls.append("formatter.set_different_odd_even(True)")
    return calls


def _header_footer_calls(profile: dict[str, Any]) -> tuple[list[str], list[str]]:
    calls: list[str] = []
    pending: list[str] = []
    for item in profile["document"].get("headers_footers") or []:
        section_index = item.get("section", 0)
        for kind in ("header", "footer"):
            for variant, data in (item.get(kind) or {}).items():
                if data.get("linked_to_previous"):
                    continue
                paragraphs = data.get("paragraphs") or []
                has_tables = bool(data.get("tables"))
                if not paragraphs:
                    if has_tables:
                        pending.append(f"{kind}/{variant}：含表格")
                    else:
                        calls.append(
                            f"formatter.clear_{kind}(doc.sections[{section_index}], "
                            f"variant={_py_string(variant)})",
                        )
                    continue
                has_rich = has_tables or any(
                    paragraph.get("fields") or paragraph.get("images")
                    for paragraph in paragraphs
                )
                if len(paragraphs) != 1 or has_rich:
                    pending.append(
                        f"{kind}/{variant}：{len(paragraphs)} 段"
                        + ("，含域/图片/表格" if has_rich else "")
                    )
                    continue

                paragraph = paragraphs[0]
                text = paragraph.get("text") or ""
                font = paragraph.get("font") or {}
                arguments = [f"variant={_py_string(variant)}"]
                if "\n" not in text:
                    parts = text.split("\t")
                    if len(parts) == 3:
                        arguments.extend(
                            f"{key}={_py_string(value)}"
                            for key, value in zip(("left", "center", "right"), parts)
                        )
                        method = f"set_{kind}_parts"
                    else:
                        method = f"set_{kind}"
                        arguments.insert(0, f"text={_py_string(text)}")
                        if paragraph.get("alignment"):
                            arguments.append(
                                f"alignment={_py_string(paragraph['alignment'])}",
                            )
                else:
                    pending.append(f"{kind}/{variant}：多行文本")
                    continue

                if font.get("cn"):
                    arguments.append(f"font_name={_py_string(font['cn'])}")
                if font.get("size_pt") is not None:
                    arguments.append(f"font_size={_py_number(font['size_pt'])}")
                border_key = "bottom_border" if kind == "header" else "top_border"
                arguments.append(
                    f"{border_key}={bool(paragraph.get(border_key))!r}"
                )
                calls.append(
                    f"formatter.{method}(doc.sections[{section_index}], "
                    + ", ".join(arguments) + ")",
                )
    return calls, pending


def _numbering_calls(profile: dict[str, Any]) -> list[str]:
    definitions = (profile["document"].get("numbering") or {}).get("definitions") or []
    if not definitions:
        return []
    lines = ["# 自动编号（现有写接口不支持重建，以下为源文档定义，需在 Word 中手工建多级列表）"]
    for definition in definitions:
        numbers = definition.get("levels") or []
        sample = numbers[0] if numbers else {}
        lines.append(
            f"#   numId={definition.get('num_id')} "
            f"format={sample.get('format')} text={sample.get('text')}"
        )
    return lines


def _build_code(profile: dict[str, Any], *, include_tables: bool = False) -> str:
    lines = [
        "# 由 Word 文档提取的格式配置（schema_version=%s）" % profile["schema_version"],
        "from wc_word_report_tool import WordFormatter",
        "",
        "formatter = WordFormatter(doc)",
        "",
        "# 1) 段落样式（生效值；同名样式会被更新）",
    ]

    style_calls = 0
    skipped = 0
    for name, style in (profile.get("styles") or {}).items():
        if style.get("used_count", 0) == 0 and name != "Normal":
            skipped += 1
            continue
        call, unsupported = _style_call(name, style)
        if call:
            lines.append(call)
            style_calls += 1
        if unsupported:
            lines.append(
                f"# 未覆盖（{name}）：{'、'.join(unsupported)}，见 JSON styles[{name!r}]"
            )
    if skipped:
        lines.append(f"# 另有 {skipped} 个未使用的样式仅保留在 JSON 中")
    if not style_calls:
        lines.append("# （没有可生成的样式调用）")

    page_calls = _page_calls(profile)
    if page_calls:
        lines.extend(["", "# 2) 页面设置"])
        lines.extend(page_calls)

    header_calls, pending = _header_footer_calls(profile)
    if header_calls or pending:
        lines.extend(["", "# 3) 页眉页脚（仅纯文本可直接还原）"])
        lines.extend(header_calls)
        for item in pending:
            lines.append(f"# 需手工还原：{item}，见 JSON document.headers_footers")

    if include_tables:
        table_lines = []
        for table in (profile.get("observed") or {}).get("tables") or []:
            index = table.get("index", 0)
            widths = [value for value in table.get("column_widths_cm") or [] if value]
            if widths and len(widths) == table.get("columns"):
                table_lines.append(
                    f"formatter.set_table_columns(doc.tables[{index}], "
                    + "[" + ", ".join(_py_number(value) for value in widths) + "])"
                )
            borders = table.get("borders") or {}
            if borders:
                edge = next(iter(borders.values()))
                color = edge.get("color") or "000000"
                size = edge.get("size_pt")
                size_eighth = int(round(size * 8)) if size is not None else 4
                table_lines.append(
                    f"formatter.set_table_borders(doc.tables[{index}], "
                    f"color={_py_string(color)}, size={size_eighth})"
                )
        if table_lines:
            lines.extend(["", "# 4) 表格（仅当目标文档已有同下标表格）"])
            lines.extend(table_lines)

    numbering_lines = _numbering_calls(profile)
    if numbering_lines:
        lines.extend(["", "# 5) 自动编号"])
        lines.extend(numbering_lines)

    code = "\n".join(lines) + "\n"
    compile(code, "<format-profile>", "exec")
    return code


# ====================================================================
# notes 与入口
# ====================================================================


def _build_notes(profile: dict[str, Any]) -> list[str]:
    notes: list[str] = []
    document = profile["document"]

    if (document.get("numbering") or {}).get("definitions"):
        notes.append(
            "自动编号只读：JSON 保留 numFmt/lvlText/start 与可选原始 XML，代码片段不重建编号。"
        )

    rich_parts = []
    images = 0
    fields = []
    header_tables = 0
    mixed_fonts = 0
    for item in document.get("headers_footers") or []:
        for kind in ("header", "footer"):
            for variant, data in (item.get(kind) or {}).items():
                paragraphs = data.get("paragraphs") or []
                header_tables += data.get("tables", 0)
                if len(paragraphs) > 1:
                    rich_parts.append(f"{kind}/{variant} 多段")
                for paragraph in paragraphs:
                    images += len(paragraph.get("images") or [])
                    fields.extend(paragraph.get("fields") or [])
                    run_styles = {
                        (run.get("cn"), run.get("size_pt"), run.get("bold"))
                        for run in paragraph.get("runs") or []
                    }
                    if len(run_styles) > 1:
                        mixed_fonts += 1
    if images:
        notes.append(f"页眉页脚含 {images} 张图片，需手工还原。")
    if fields:
        notes.append(
            "页眉页脚含域：" + "、".join(sorted(set(fields))) + "，需手工还原。"
        )
    if header_tables:
        notes.append(f"页眉页脚含 {header_tables} 个表格，只报告数量，不展开内容。")
    if mixed_fonts:
        notes.append(
            f"页眉页脚有 {mixed_fonts} 个段落包含多种字体，代码片段只还原段落主字体。"
        )
    if rich_parts:
        notes.append("页眉页脚存在多段结构：" + "、".join(rich_parts) + "。")

    grids = [
        section for section in (document.get("page") or {}).get("sections") or []
        if section.get("doc_grid")
        and (
            section["doc_grid"].get("type") not in (None, "default")
            or section["doc_grid"].get("char_space")
        )
    ]
    if grids:
        notes.append("文档使用了 w:docGrid（行网格），现有接口只读不写，需在 Word 中设置。")

    multi_column = [
        section for section in (document.get("page") or {}).get("sections") or []
        if (section.get("columns") or {}).get("count", 1) > 1
    ]
    if multi_column:
        notes.append("文档使用了分栏，现有接口没有分栏设置方法，需在 Word 中设置。")

    vertical = [
        section for section in (document.get("page") or {}).get("sections") or []
        if section.get("vertical_alignment")
    ]
    if vertical:
        notes.append("页面垂直对齐方式（w:vAlign）只读不写，需在 Word 中设置。")

    unsupported = []
    for name, style in (profile.get("styles") or {}).items():
        if style.get("used_count", 0) <= 0:
            continue
        items = _unsupported_style_items(style)
        if items:
            unsupported.append(f"{name}（{'、'.join(items)}）")
    if unsupported:
        notes.append("样式中有现有接口未覆盖的属性：" + "；".join(unsupported) + "。")

    conditional = [
        table for table in (profile.get("observed") or {}).get("tables") or []
        if table.get("conditional_styles")
    ]
    if conditional:
        notes.append(
            "表格使用了条件样式（首行/镶边等），只报告类型，不级联展开。"
        )

    return notes


def _json_safe(value):
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return round(value, 4)
    if isinstance(value, str):
        return value
    if hasattr(value, "pt"):
        return _length_pt(value)
    return str(value)


def extract_format_profile(
    source,
    *,
    include_observed: bool = True,
    include_code: bool = True,
    include_raw: bool = False,
    include_all_styles: bool = False,
    include_sources: bool = True,
    max_observed_groups: int = 50,
    max_samples: int = 10,
) -> dict[str, Any]:
    """读取 Word 文档格式，返回可复刻的 JSON 配置与代码片段。

    :param source: ``docx.Document`` 对象，或 ``.docx`` 文件路径。
    :param include_observed: 是否逐块实测正文格式并聚类，默认是。
    :param include_code: 是否生成可直接执行的 Python 片段，默认是。
    :param include_raw: 是否附带编号/条件格式/样式的原始 XML，默认否。
    :param include_all_styles: 是否输出全部段落样式；默认只输出「被使用 + 关键样式
                               + 带大纲级别」的样式。
    :param include_sources: 是否输出 ``resolved_from`` 来源标注，默认是。
    :param max_observed_groups: 实测分组上限，默认 50。
    :param max_samples: 每组保留的段落下标数量，默认 10。
    :return: 含 ``schema_version / document / styles / observed / python_code / notes``
             的字典，所有值均可 ``json.dumps``。
    """
    if isinstance(source, (str, Path)):
        doc = Document(str(source))
    elif hasattr(source, "part") and hasattr(source, "paragraphs"):
        doc = source
    else:
        raise TypeError("source 必须是 docx.Document 或 .docx 文件路径")

    theme = _theme_fonts(doc)
    _r_defaults, _p_defaults, rpr, ppr = _default_fields(doc, theme)
    ctx = _Context(
        doc, theme, rpr, ppr,
        include_sources=include_sources, include_raw=include_raw,
    )
    ctx.style_usage = _collect_style_usage(doc)
    _collect_numbering_usage(doc, ctx)

    profile: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "document": {
            "defaults": {
                "font": _font_dict(_r_defaults),
                "paragraph": _paragraph_dict(_p_defaults),
            },
            "page": _describe_page(doc),
            "headers_footers": _describe_headers_footers(doc, ctx),
            "numbering": _describe_numbering(doc, ctx),
        },
        "styles": _describe_styles(doc, ctx, include_all=include_all_styles),
    }

    if include_observed:
        profile["observed"] = {
            "groups": _describe_observed(
                doc, ctx,
                max_groups=max_observed_groups, max_samples=max_samples,
            ),
            "tables": _describe_tables(doc, ctx),
        }

    profile["notes"] = _build_notes(profile)
    if include_code:
        profile["python_code"] = _build_code(profile)
    return _json_safe(profile)
