"""MCP 工具函数。

每个函数都是普通 Python 函数，模块本身不依赖 ``mcp`` SDK，因此可以直接单元测试；
由 :func:`register` 统一注册到 MCPServer 实例。

命名统一加 ``word_`` 前缀：MCP 客户端会把多个 server 的工具名拍平到同一命名空间，
前缀可以避免与其他文档类 server 冲突。

设计分层：

* 会话工具 —— 创建 / 打开 / 保存 / 关闭 / 查看文档
* 语义工具 —— 覆盖 AI 生成报告的绝大多数路径，参数带完整默认值
* 兜底工具 —— ``word_describe_api`` + ``word_call`` 保证对 ``WordFormatter`` 的
  全量覆盖，长尾方法不会因为没被单独包装而不可用
"""

import time
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Union

from docx.text.paragraph import Paragraph

from ..numbers import number_to_chinese_upper
from ..word import (
    DEFAULT_CN_FONT,
    DEFAULT_EN_FONT,
    DEFAULT_HEADER_FOOTER_FONT,
    DEFAULT_HEADER_FOOTER_SIZE,
    WordFormatter,
    describe_font_size,
)
from .api_help import formatter_api_text
from .errors import CONVERTIBLE_ERRORS, ToolError, to_tool_error
from .session import (
    ReportSessionError,
    Session,
    SessionStore,
    default_logo_path,
    materialize_image,
    resolve_input_path,
    resolve_output_path,
)

FontSize = Union[float, str]

_STORE = SessionStore()


def reset_store() -> None:
    """清空全部会话，供测试使用。"""
    _STORE.close_all()


@contextmanager
def _tool_errors():
    """把工具层的可预期异常统一转换成 MCP 可见的 ToolError。"""
    try:
        yield
    except ToolError:
        raise
    except CONVERTIBLE_ERRORS as exc:
        raise to_tool_error(exc) from exc


@contextmanager
def _op(doc_id: str, *, mark_dirty: bool = True) -> Iterator[Session]:
    """取出会话并在全局锁内执行文档改写。

    可预期的异常在这里统一转成 ``ToolError``，否则 mcp SDK 会把消息掩码成
    ``Error executing tool <name>``，AI 就拿不到可自纠的中文提示了。
    """
    with _tool_errors():
        with _STORE.lock:
            session = _STORE.get(doc_id)
            try:
                yield session
            except Exception:
                raise
            else:
                if mark_dirty:
                    session.dirty = True
            finally:
                session.last_used = time.time()


def _result(session: Session, **extra: Any) -> Dict[str, Any]:
    payload = session.describe()
    payload.update(extra)
    events = _STORE.drain_events()
    if events:
        payload["maintenance_events"] = events
    return payload


def _section(session: Session, index: int):
    sections = session.doc.sections
    resolved = index + len(sections) if index < 0 else index
    if resolved < 0 or resolved >= len(sections):
        raise ReportSessionError(
            f"section_index {index} 超出范围"
            f"（共 {len(sections)} 节，0~{len(sections) - 1}）"
        )
    return sections[resolved]


def _table(session: Session, index: int):
    tables = session.doc.tables
    if not tables:
        raise ReportSessionError("文档中还没有表格，请先调用 word_add_table。")
    resolved = index + len(tables) if index < 0 else index
    if resolved < 0 or resolved >= len(tables):
        raise ReportSessionError(
            f"table_index {index} 超出范围"
            f"（共 {len(tables)} 个表格，0~{len(tables) - 1}）"
        )
    return tables[resolved]


def _paragraph(session: Session, index: int):
    paragraphs = session.doc.paragraphs
    if not paragraphs:
        raise ReportSessionError("文档中还没有正文段落。")
    resolved = index + len(paragraphs) if index < 0 else index
    if resolved < 0 or resolved >= len(paragraphs):
        raise ReportSessionError(
            f"paragraph_index {index} 超出范围"
            f"（共 {len(paragraphs)} 段，0~{len(paragraphs) - 1}）"
        )
    return paragraphs[resolved]


def _resolve_logo(
    session: Session,
    image_path: Optional[str],
    image_base64: Optional[str],
    label: str = "Logo 图片",
) -> Path:
    if image_path or image_base64:
        return materialize_image(
            session, image_path=image_path, image_base64=image_base64, label=label,
        )
    fallback = default_logo_path()
    if fallback is not None:
        return fallback
    raise ReportSessionError(
        f"缺少{label}：请提供 image_path 或 image_base64，"
        "或在服务端配置环境变量 WC_REPORT_MCP_DEFAULT_LOGO 指向一个默认 Logo。"
    )


# ====================================================================
# Layer 1 · 会话生命周期
# ====================================================================


def word_create_report(
    title: Optional[str] = None,
    author: Optional[str] = None,
    apply_default_margins: bool = True,
    language: Optional[str] = "zh-CN",
) -> Dict[str, Any]:
    """新建一份空白 Word 报告，返回后续所有操作都要带上的 doc_id。

        :param title: 文档标题，写入 docx 核心属性。
        :param author: 作者，写入 docx 核心属性。
        :param apply_default_margins: 是否套用公文默认页边距（上 3.7 / 下 3.5 / 左 2.8 / 右 2.6 cm）。
        :param language: 文档校对语言，默认 zh-CN；传 None 则不设置。

    AI 调用说明：
    适用：从空白开始生成新报告，或需要与现有文档隔离的新会话。
    不要用于：已有 .docx 模板的填写；这种情况用 word_open_report。
    调用前后：先创建，后续所有 Word 工具都带返回的 doc_id；结束时保存并关闭。
    失败处理：参数非法会抛异常；创建失败时不要猜测或复用旧 doc_id。
    """
    with _tool_errors():
        session = _STORE.create()
        formatter = session.formatter
        if apply_default_margins:
            formatter.setup_defaults()
        if language:
            formatter.set_document_language(language)
        if title:
            session.doc.core_properties.title = title
        if author:
            session.doc.core_properties.author = author
        return _result(session)


def word_open_report(path: str) -> Dict[str, Any]:
    """打开一份已存在的 .docx 继续编辑，返回新的 doc_id。

        :param path: 文件路径；相对路径先按当前目录解析，再按 WC_REPORT_MCP_OUTPUT_DIR 解析。

    AI 调用说明：
    适用：打开现有 .docx 后原位填写或检查格式。
    不要用于：需要全新空白文档时；这种情况用 word_create_report。
    调用前后：通常在 word_list_blocks / word_describe_format_profile 之前，操作完成后 save 和 close。
    失败处理：路径不存在或无权限会失败；确认绝对路径和文件可读后重试。
    """
    with _tool_errors():
        resolved = resolve_input_path(path, what="Word 文档")
        session = _STORE.create(source_path=resolved)
        return _result(session, opened=str(resolved))


def word_save_report(
    doc_id: str,
    path: Optional[str] = None,
    refresh_fields_on_open: bool = True,
) -> Dict[str, Any]:
    """把当前文档保存为 .docx 并返回落盘路径。

        :param doc_id: 会话 id。
        :param path: 保存路径；相对路径挂在 WC_REPORT_MCP_OUTPUT_DIR（默认 ~/wc-reports）下。
                     省略时复用上次保存路径，从未保存过则自动生成带时间戳的文件名。
        :param refresh_fields_on_open: 是否让 Word 打开文件时自动刷新目录与页码域。

    AI 调用说明：
    适用：把当前会话落盘、交付检查或关闭前保存。
    不要用于：只想临时查看结构；查看用 word_describe_report。
    调用前后：保存不会关闭会话，之后仍可继续修改；最终结束时再 close。
    失败处理：目标路径无权限或扩展名错误会失败；换可写路径重试，不要丢弃内存中的会话。
    """
    with _op(doc_id, mark_dirty=False) as session:
        if path:
            target = resolve_output_path(path)
        elif session.saved_path is not None:
            target = session.saved_path
        else:
            target = resolve_output_path(None)
        session.save(target, refresh_fields_on_open=refresh_fields_on_open)
        return _result(session, saved_to=str(target), size_bytes=target.stat().st_size)


def word_close_report(doc_id: str) -> Dict[str, Any]:
    """释放会话占用的内存。

        :param doc_id: 会话 id。关闭后该 doc_id 立即失效。

    AI 调用说明：
    适用：明确结束某个文档会话，释放内存。
    不要用于：还需要继续修改或保存同一文档时。
    调用前后：先确保已保存或确认自动保存结果，再关闭；关闭后 doc_id 失效。
    失败处理：未知 doc_id 会失败；不要继续使用已关闭会话，必要时重新 open。
    """
    with _tool_errors():
        return _STORE.close(doc_id)


def word_describe_report(doc_id: str) -> Dict[str, Any]:
    """查看当前文档的结构摘要（段落数、节数、表格数及其形状）。

        :param doc_id: 会话 id。

    AI 调用说明：
    适用：快速查看段落、节、表格数量和保存状态，做会话诊断。
    不要用于：需要逐段文本、run 下标或下划线信息；这种情况用 word_list_blocks。
    调用前后：可在写入前后调用；本工具不修改文档。
    失败处理：doc_id 无效时重新 open；不要根据失败结果推断文档为空。
    """
    with _op(doc_id, mark_dirty=False) as session:
        payload = _result(session)
        # python-docx 的 Section 没有 paragraphs 属性，只能按节遍历正文内容
        payload["paragraphs_per_section"] = [
            sum(
                1
                for item in section.iter_inner_content()
                if isinstance(item, Paragraph)
            )
            for section in session.doc.sections
        ]
        payload["table_shapes"] = [
            {"rows": len(table.rows), "columns": len(table.columns)}
            for table in session.doc.tables
        ]
        return payload


# ====================================================================
# Layer 2 · 文档级设置
# ====================================================================


def word_set_default_font(
    doc_id: str,
    font_size: FontSize = 14,
    cn_font: str = DEFAULT_CN_FONT,
    en_font: str = DEFAULT_EN_FONT,
) -> Dict[str, Any]:
    """设置文档 Normal 样式的默认字号与中英文字体。

        :param doc_id: 会话 id。
        :param font_size: 字号，支持数字（pt）或中文字号字符串（三号/小四/小五…）。
        :param cn_font: 中文字体名。注意：字体必须在运行本服务端的机器上已安装。
        :param en_font: 西文字体名。

    AI 调用说明：
    适用：在开始追加正文前设置整份文档的 Normal 默认字体字号。
    不要用于：修改已有 run 的显式格式；已有内容替换用 word_replace_text。
    调用前后：通常紧跟 create/open 之后、添加正文之前调用。
    失败处理：字号或字体参数非法会失败；修正参数后重试，不要假定部分内容已改。
    """
    with _op(doc_id) as session:
        session.formatter.set_default_font(
            font_size=font_size, cn_font=cn_font, en_font=en_font,
        )
        return _result(session)


def word_set_page_margins(
    doc_id: str,
    top: float = 3.7,
    bottom: float = 3.5,
    left: float = 2.8,
    right: float = 2.6,
    gutter: float = 0,
    horizontal_alignment: str = "L",
) -> Dict[str, Any]:
    """设置全文页边距（单位 cm）。

        :param doc_id: 会话 id。
        :param top: 上边距，单位 cm，默认 3.7。
        :param bottom: 下边距，单位 cm，默认 3.5。
        :param left: 左边距，单位 cm，默认 2.8。
        :param right: 右边距，单位 cm，默认 2.6。
        :param gutter: 装订线宽度，单位 cm。
        :param horizontal_alignment: 页面水平对齐，支持 L/C/R 或 左对齐/居中/右对齐。

    AI 调用说明：
    适用：统一设置全文页边距和装订线。
    不要用于：只调整某一节的页面；这种情况用 word_set_section_page。
    调用前后：建议在排版开始前调用；后续插入节可能继承当前页面设置。
    失败处理：越界尺寸会失败；根据异常修正 cm 值，不要传负值。
    """
    with _op(doc_id) as session:
        session.formatter.set_page_margins(
            top=top, bottom=bottom, left=left, right=right,
            gutter=gutter, horizontal_alignment=horizontal_alignment,
        )
        return _result(session)


def word_set_document_language(doc_id: str, lang: str = "zh-CN") -> Dict[str, Any]:
    """设置文档校对语言，避免 Word/OnlyOffice 按英文处理中文排版。

        :param doc_id: 会话 id。
        :param lang: BCP 47 语言标签，默认 zh-CN。

    AI 调用说明：
    适用：让 Word/OnlyOffice 按中文校对和排版。
    不要用于：修改 OnlyOffice 编辑器界面语言或区域格式；本工具只改文档语言。
    调用前后：可在 create/open 后尽早调用。
    失败处理：不支持的语言代码会失败；使用 BCP 47 形式如 zh-CN。
    """
    with _op(doc_id) as session:
        session.formatter.set_document_language(lang)
        return _result(session)


# ====================================================================
# Layer 2 · 段落与标题
# ====================================================================


def word_add_heading(
    doc_id: str,
    text: str,
    level: int = 1,
    font_name: str = "黑体",
    font_size: FontSize = "三号",
    bold: bool = False,
    indent: bool = False,
    alignment: Optional[str] = None,
    line_spacing: float = 1.5,
    line_spacing_pt: float = 0,
    left_indent_pt: float = 0,
    space_before: Optional[float] = None,
    space_after: Optional[float] = None,
    page_break_before: bool = False,
) -> Dict[str, Any]:
    """添加标题段落。

        :param doc_id: 会话 id。
        :param text: 标题文本。
        :param level: 标题层级 1~9，默认 1。层级会写入 Word 大纲级别，也是目录的收集依据。
        :param font_name: 中文字体名，默认黑体。
        :param font_size: 字号，默认三号。
        :param bold: 是否加粗。
        :param indent: 是否首行缩进 2 字符。
        :param alignment: 对齐方式，省略时沿用标题样式的居左。做居中的封面大标题时传 "居中"。
        :param line_spacing: 行距倍数，默认 1.5。
        :param space_before: 段前空白（pt）；省略时按字号的 0.5 行计算。
        :param space_after: 段后空白（pt）；省略时按字号的 0.5 行计算。
        :param page_break_before: 是否在本标题之前分页。模板类文档「一个表单起一页」时用它；
            该属性在标题已处于页首时自动失效，不会多顶出空白页。

    AI 调用说明：
    适用：按顺序生成新报告的标题段落。
    不要用于：已有模板中的标题替换；这种情况用 word_replace_text。
    调用前后：调用前确认当前插入位置和标题级别，后续正文会追加在它之后。
    失败处理：字体或行距参数非法会失败；修正后重试，不要用正文工具替代标题。
    """
    with _op(doc_id) as session:
        session.formatter.heading(
            text, level=level, font_name=font_name, font_size=font_size,
            bold=bold, indent=indent, line_spacing=line_spacing,
            line_spacing_pt=line_spacing_pt, left_indent_pt=left_indent_pt,
            alignment=alignment, space_before=space_before, space_after=space_after,
            page_break_before=page_break_before,
        )
        return _result(session)


def word_add_body(
    doc_id: str,
    text: str,
    font_name: str = "宋体",
    font_size: FontSize = "四号",
    indent: bool = True,
    bold: bool = False,
    alignment: str = "两端对齐",
    line_spacing: float = 1.5,
    line_spacing_pt: float = 0,
    left_indent_pt: float = 0,
    space_before: float = 0,
    space_after: float = 0,
    page_break_before: bool = False,
    highlight: bool = False,
    underline: bool = False,
) -> Dict[str, Any]:
    """添加一个正文段落（默认宋体四号、首行缩进 2 字符、两端对齐、1.5 倍行距）。

        :param doc_id: 会话 id。
        :param text: 正文内容。
        :param font_name: 中文字体名。
        :param font_size: 字号，默认四号。
        :param indent: 是否首行缩进 2 字符。
        :param bold: 是否加粗。
        :param alignment: 对齐方式，支持 两端对齐/居中/居左/居右 或 left/center/right/justify。
        :param line_spacing: 行距倍数。
        :param line_spacing_pt: 行距固定值，单位 pt；给了它覆盖 line_spacing。
        :param left_indent_pt: 整段左缩进，单位 pt，默认 0。用来还原「整段右移」的表单行
                              （如「联合体牵头人名称：____（盖单位公章）」），它既不是居中，
                              也不是首行缩进。
                                从 PDF 还原版式时用它——PDF 只能量出「相邻行基线差多少 pt」，
                                换算成倍数要依赖字体自身行高，直接写 pt 才精确。
        :param space_before: 段前空白，单位 pt。
        :param space_after: 段后空白，单位 pt。
        :param page_break_before: 是否在本段之前分页。
        :param highlight: 是否黄色高亮，默认否。用于标出由 AI 填写的临时数据。
        :param underline: 整段是否加下划线，默认否。用于还原模板里的填空线；
                          只对段落一部分划线时请改用 word_add_body_segments。

    AI 调用说明：
    适用：从空白生成报告时追加一个正文段落。
    不要用于：填现有模板；追加会跑到文档末尾，应改用 word_replace_text。
    调用前后：按阅读顺序调用；批量正文优先用 word_add_body_list。
    失败处理：参数非法会失败；若位置写错不能靠重复调用修复，应重新打开或调整内容。
    """
    with _op(doc_id) as session:
        session.formatter.body(
            text, font_name=font_name, font_size=font_size, indent=indent,
            bold=bold, alignment=alignment, line_spacing=line_spacing,
            line_spacing_pt=line_spacing_pt, left_indent_pt=left_indent_pt,
            space_before=space_before, space_after=space_after,
            page_break_before=page_break_before, highlight=highlight,
            underline=underline,
        )
        return _result(session)


def word_add_body_list(
    doc_id: str,
    texts: List[str],
    font_name: str = "宋体",
    font_size: FontSize = "四号",
    indent: bool = True,
    bold: bool = False,
    alignment: str = "两端对齐",
    line_spacing: float = 1.5,
    line_spacing_pt: float = 0,
    left_indent_pt: float = 0,
    space_before: float = 0,
    space_after: float = 0,
    page_break_before: bool = False,
    highlight: bool = False,
    underline: bool = False,
) -> Dict[str, Any]:
    """批量添加多个正文段落，共用同一套格式参数（长报告建议用它减少调用次数）。

        :param doc_id: 会话 id。
        :param texts: 正文段落列表，按顺序逐段写入。
        :param page_break_before: 只对**第一段**生效，是否在它之前分页。
        :param highlight: 是否把这几段整体黄色高亮，默认否。
        :param underline: 是否把这几段整体加下划线，默认否。
        其余参数含义同 word_add_body。

    AI 调用说明：
    适用：从空白生成报告时批量追加多段正文。
    不要用于：模板原位填写或多段需要不同格式的场景。
    调用前后：一次传入完整段落数组；后续工具继续按文档顺序追加。
    失败处理：空数组或非法格式会失败；修正数组后重试，不要重复追加已成功的部分。
    """
    with _op(doc_id) as session:
        for index, text in enumerate(texts):
            session.formatter.body(
                text, font_name=font_name, font_size=font_size, indent=indent,
                bold=bold, alignment=alignment, line_spacing=line_spacing,
                line_spacing_pt=line_spacing_pt, left_indent_pt=left_indent_pt,
                space_before=space_before, space_after=space_after,
                page_break_before=page_break_before and index == 0,
                highlight=highlight, underline=underline,
            )
        return _result(session, added=len(texts))


def word_add_body_segments(
    doc_id: str,
    segments: List[Dict[str, Any]],
    font_name: str = "宋体",
    font_size: FontSize = "四号",
    indent: bool = True,
    alignment: str = "两端对齐",
    line_spacing: float = 1.5,
    line_spacing_pt: float = 0,
    left_indent_pt: float = 0,
    page_break_before: bool = False,
) -> Dict[str, Any]:
    """写一个「由多段拼接」的正文段落，每段可单独高亮、单独加下划线。

        典型用法是「标签 + 填值」只标记填值部分：
        `[{"text": "项目名称：", "highlight": false},
          {"text": "雄安新区示范工程", "highlight": true}]`
        只把填进去的值标黄，标签保持原样。

        招投标模板里的填空线同理，只画在填进去的值下面：
        `[{"text": "中国雄安集团生态建设投资有限公司", "underline": true},
          {"text": "（采购人名称）"}]`

        :param doc_id: 会话 id。
        :param segments: 片段列表，每项形如
                         {"text": "…", "highlight": true, "underline": true}；
                         highlight / underline 都可省略。
        其余参数含义同 word_add_body。

    AI 调用说明：
    适用：一个段落里需要多种高亮或下划线片段。
    不要用于：普通纯文本正文；用 word_add_body 更简单。
    调用前后：按 segments 顺序写入，后续段落继续追加。
    失败处理：片段结构非法会失败；确保每项是元组或含 text 的字典。
    """
    with _op(doc_id) as session:
        session.formatter.body_segments(
            segments, font_name=font_name, font_size=font_size, indent=indent,
            alignment=alignment, line_spacing=line_spacing,
            line_spacing_pt=line_spacing_pt, left_indent_pt=left_indent_pt,
            page_break_before=page_break_before,
        )
        return _result(session, segments=len(segments))


def word_add_blank_lines(doc_id: str, count: int = 1) -> Dict[str, Any]:
    """添加若干空行（封面、签发页等留白用）。

        :param doc_id: 会话 id。
        :param count: 空行数量，默认 1。

    AI 调用说明：
    适用：在封面、签发页等位置追加视觉留白。
    不要用于：精确定位分页；分页用 word_add_page_break。
    调用前后：按当前文档末尾位置追加。
    失败处理：count 非正数会失败；使用正整数重试。
    """
    with _op(doc_id) as session:
        session.formatter.blank_lines(count)
        return _result(session, added=count)


def word_add_page_break(doc_id: str, text: str = "") -> Dict[str, Any]:
    """插入一个分页符（可同时带一段文字），用于「一个表单起一页」这类版式。

        模板类文档的分页通常是固定排版的一部分；只靠内容自然流排的话页数会对不上原稿。

        :param doc_id: 会话 id。
        :param text: 分页符之后要写的文字，默认空（纯分页）。

    AI 调用说明：
    适用：在报告追加流程中开始新页，可附带标题文字。
    不要用于：已有模板中替换分页或定位内容。
    调用前后：在希望新页开始的位置调用；之后内容从新页追加。
    失败处理：参数非法会失败；不要用多个空段落模拟分页。
    """
    with _op(doc_id) as session:
        session.formatter.add_page_break(text)
        return _result(session)


def word_add_cover_text(
    doc_id: str,
    text: str,
    font_name: str = "方正小标宋简体",
    font_size: FontSize = 24,
    bold: bool = True,
    alignment: str = "居中",
    line_spacing: float = 2,
    space_before: Optional[float] = None,
    space_after: Optional[float] = None,
) -> Dict[str, Any]:
    """添加封面文本（默认居中加粗、2 倍行距）。

        :param doc_id: 会话 id。
        :param text: 封面文字。
        :param font_name: 中文字体名。
        :param font_size: 字号。
        :param bold: 是否加粗。
        :param alignment: 对齐方式。
        :param line_spacing: 行距倍数。
        :param space_before: 段前空白（pt）；省略时按字号的 1 行计算。
        :param space_after: 段后空白（pt）；省略时按字号的 1 行计算。

    AI 调用说明：
    适用：从空白生成封面时追加封面文字。
    不要用于：填写已有封面模板；这种情况用 word_replace_text。
    调用前后：按封面行顺序调用，通常先 project/title 再日期和落款。
    失败处理：非法字号或对齐会失败；修正后重试。
    """
    with _op(doc_id) as session:
        session.formatter.cover_text(
            text, font_name=font_name, font_size=font_size, bold=bold,
            alignment=alignment, line_spacing=line_spacing,
            space_before=space_before, space_after=space_after,
        )
        return _result(session)


def word_add_right_text(
    doc_id: str,
    text: str,
    font_name: str = DEFAULT_CN_FONT,
    font_size: FontSize = 14,
    indent: bool = True,
) -> Dict[str, Any]:
    """添加右对齐段落（公司名、落款等）。

        :param doc_id: 会话 id。
        :param text: 段落文本。
        :param font_name: 中文字体名。
        :param font_size: 字号。
        :param indent: 是否首行缩进 2 字符。

    AI 调用说明：
    适用：追加右对齐的公司名、落款或日期行。
    不要用于：填写模板中的现有落款空位。
    调用前后：放在正文和签发内容之后按顺序追加。
    失败处理：参数非法会失败；检查对齐和字号后重试。
    """
    with _op(doc_id) as session:
        session.formatter.right_text(
            text, font_name=font_name, font_size=font_size, indent=indent,
        )
        return _result(session)


def word_add_date(
    doc_id: str,
    value: Optional[str] = None,
    style: str = "chinese",
    font_name: str = DEFAULT_CN_FONT,
    font_size: FontSize = 14,
    bold: bool = False,
    alignment: str = "居中",
) -> Dict[str, Any]:
    """添加日期文本。

        :param doc_id: 会话 id。
        :param value: 日期，支持 2026年07月16日 / 2026-07-16 / 2026/07/16 / 2026.07.16；省略时用当天。
        :param style: chinese = 二〇二六年七月十六日；year_month = 二〇二六年七月；
                      numeric = 2026年07月16日（右对齐）。
        :param font_name: 中文字体名。
        :param font_size: 字号。
        :param bold: 是否加粗。
        :param alignment: 对齐方式（numeric 样式固定右对齐，该参数不生效）。

    AI 调用说明：
    适用：向新报告追加中文、年月或数字格式日期。
    不要用于：替换模板中的日期占位；用 word_replace_text。
    调用前后：按封面或落款顺序调用。
    失败处理：无法解析日期或未知 style 会失败；改用受支持日期或 chinese/year_month/numeric。
    """
    with _op(doc_id) as session:
        formatter = session.formatter
        if style == "chinese":
            formatter.chinese_date(
                value, font_name=font_name, font_size=font_size,
                bold=bold, alignment=alignment,
            )
        elif style == "year_month":
            formatter.chinese_year_month(
                value, font_name=font_name, font_size=font_size,
                bold=bold, alignment=alignment,
            )
        elif style == "numeric":
            formatter.created_time(
                value, font_name=font_name, font_size=font_size, bold=bold,
            )
        else:
            raise ReportSessionError(
                f"未知的日期样式 {style!r}，可选 chinese / year_month / numeric。"
            )
        return _result(session, style=style)


def word_insert_image(
    doc_id: str,
    width_cm: Optional[float] = None,
    height_cm: Optional[float] = None,
    image_path: Optional[str] = None,
    image_base64: Optional[str] = None,
    alignment: str = "居左",
    floating: bool = False,
    y_offset_pt: float = 0,
) -> Dict[str, Any]:
    """插入一张图片到正文。

        ``width_cm`` 与 ``height_cm`` 都不传时按图片自身的 DPI 使用原始尺寸，
        因此从 PDF/扫描件裁切出来的图片可以直接插入，不需要先量一个宽度。

        :param doc_id: 会话 id。
        :param width_cm: 图片宽度，单位 cm；省略时按图片原始尺寸。
        :param height_cm: 图片高度，单位 cm；省略时按图片原始尺寸。
        :param image_path: 图片路径（与 image_base64 二选一）。
        :param image_base64: 图片的 base64，支持裸串或 data:image/png;base64,... 形式。
        :param alignment: 段落对齐方式。
        :param floating: 是否让图片浮于文字上方。印章要压在落款文字上时传 True，
                         从 PDF 裁出的图表、公式同理。
        :param y_offset_pt: 浮动图片的垂直偏移，单位 pt，正数上移、负数下移。

    AI 调用说明：
    适用：把图片追加到正文末尾，例如独立插图、公章浮层。
    不要用于：贴入表格单元格；身份证、营业执照等用 word_insert_cell_image。
    调用前后：图片参数使用 image_path 或 image_base64，二选一；后续内容继续追加。
    失败处理：文件不存在、base64 无效或尺寸非法会失败；修正图片输入后重试。
    """
    with _op(doc_id) as session:
        path = materialize_image(
            session, image_path=image_path, image_base64=image_base64, label="正文图片",
        )
        session.formatter.insert_img(
            str(path), width_cm, height=height_cm, alignment=alignment,
            floating=floating, y_offset_pt=y_offset_pt,
        )
        return _result(
            session,
            width_cm=width_cm,
            height_cm=height_cm,
            auto_sized=width_cm is None and height_cm is None,
            floating=floating,
        )


# ====================================================================
# Layer 2 · 表格
# ====================================================================


def word_add_table(
    doc_id: str,
    headers: List[str],
    rows: List[List[Any]],
    col_widths: Optional[List[float]] = None,
    font_size: FontSize = 12,
    font_name: str = DEFAULT_CN_FONT,
    header_bold: bool = True,
    alignment: str = "居中",
    border_color: str = "000000",
    border_size: int = 4,
    merges: Optional[List[Dict[str, int]]] = None,
    row_heights: Optional[List[Optional[float]]] = None,
    table_width_cm: Optional[float] = None,
    cell_font_names: Optional[List[Optional[List[Optional[str]]]]] = None,
    cell_highlights: Optional[List[Optional[List[Optional[bool]]]]] = None,
) -> Dict[str, Any]:
    """在末尾追加一个带表头的表格。

        :param doc_id: 会话 id。
        :param headers: 表头文本列表，例如 ["序号", "项目", "金额（万元）"]。
        :param rows: 二维数据列表，每个子列表是一行。
        :param col_widths: 每列宽度列表，单位 cm，可选。
        :param font_size: 表格字号。
        :param font_name: 单元格中文字体名。
        :param header_bold: 表头是否加粗。
        :param alignment: 单元格对齐方式。
        :param border_color: 边框颜色，6 位十六进制 RGB。
        :param border_size: 边框粗细，单位 1/8 pt（4 = 0.5pt，8 = 1pt）。
        :param merges: 合并区域列表，例如
                       [{"row": 0, "col": 0, "rowspan": 1, "colspan": 4}]。
                       行列索引从 0 开始、按完整网格计数；合并区域的内容取左上角单元格的值。
        :param row_heights: 每行高度，单位 cm，可选；按「最小值」规则设置，内容超出会自动撑高。
        :param table_width_cm: 表格总宽度，单位 cm，可选。
        :param cell_font_names: 逐单元格中文字体名的二维列表（含表头行），可选；
                                未覆盖的位置回落到 font_name。
        :param cell_highlights: 逐单元格「是否黄色高亮」的二维布尔列表（含表头行），可选。
                                用来标出由 AI 填写的临时数据，便于人工复核。

    AI 调用说明：
    适用：从空白生成报告时在末尾追加新表格。
    不要用于：修改已有表格；已有表格用 merge/format_cell，图片用 insert_cell_image。
    调用前后：先准备 headers/rows 和列宽，再按顺序追加后续内容。
    失败处理：列宽数量不匹配或行列非法会失败；修正行列结构后重试。
    """
    with _op(doc_id) as session:
        table = session.formatter.add_table(
            headers, rows, col_widths=col_widths, font_size=font_size,
            font_name=font_name, header_bold=header_bold, alignment=alignment,
            border_color=border_color, border_size=border_size,
            merges=merges, row_heights=row_heights,
            table_width_cm=table_width_cm, cell_font_names=cell_font_names,
            cell_highlights=cell_highlights,
        )
        return _result(
            session,
            table_index=len(session.doc.tables) - 1,
            rows=len(table.rows),
            columns=len(table.columns),
        )


def word_merge_cells(
    doc_id: str,
    merges: List[Dict[str, int]],
    table_index: int = -1,
) -> Dict[str, Any]:
    """合并指定表格里的单元格区域（对已存在的表格操作）。

        在 ``word_add_table`` 之后需要补合并时用它；新建表格时也可以直接在
        ``word_add_table`` 里传 ``merges``。

        :param doc_id: 会话 id。
        :param merges: 合并区域列表，例如
                       [{"row": 0, "col": 0, "rowspan": 1, "colspan": 3}]。
                       行列索引从 0 开始、按完整网格计数。
        :param table_index: 表格索引，默认 -1 表示文档中最后一个表格。

    AI 调用说明：
    适用：对已有表格执行单元格合并。
    不要用于：创建新表格或填写单元格文字。
    调用前后：先确认 table_index 和坐标范围，合并前完成需要独立保留的单元格写入。
    失败处理：越界或重叠非法会失败；不要假定部分合并已成功，必要时重新打开文档。
    """
    with _op(doc_id) as session:
        table = _table(session, table_index)
        WordFormatter.merge_cells(table, merges)
        return _result(session, table_index=table_index, merged=len(merges or []))


def word_format_cell(
    doc_id: str,
    row: int,
    col: int,
    text: str,
    table_index: int = -1,
    font_name: str = DEFAULT_CN_FONT,
    font_size: FontSize = 12,
    bold: bool = False,
    alignment: str = "居中",
    line_spacing: float = 1.25,
    highlight: bool = False,
) -> Dict[str, Any]:
    """改写指定表格单元格的内容与格式。

        :param doc_id: 会话 id。
        :param row: 行索引，从 0 开始。
        :param col: 列索引，从 0 开始。
        :param text: 新的单元格文本（会清空原有内容）。
        :param table_index: 表格索引，默认 -1 表示文档中最后一个表格。
        :param font_name: 中文字体名。
        :param font_size: 字号。
        :param bold: 是否加粗。
        :param alignment: 对齐方式。
        :param line_spacing: 行距倍数。

    AI 调用说明：
    适用：改写指定单元格的文字和字体格式。
    不要用于：向单元格贴图；用 word_insert_cell_image。
    调用前后：先确认 table_index/row/col，单元格原内容会被覆盖。
    失败处理：坐标越界会失败；修正坐标后重试，重要内容先备份或读取。
    """
    with _op(doc_id) as session:
        table = _table(session, table_index)
        row_count, col_count = len(table.rows), len(table.columns)
        if row < 0 or row >= row_count or col < 0 or col >= col_count:
            raise ReportSessionError(
                f"单元格 ({row}, {col}) 超出范围（该表共 {row_count} 行 × {col_count} 列）。"
            )
        WordFormatter.set_cell(
            table.cell(row, col), text, font_name=font_name, font_size=font_size,
            bold=bold, alignment=alignment, line_spacing=line_spacing,
            highlight=highlight,
        )
        return _result(session, table_index=table_index, row=row, col=col)


def word_set_table_borders(
    doc_id: str,
    table_index: int = -1,
    color: str = "000000",
    size: int = 4,
) -> Dict[str, Any]:
    """为指定表格添加四边及内部横竖边框。

        :param doc_id: 会话 id。
        :param table_index: 表格索引，默认 -1 表示最后一个表格。
        :param color: 边框颜色，6 位十六进制 RGB。
        :param size: 边框粗细，单位 1/8 pt。

    AI 调用说明：
    适用：统一设置已有表格四边和内部边框。
    不要用于：填充内容或合并单元格。
    调用前后：通常在表格结构和内容确定后调用。
    失败处理：颜色或粗细非法会失败；使用 6 位十六进制颜色和非负粗细。
    """
    with _op(doc_id) as session:
        WordFormatter.set_table_borders(
            _table(session, table_index), color=color, size=size,
        )
        return _result(session, table_index=table_index)


# ====================================================================
# Layer 2 · 模板填充（在已有文档上原位改写）
# ====================================================================


def word_list_blocks(
    doc_id: str,
    include_empty: bool = False,
    max_text_chars: int = 200,
) -> Dict[str, Any]:
    """列出文档的正文块（段落 / 表格），用来定位占位符在哪一段、哪一格。

        生成报告时不需要它——按顺序 add 就行；但**填模板**必须先把已有文档看清楚：
        段落块返回它在 ``doc.paragraphs`` 里的下标与逐 run 的字体字号，表格块返回
        ``doc.tables`` 的下标与单元格文字网格，两者分别对应 word_replace_text 的
        ``paragraph_index`` 和 ``table_index`` / ``row`` / ``col``。

        :param doc_id: 会话 id。
        :param include_empty: 是否连空段落一起返回，默认否（模板留白通常不用管）。
        :param max_text_chars: 单段/单格文本的截断长度，默认 200，防止长文档回传过大。

    AI 调用说明：
    适用：填模板前读取段落、表格、run 下标、字体、加粗和下划线状态。
    不要用于：直接修改文档；本工具只读。
    调用前后：通常紧跟 word_open_report，之后据此选择 replace_text 或 fill_underline。
    失败处理：doc_id 无效时重新 open；不要凭猜测构造 paragraph_index/run_index。
    """
    with _op(doc_id, mark_dirty=False) as session:
        blocks = session.formatter.describe_blocks(
            include_empty=include_empty, max_text_chars=max_text_chars,
        )
        return _result(session, blocks=blocks, block_count=len(blocks))


def word_describe_format_profile(
    doc_id: str,
    include_observed: bool = True,
    include_code: bool = True,
    include_raw: bool = False,
    max_groups: int = 50,
) -> Dict[str, Any]:
    """读取文档格式，返回标题/正文等可复刻的配置与 Python 代码。

        用于「照着一份现成 .docx 复刻版式」：给出标题一/标题二/正文的字体、字号、
        行距、首行缩进、左右缩进、段前段后、对齐，以及页面设置、页眉页脚、
        自动编号和表格格式。返回 ``format_profile`` 与 ``schema_version``。

        自动编号、页眉页脚中的域/图片、文档网格等现有接口无法重建的部分，
        会在 ``format_profile["notes"]`` 中列出。

        :param doc_id: 会话 id。
        :param include_observed: 是否逐块实测并聚类正文格式，默认是。
        :param include_code: 是否生成可直接执行的 Python 片段，默认是。
        :param include_raw: 是否附带编号/条件格式的原始 XML，默认否。
        :param max_groups: 实测分组上限，默认 50，避免长文档回传过大。

    AI 调用说明：
    适用：读取现成文档的完整版式配置，准备复刻格式。
    不要用于：填写字段或快速定位单个占位符；后者用 word_list_blocks。
    调用前后：只读；先 open，读取后可按 python_code 或工具调用复刻。
    失败处理：复杂域无法解析时会通过 notes 提示；不要假定缺失项已恢复。
    """
    with _op(doc_id, mark_dirty=False) as session:
        profile = session.formatter.describe_format_profile(
            include_observed=include_observed,
            include_code=include_code,
            include_raw=include_raw,
            max_observed_groups=max_groups,
        )
        return _result(
            session,
            format_profile=profile,
            schema_version=profile.get("schema_version"),
        )


def word_replace_text(
    doc_id: str,
    old: str,
    new: str,
    paragraph_index: Optional[int] = None,
    table_index: Optional[int] = None,
    row: Optional[int] = None,
    col: Optional[int] = None,
    count: int = 1,
    highlight: bool = False,
) -> Dict[str, Any]:
    """把已有内容里的 ``old`` 原位替换为 ``new``，**格式沿用原处**。

        这是填模板的主工具：它只改写原有 run 的文本，不新建段落，所以填完的字体、字号、
        加粗与原文一致。用 ``word_add_body`` 之类的追加型工具会把内容写到文档末尾，
        填模板时用错位置就全乱了。

        定位方式二选一：

        * ``paragraph_index``：正文段落，下标取自 word_list_blocks 段落块的 ``index``；
        * ``table_index`` + ``row`` + ``col``：表格单元格，坐标取自同一份 blocks。

        ``old`` 传空串表示在段尾追加 ``new``，用于「委托代理人身份证号码：____」这类
        只有标签、值要补在标签后面的填空。

        :param doc_id: 会话 id。
        :param old: 要被替换掉的原文；空串表示追加到段尾。
        :param new: 新写入的文本。
        :param paragraph_index: 正文段落下标，支持负数（-1 表示最后一段）；与表格坐标二选一。
        :param table_index: 表格下标，-1 表示最后一个表格；须与 row/col 同时给出。
        :param row: 单元格行号，从 0 开始。
        :param col: 单元格列号，从 0 开始。
        :param count: 最多替换几处，默认 1；传 0 表示全部替换。
        :param highlight: 是否把填进去的内容标黄，便于人工复核，默认否。

    AI 调用说明：
    适用：替换可见文字占位，或向段尾追加普通文本。
    不要用于：只由空格组成且带下划线的空位；这种情况必须用 word_fill_underline。
    调用前后：先 list_blocks 获取 paragraph_index 或表格坐标；old 必须来自原文，old="" 表示段尾追加。
    失败处理：位置越界或 old 未找到会失败；重新读取 blocks 核对原文，不要反复盲试。
    """
    with _op(doc_id) as session:
        if paragraph_index is not None:
            paragraphs = [_paragraph(session, paragraph_index)]
            location = {"paragraph_index": paragraph_index}
        elif table_index is not None and row is not None and col is not None:
            table = _table(session, table_index)
            row_count, col_count = len(table.rows), len(table.columns)
            if row < 0 or row >= row_count or col < 0 or col >= col_count:
                raise ReportSessionError(
                    f"单元格 ({row}, {col}) 超出范围（该表共 {row_count} 行 × {col_count} 列）。"
                )
            paragraphs = list(table.cell(row, col).paragraphs)
            location = {"table_index": table_index, "row": row, "col": col}
        else:
            raise ReportSessionError(
                "需要指定位置：paragraph_index，或同时给 table_index / row / col。"
                "可先用 word_list_blocks 查看文档里有哪些段落和单元格。"
            )

        if old:
            target = next((item for item in paragraphs if old in item.text), None)
        else:
            target = paragraphs[0] if paragraphs else None
        if target is None:
            raise ReportSessionError(
                f"指定位置里没有找到 {old!r}。请先用 word_list_blocks 核对原文"
                "（Word 常把一句话拆成多个 run，这里按整段文本匹配，不受拆分影响）。"
            )
        replaced = WordFormatter.replace_text(
            target, old, new, count=count, highlight=highlight,
        )
        return _result(session, replaced=replaced, old=old, new=new, **location)


def word_fill_underline(
    doc_id: str,
    run_index: int,
    new: str,
    paragraph_index: Optional[int] = None,
    table_index: Optional[int] = None,
    row: Optional[int] = None,
    col: Optional[int] = None,
    cell_paragraph_index: int = 0,
    trailing_spaces: Optional[int] = None,
) -> Dict[str, Any]:
    """填写段落里由带下划线空格组成的固定空位。

        ``word_list_blocks`` 的段落 run 现在会返回 ``run_index`` 和 ``underline``。
        找到只由空格组成且 ``underline=true`` 的 run 后调用本工具，方法会把值放进
        该 run，并自动保留一个前导空格和补足原宽度的尾随空格。

        :param doc_id: 会话 id。
        :param run_index: 段内 run 下标。
        :param new: 要填入的业务值，不要自己附加空格或下划线。
    :param paragraph_index: 正文段落下标；与表格坐标二选一。
    :param table_index/row/col: 表格单元格坐标。
    :param cell_paragraph_index: 单元格内段落下标，默认 0。
    :param trailing_spaces: 尾随下划线空格数；None 表示按原空位宽度补齐。

    AI 调用说明：
    适用：填写只由空格组成且 underline=true 的下划线空位。
    不要用于：包含可见文字的普通占位；这种情况用 word_replace_text。
    调用前后：先用 word_list_blocks 获取 paragraph_index 和 run_index；new 只传业务值，不要自己补空格。
    失败处理：run 非空白或未加下划线会失败；重新读取 blocks，确认 run_index 后重试。
    """
    with _op(doc_id) as session:
        if paragraph_index is not None:
            paragraph = _paragraph(session, paragraph_index)
            location = {"paragraph_index": paragraph_index}
        elif table_index is not None and row is not None and col is not None:
            table = _table(session, table_index)
            row_count, col_count = len(table.rows), len(table.columns)
            if row < 0 or row >= row_count or col < 0 or col >= col_count:
                raise ReportSessionError(
                    f"单元格 ({row}, {col}) 超出范围（该表共 {row_count} 行 × {col_count} 列）。"
                )
            paragraphs = table.cell(row, col).paragraphs
            if cell_paragraph_index < 0 or cell_paragraph_index >= len(paragraphs):
                raise ReportSessionError(
                    f"cell_paragraph_index {cell_paragraph_index} 超出单元格段落范围"
                )
            paragraph = paragraphs[cell_paragraph_index]
            location = {
                "table_index": table_index,
                "row": row,
                "col": col,
                "cell_paragraph_index": cell_paragraph_index,
            }
        else:
            raise ReportSessionError(
                "word_fill_underline 需要 paragraph_index，或 table_index/row/col。"
            )
        result = WordFormatter.fill_underline(
            paragraph,
            run_index,
            new,
            trailing_spaces=trailing_spaces,
        )
        return _result(
            session,
            paragraph_text=paragraph.text,
            new=str(new or "").strip(),
            **location,
            **result,
        )


def word_replace_underline_text(
    doc_id: str,
    paragraph_index: int,
    run_index: int,
    old: str,
    new: str,
) -> Dict[str, Any]:
    """替换带下划线 run 里的可见占位文字。

    用于 `（招标项目名称、招标编号）` 这类本身有文字、同时带下划线格式的占位。
    与 word_fill_underline 的区别是：本工具要求 run 中有可见的 old 文字。

    AI 调用说明：
    适用：下划线 run 内包含可见占位词，需要保留下划线原位替换。
    不要用于：只由空格组成的下划线空位；这种情况用 word_fill_underline。
    调用前后：先用 word_list_blocks 找到 paragraph_index/run_index，再传原文 old。
    失败处理：run 无下划线或 old 未找到会失败；重新读取 blocks 后修正参数。
    """
    with _op(doc_id) as session:
        paragraph = _paragraph(session, paragraph_index)
        result = WordFormatter.replace_underline_text(
            paragraph, run_index, old, new
        )
        return _result(
            session,
            paragraph_index=paragraph_index,
            paragraph_text=paragraph.text,
            **result,
        )


def word_insert_cell_image(
    doc_id: str,
    row: int,
    col: int,
    table_index: int = -1,
    image_path: Optional[str] = None,
    image_base64: Optional[str] = None,
    width_cm: Optional[float] = None,
    height_cm: Optional[float] = None,
    alignment: str = "居中",
    keep_text: bool = False,
) -> Dict[str, Any]:
    """把图片放进指定表格单元格（身份证、营业执照、印章这类「贴格子」场景）。

        ``word_insert_image`` 只能把图片追加到正文，落不进一个已有的格子；而表单里的证件照
        恰恰要贴在「此格附身份证正面」那个格子里，所以需要这个工具。
        ``width_cm`` / ``height_cm`` 都省略时按图片自身 DPI 使用原始尺寸。

        :param doc_id: 会话 id。
        :param row: 单元格行号，从 0 开始。
        :param col: 单元格列号，从 0 开始。
        :param table_index: 表格下标，默认 -1 表示文档中最后一个表格。
        :param image_path: 图片路径（与 image_base64 二选一）。
        :param image_base64: 图片的 base64，支持裸串或 data:image/png;base64,... 形式。
        :param width_cm: 图片宽度，单位 cm，可选。
        :param height_cm: 图片高度，单位 cm，可选。
        :param alignment: 单元格内对齐方式，默认居中。
        :param keep_text: 是否保留格子里已有的文字，默认否（清空后再贴图）。

    AI 调用说明：
    适用：把证件照、营业执照、印章等图片贴进已有表格单元格。
    不要用于：向正文追加独立图片；正文图片用 word_insert_image。
    调用前后：先确认 table_index/row/col；默认清空格内文字，需要保留时设 keep_text=true。
    失败处理：图片无效或坐标越界会失败；修正图片或坐标后重试。
    """
    with _op(doc_id) as session:
        table = _table(session, table_index)
        row_count, col_count = len(table.rows), len(table.columns)
        if row < 0 or row >= row_count or col < 0 or col >= col_count:
            raise ReportSessionError(
                f"单元格 ({row}, {col}) 超出范围（该表共 {row_count} 行 × {col_count} 列）。"
            )
        path = materialize_image(
            session, image_path=image_path, image_base64=image_base64,
            label="单元格图片",
        )
        WordFormatter.insert_cell_image(
            table.cell(row, col), str(path), width=width_cm, height=height_cm,
            alignment=alignment, keep_text=keep_text,
        )
        return _result(
            session, table_index=table_index, row=row, col=col,
            width_cm=width_cm, height_cm=height_cm,
        )


# ====================================================================
# Layer 2 · 页眉页脚
# ====================================================================


def word_set_header(
    doc_id: str,
    text: str,
    section_index: int = -1,
    variant: str = "primary",
    alignment: str = "居中",
    font_name: str = DEFAULT_HEADER_FOOTER_FONT,
    font_size: FontSize = DEFAULT_HEADER_FOOTER_SIZE,
    bottom_border: bool = True,
    line_length: Optional[float] = None,
    line_alignment: str = "居左",
) -> Dict[str, Any]:
    """设置某节的页眉文字（默认楷体小五号，页眉默认带下划线）。

        :param doc_id: 会话 id。
        :param text: 页眉文本。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param variant: 页眉变体：primary（默认/奇数页）、first（首页）、even（偶数页）。
                        使用 first 前需先用 word_set_different_first_page 打开本节开关；
                        使用 even 前需先用 word_set_different_odd_even 打开全文开关。
        :param alignment: 文本对齐方式。
        :param font_name: 中文字体名。
        :param font_size: 字号。
        :param bottom_border: 是否给页眉段落加下横线。
        :param line_length: 横线长度，单位 cm；省略时使用页面可用宽度。
        :param line_alignment: 横线对齐方式。

    AI 调用说明：
    适用：设置某节单一文本页眉。
    不要用于：需要左右中分栏页眉时；用 word_set_header_parts。
    调用前后：先确认 section_index，后续页眉设置可覆盖当前值。
    失败处理：节越界或图片/字体参数错误会失败；修正节下标后重试。
    """
    with _op(doc_id) as session:
        section = _section(session, section_index)
        WordFormatter.set_header(
            section, text, variant=variant, alignment=alignment, font_name=font_name,
            font_size=font_size, bottom_border=bottom_border,
            line_length=line_length, line_alignment=line_alignment,
        )
        return _result(session, section_index=section_index, variant=variant)


def word_set_footer(
    doc_id: str,
    text: str,
    section_index: int = -1,
    variant: str = "primary",
    alignment: str = "居中",
    font_name: str = DEFAULT_HEADER_FOOTER_FONT,
    font_size: FontSize = DEFAULT_HEADER_FOOTER_SIZE,
    top_border: bool = True,
    line_length: Optional[float] = None,
    line_alignment: str = "居左",
) -> Dict[str, Any]:
    """设置某节的页脚文字（默认楷体小五号，页脚默认带上横线）。

        :param doc_id: 会话 id。
        :param text: 页脚文本。已有的页码域会被保留。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param variant: 页脚变体：primary（默认/奇数页）、first（首页）、even（偶数页）。
        :param alignment: 文本对齐方式。
        :param font_name: 中文字体名。
        :param font_size: 字号。
        :param top_border: 是否给页脚段落加上横线。
        :param line_length: 横线长度，单位 cm。
        :param line_alignment: 横线对齐方式。

    AI 调用说明：
    适用：设置某节单一文本页脚。
    不要用于：需要左右中分栏页脚时；用 word_set_footer_parts。
    调用前后：先确认 section_index；页码域应使用页码工具。
    失败处理：节越界会失败；修正 section_index 后重试。
    """
    with _op(doc_id) as session:
        section = _section(session, section_index)
        WordFormatter.set_footer(
            section, text, variant=variant, alignment=alignment, font_name=font_name,
            font_size=font_size, top_border=top_border,
            line_length=line_length, line_alignment=line_alignment,
        )
        return _result(session, section_index=section_index, variant=variant)


def word_set_header_parts(
    doc_id: str,
    left: str = "",
    center: str = "",
    right: str = "",
    section_index: int = -1,
    variant: str = "primary",
    font_name: str = DEFAULT_HEADER_FOOTER_FONT,
    font_size: FontSize = DEFAULT_HEADER_FOOTER_SIZE,
    bottom_border: bool = True,
    line_length: Optional[float] = None,
    line_alignment: str = "居左",
) -> Dict[str, Any]:
    """在一行内分左、中、右三段设置页眉（中文报告最常见的页眉形式）。

        用制表位实现：左段贴版心左边界、中段居中、右段贴版心右边界。

        :param doc_id: 会话 id。
        :param left: 左段文本，如公司名。
        :param center: 中段文本，如报告名。
        :param right: 右段文本，如页码或密级。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param variant: primary / first / even。
        :param font_name: 中文字体名。
        :param font_size: 字号。
        :param bottom_border: 是否给页眉段落加下横线。
        :param line_length: 横线长度，单位 cm。
        :param line_alignment: 横线对齐方式。

    AI 调用说明：
    适用：页眉需要左、中、右三段内容或页码组合。
    不要用于：单段简单页眉；用 word_set_header。
    调用前后：先确认 section_index 和 variant，再设置各段内容。
    失败处理：参数或节越界会失败；修正后重试。
    """
    with _op(doc_id) as session:
        WordFormatter.set_header_parts(
            _section(session, section_index),
            left=left, center=center, right=right, variant=variant,
            font_name=font_name, font_size=font_size, bottom_border=bottom_border,
            line_length=line_length, line_alignment=line_alignment,
        )
        return _result(session, section_index=section_index, variant=variant)


def word_set_footer_parts(
    doc_id: str,
    left: str = "",
    center: str = "",
    right: str = "",
    section_index: int = -1,
    variant: str = "primary",
    font_name: str = DEFAULT_HEADER_FOOTER_FONT,
    font_size: FontSize = DEFAULT_HEADER_FOOTER_SIZE,
    top_border: bool = True,
    line_length: Optional[float] = None,
    line_alignment: str = "居左",
) -> Dict[str, Any]:
    """在一行内分左、中、右三段设置页脚；参数含义同 word_set_header_parts。

        :param top_border: 是否给页脚段落加上横线。

    AI 调用说明：
    适用：页脚需要左、中、右三段内容或页码组合。
    不要用于：单段简单页脚；用 word_set_footer。
    调用前后：先确认 section_index 和 variant，再设置各段内容。
    失败处理：参数或节越界会失败；修正后重试。
    """
    with _op(doc_id) as session:
        WordFormatter.set_footer_parts(
            _section(session, section_index),
            left=left, center=center, right=right, variant=variant,
            font_name=font_name, font_size=font_size, top_border=top_border,
            line_length=line_length, line_alignment=line_alignment,
        )
        return _result(session, section_index=section_index, variant=variant)


def word_set_different_first_page(
    doc_id: str,
    enabled: bool = True,
    section_index: int = -1,
) -> Dict[str, Any]:
    """设置某节首页是否使用独立的页眉页脚（w:titlePg）。

        打开后即可用 variant="first" 单独设置首页页眉页脚，常见于「首页不显示页码」。

        :param doc_id: 会话 id。
        :param enabled: 是否启用独立首页页眉页脚。
        :param section_index: 节索引，默认 -1 表示最后一节。

    AI 调用说明：
    适用：首页需要独立页眉页脚（如封面不显示页眉）。
    不要用于：普通首页与后续页完全相同的场景。
    调用前后：先设置开关，再分别设置首页和默认页眉页脚。
    失败处理：节越界会失败；修正 section_index 后重试。
    """
    with _op(doc_id) as session:
        WordFormatter.set_different_first_page(
            _section(session, section_index), enabled,
        )
        return _result(session, section_index=section_index, enabled=enabled)


def word_set_different_odd_even(
    doc_id: str,
    enabled: bool = True,
) -> Dict[str, Any]:
    """设置全文奇偶页是否使用不同的页眉页脚（文档级 w:evenAndOddHeaders）。

        打开后 variant="even" 的页眉页脚才会生效。注意这是**文档级**开关，
        一旦打开，未单独设置偶页页眉的节会显示为空页眉。

        :param doc_id: 会话 id。
        :param enabled: 是否启用奇偶页不同。

    AI 调用说明：
    适用：全文奇偶页需要不同页眉页脚。
    不要用于：仅单节设置；本工具是文档级开关。
    调用前后：开启后分别设置奇数页和偶数页内容。
    失败处理：参数非布尔会失败；使用 true/false 后重试。
    """
    with _op(doc_id) as session:
        session.formatter.set_different_odd_even(enabled)
        return _result(session, enabled=enabled)


def word_clear_header_footer(
    doc_id: str,
    target: str = "header",
    section_index: int = -1,
    variant: str = "primary",
) -> Dict[str, Any]:
    """清空某节的页眉或页脚。

        :param doc_id: 会话 id。
        :param target: header / footer / both。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param variant: primary / first / even。

    AI 调用说明：
    适用：清空某节页眉或页脚内容。
    不要用于：只想替换文字；直接设置新页眉页脚即可。
    调用前后：先确认 target、section_index 和 variant。
    失败处理：目标或节不存在会失败；修正参数后重试。
    """
    with _op(doc_id) as session:
        if target not in ("header", "footer", "both"):
            raise ReportSessionError(
                f"target 只能是 header/footer/both，收到 {target!r}。"
            )
        section = _section(session, section_index)
        if target in ("header", "both"):
            WordFormatter.clear_header(section, variant=variant)
        if target in ("footer", "both"):
            WordFormatter.clear_footer(section, variant=variant)
        return _result(session, section_index=section_index,
                       variant=variant, cleared=target)


def word_set_header_image(
    doc_id: str,
    target: str = "header",
    section_index: int = -1,
    variant: str = "primary",
    width_cm: Optional[float] = None,
    height_cm: Optional[float] = None,
    image_path: Optional[str] = None,
    image_base64: Optional[str] = None,
    alignment: str = "居中",
    floating: bool = True,
    bottom_border: bool = True,
    y_offset_pt: float = 0,
) -> Dict[str, Any]:
    """设置页眉（或页脚）的图片；已有文字会保留，已有图片会被替换。

        :param doc_id: 会话 id。
        :param target: header 或 footer。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param variant: primary / first / even。
        :param width_cm: 图片宽度，单位 cm（与 height_cm 至少传一个）。
        :param height_cm: 图片高度，单位 cm。
        :param image_path: 图片路径（与 image_base64 二选一）。
        :param image_base64: 图片的 base64。
        :param alignment: 对齐方式。
        :param floating: 是否浮于文字上方（Word 兼容性更好）。
        :param bottom_border: 是否加下横线。
        :param y_offset_pt: 垂直偏移，单位 pt，正数上移、负数下移。

    AI 调用说明：
    适用：页眉或页脚需要 Logo/图片。
    不要用于：正文图片或单元格图片。
    调用前后：先确认节和 variant；已有文字默认保留，已有图片会被替换。
    失败处理：图片无效或节越界会失败；修正图片和节参数后重试。
    """
    with _op(doc_id) as session:
        if target not in ("header", "footer"):
            raise ReportSessionError(
                f"target 只能是 header/footer，收到 {target!r}。"
            )
        section = _section(session, section_index)
        path = materialize_image(
            session, image_path=image_path, image_base64=image_base64,
            label="页眉页脚图片",
        )
        setter = (
            WordFormatter.set_header_image
            if target == "header"
            else WordFormatter.set_footer_image
        )
        setter(
            section, str(path), variant=variant, width=width_cm, height=height_cm,
            alignment=alignment, floating=floating, bottom_border=bottom_border,
            y_offset_pt=y_offset_pt,
        )
        return _result(session, section_index=section_index,
                       variant=variant, target=target)


# ====================================================================
# Layer 2 · 节与页码
# ====================================================================


def word_insert_section(
    doc_id: str,
    add_page_number: bool = False,
    restart_page_number: bool = False,
    start_page_number: int = 1,
    inherit_header: bool = True,
    inherit_footer: bool = True,
    prefix: str = "",
    suffix: str = "",
    alignment: str = "居中",
    number_format: Optional[str] = None,
    orientation: Optional[str] = None,
    page_width_cm: Optional[float] = None,
    page_height_cm: Optional[float] = None,
    margin_top_cm: Optional[float] = None,
    margin_bottom_cm: Optional[float] = None,
    margin_left_cm: Optional[float] = None,
    margin_right_cm: Optional[float] = None,
) -> Dict[str, Any]:
    """插入一个新节（默认从新页开始）。

        横向插页与纵向正文混排时，用 orientation / margin_* 直接给本节设定纸张与边距，
        不必先插节再单独调用 word_set_section_page。

        :param doc_id: 会话 id。
        :param add_page_number: 是否在本节设置页码。
        :param restart_page_number: 是否从 start_page_number 重新编号。
        :param start_page_number: 重新编号的起始页码。
        :param inherit_header: 是否继承上一节页眉。
        :param inherit_footer: 是否继承上一节页脚。
        :param prefix: 页码前缀，如 "第 "。
        :param suffix: 页码后缀，如 " 页"。
        :param alignment: 页码对齐方式。
        :param number_format: 页码编号格式，如 "upperRoman"/"大写罗马"/"chineseCounting"。
        :param orientation: 本节纸张方向，"纵向"/"横向" 或 portrait/landscape。
        :param page_width_cm/page_height_cm: 本节纸张宽高，单位 cm，可选。
        :param margin_top_cm/margin_bottom_cm/margin_left_cm/margin_right_cm: 本节页边距，单位 cm，可选。

    AI 调用说明：
    适用：需要新节、横向插页、独立页码或不同页眉页脚。
    不要用于：只需普通分页；普通分页用 word_add_page_break。
    调用前后：先规划节属性，再在新节中添加内容并设置页码/页眉页脚。
    失败处理：页面或节参数非法会失败；修正参数，不要假定部分节已建立。
    """
    with _op(doc_id) as session:
        session.formatter.insert_section(
            add_page_number=add_page_number,
            restart_page_number=restart_page_number,
            start_page_number=start_page_number,
            inherit_header=inherit_header, inherit_footer=inherit_footer,
            prefix=prefix, suffix=suffix, alignment=alignment,
            number_format=number_format, orientation=orientation,
            page_width=page_width_cm, page_height=page_height_cm,
            top=margin_top_cm, bottom=margin_bottom_cm,
            left=margin_left_cm, right=margin_right_cm,
        )
        return _result(session, section_index=len(session.doc.sections) - 1)


def word_set_section_page(
    doc_id: str,
    section_index: int = -1,
    orientation: Optional[str] = None,
    page_width_cm: Optional[float] = None,
    page_height_cm: Optional[float] = None,
    margin_top_cm: Optional[float] = None,
    margin_bottom_cm: Optional[float] = None,
    margin_left_cm: Optional[float] = None,
    margin_right_cm: Optional[float] = None,
    gutter_cm: Optional[float] = None,
    horizontal_alignment: Optional[str] = None,
) -> Dict[str, Any]:
    """单独设置某一节的纸张方向、尺寸与页边距。

        word_set_page_margins 会把边距应用到全文所有节；本节级设置有需要时用这个工具。

        :param doc_id: 会话 id。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param orientation: "纵向"/"横向" 或 portrait/landscape；与当前方向不同时自动交换宽高。
        :param page_width_cm/page_height_cm: 纸张宽高，单位 cm，可选。
        :param margin_top_cm/margin_bottom_cm/margin_left_cm/margin_right_cm: 页边距，单位 cm，可选。
        :param gutter_cm: 装订线宽度，单位 cm，可选。
        :param horizontal_alignment: 页面水平对齐，支持 L/C/R 或 左对齐/居中/右对齐。

    AI 调用说明：
    适用：单独调整某节纸张方向、尺寸和页边距。
    不要用于：统一全文页面；全文用 word_set_page_margins。
    调用前后：先确认 section_index，再设置该节内容和页眉页脚。
    失败处理：越界尺寸或节下标会失败；修正参数后重试。
    """
    with _op(doc_id) as session:
        section = _section(session, section_index)
        session.formatter.set_section_page(
            section, orientation=orientation,
            page_width=page_width_cm, page_height=page_height_cm,
            top=margin_top_cm, bottom=margin_bottom_cm,
            left=margin_left_cm, right=margin_right_cm,
            gutter=gutter_cm, horizontal_alignment=horizontal_alignment,
        )
        return _result(
            session,
            section_index=section_index,
            page_width_cm=round(section.page_width.cm, 2),
            page_height_cm=round(section.page_height.cm, 2),
        )


def word_set_page_numbers(
    doc_id: str,
    start_section_index: int = 0,
    start: int = 1,
    prefix: str = "",
    suffix: str = "",
    alignment: str = "居中",
    font_name: Optional[str] = None,
    font_size: Optional[FontSize] = None,
    number_format: Optional[str] = None,
) -> Dict[str, Any]:
    """从指定节开始编页码，之前的节不带页码（跨节页码的主入口）。

        典型用法：封面/目录在前置节（无页码），正文从第 1 节开始编号为"第 1 页"。

        :param doc_id: 会话 id。
        :param start_section_index: 从 0 开始的节索引，页码从此节开始。
        :param start: 起始页码。
        :param prefix: 页码前缀，如 "第 "。
        :param suffix: 页码后缀，如 " 页"。
        :param alignment: 页码对齐方式。
        :param font_name: 页码字体名，省略时用默认中文字体。
        :param font_size: 页码字号，省略时用默认正文字号。
        :param number_format: 页码编号格式，如 "upperRoman"/"大写罗马"、
                              "chineseCounting"/"中文数字"；省略时用阿拉伯数字。

    AI 调用说明：
    适用：从指定节开始建立跨节页码，之前的节不带页码。
    不要用于：只改已有页码格式；用 word_set_page_number_format。
    调用前后：先确认起始节，再设置后续节是否继续或重启。
    失败处理：节越界或格式非法会失败；修正 section_index/number_format 后重试。
    """
    with _op(doc_id) as session:
        session.formatter.set_page_number_from_section(
            start_section_index, start=start, prefix=prefix, suffix=suffix,
            alignment=alignment, font_name=font_name, font_size=font_size,
            number_format=number_format,
        )
        return _result(session, start_section_index=start_section_index)


def word_restart_page_numbering(
    doc_id: str,
    section_index: int = -1,
    start: int = 1,
    add_footer_number: bool = True,
    prefix: str = "",
    suffix: str = "",
    alignment: str = "居中",
    number_format: Optional[str] = None,
) -> Dict[str, Any]:
    """重启某节的页码编号。

        :param doc_id: 会话 id。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param start: 起始页码。
        :param add_footer_number: 是否同时在页脚渲染页码。
        :param prefix: 页码前缀。
        :param suffix: 页码后缀。
        :param alignment: 页码对齐方式。
        :param number_format: 页码编号格式，如 "upperRoman"/"大写罗马"；省略时沿用原格式。

    AI 调用说明：
    适用：某节页码需要从指定值重新开始。
    不要用于：只设置起始值但不添加页码域；用 word_set_page_number_start。
    调用前后：先确认节和起始值，后续节可继续编号。
    失败处理：节越界会失败；修正参数后重试。
    """
    with _op(doc_id) as session:
        session.formatter.restart_page_numbering(
            _section(session, section_index), start=start,
            add_footer_number=add_footer_number, prefix=prefix, suffix=suffix,
            alignment=alignment, number_format=number_format,
        )
        return _result(session, section_index=section_index)


def word_set_page_number_start(
    doc_id: str,
    start: int = 1,
    section_index: int = -1,
    number_format: Optional[str] = None,
) -> Dict[str, Any]:
    """设置某节页码的起始值（可选同时设置编号格式），不添加页脚页码域。

        :param doc_id: 会话 id。
        :param start: 起始页码。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param number_format: 页码编号格式，如 "upperRoman"/"大写罗马"/"chineseCounting"。

    AI 调用说明：
    适用：调整已有页码域的起始值。
    不要用于：需要新增页脚页码域；用 word_set_page_numbers/restart。
    调用前后：先确认该节已有页码设置。
    失败处理：节越界或格式非法会失败；修正参数后重试。
    """
    with _op(doc_id) as session:
        WordFormatter.set_page_number_start(
            _section(session, section_index), start, number_format=number_format,
        )
        return _result(session, section_index=section_index, start=start)


def word_set_page_number_format(
    doc_id: str,
    number_format: str,
    section_index: int = -1,
    start: Optional[int] = None,
) -> Dict[str, Any]:
    """只设置某节的页码编号格式（不改动已有的页码域）。

        封面用罗马数字、正文用阿拉伯数字时，配合 word_insert_section 分节后分别设置。

        :param doc_id: 会话 id。
        :param number_format: 编号格式，支持 decimal / upperRoman / lowerRoman /
                              chineseCounting / decimalEnclosedCircle 等，也接受
                              大写罗马 / 小写罗马 / 中文数字 / 带圈数字。
        :param section_index: 节索引，默认 -1 表示最后一节。
        :param start: 同时指定起始页码，可选。

    AI 调用说明：
    适用：只改变页码格式，不新增页码域。
    不要用于：首次创建页码；用 word_set_page_numbers。
    调用前后：先确认目标节已有页码，再改格式。
    失败处理：未知 number_format 会失败；使用受支持格式后重试。
    """
    with _op(doc_id) as session:
        WordFormatter.set_page_number_format(
            _section(session, section_index), number_format, start=start,
        )
        return _result(session, section_index=section_index,
                       number_format=number_format, start=start)


# ====================================================================
# Layer 2 · 目录与样式
# ====================================================================


def word_add_toc(
    doc_id: str,
    title: Optional[str] = "目录",
    levels_min: int = 1,
    levels_max: int = 3,
    title_font_name: str = "黑体",
    title_font_size: FontSize = 16,
    title_bold: bool = True,
    title_alignment: str = "居中",
    toc_level_styles: Optional[Dict[str, Any]] = None,
    use_hyperlinks: bool = True,
) -> Dict[str, Any]:
    """插入 Word 目录域（TOC）。

        注意：目录内容需要 Word/OnlyOffice 打开后刷新域才会显示页码。本工具会写入
        updateFields 标记，让 Word 打开时自动刷新；LibreOffice/Google Docs 不保证。

        :param doc_id: 会话 id。
        :param title: 目录标题；传 None 则不生成标题。
        :param levels_min: 收集的最小标题层级。
        :param levels_max: 收集的最大标题层级。
        :param title_font_name: 目录标题字体。
        :param title_font_size: 目录标题字号。
        :param title_bold: 目录标题是否加粗。
        :param title_alignment: 目录标题对齐方式。
        :param toc_level_styles: 各级目录样式，格式 {"1": {"font_name": "黑体", "font_size": 14}}。
        :param use_hyperlinks: 是否生成超链接。

    AI 调用说明：
    适用：生成或追加 Word TOC 目录域。
    不要用于：已有模板中的静态目录替换。
    调用前后：标题层级确定后调用；最终需在 Word/OnlyOffice 刷新域。
    失败处理：参数非法会失败；目录页码不能由工具即时校验，需提醒用户刷新。
    """
    with _op(doc_id) as session:
        styles = None
        if toc_level_styles:
            styles = {int(level): options for level, options in toc_level_styles.items()}
        session.formatter.add_toc(
            title=title, levels=(levels_min, levels_max),
            title_font_name=title_font_name, title_font_size=title_font_size,
            title_bold=title_bold, title_alignment=title_alignment,
            toc_level_styles=styles, use_hyperlinks=use_hyperlinks,
        )
        return _result(session, levels=[levels_min, levels_max])


def word_insert_toc(
    doc_id: str,
    anchor_paragraph_index: int,
    entries: List[Dict[str, Any]],
    levels_min: int = 1,
    levels_max: int = 3,
    use_hyperlinks: bool = True,
) -> Dict[str, Any]:
    """在模板指定位置插入 Word TOC 目录域，并设置识别出的大纲级别。

    AI 先根据文档标题、章节和上下文识别目录项，再把每个目录项对应的
    ``paragraph_index`` 与 ``level`` 传入。方法会把这些段落标记为 outline
    level，并用 TOC 域替换 anchor 段落，适合“目录”页已有提示文字的模板。

    :param doc_id: 会话 id。
    :param anchor_paragraph_index: 被目录替换的锚点段落下标，例如“请根据文件内容，自行加入目录”。
    :param entries: 目录项数组，元素形如 {"paragraph_index": 25, "level": 1}。
    :param levels_min: 目录收集的最小层级。
    :param levels_max: 目录收集的最大层级。
    :param use_hyperlinks: 是否生成目录超链接。

    AI 调用说明：
    适用：已有目录占位段，需要 AI 通读全文后识别章节并生成目录。
    不要用于：从空白文档末尾追加目录；这种情况用 word_add_toc。
    调用前后：先读取并填写正文，再通读章节标题，最后传入目录项；目录页码需刷新域。
    失败处理：锚点或目录项下标越界、层级非法会失败；重新读取 blocks 后修正 entries。
    """
    with _op(doc_id) as session:
        anchor = _paragraph(session, anchor_paragraph_index)
        toc_par, normalized = session.formatter.insert_toc(
            anchor,
            entries,
            levels=(levels_min, levels_max),
            use_hyperlinks=use_hyperlinks,
        )
        return _result(
            session,
            anchor_paragraph_index=anchor_paragraph_index,
            entries=normalized,
            toc_paragraph_index=anchor_paragraph_index,
            field_instruction=toc_par._p.xml,
        )


def word_finalize_document(
    doc_id: str,
    anchor_paragraph_index: int,
    entries: List[Dict[str, Any]],
    format_plan: Dict[str, Any],
    page_start_section_index: int = 0,
    start_page_number: int = 1,
    number_format: str = "decimal",
    prefix: str = "",
    suffix: str = "",
) -> Dict[str, Any]:
    """统一整理模板目录、标题层级格式和分节页码。

    AI 应先读取 `word_describe_format_profile` 和 `word_list_blocks`，根据原文档样式
    生成 format_plan，再调用本工具。整理完成后仍需在 OnlyOffice 中刷新 TOC/PAGE 域。

    :param doc_id: 会话 id。
    :param anchor_paragraph_index: 目录占位段落下标。
    :param entries: AI 识别出的目录项，元素含 paragraph_index / level。
    :param format_plan: 层级格式计划，元素形如 {"levels":{"1":{...},"2":{...}},"toc":{...}}。
    :param page_start_section_index: 从哪一节开始显示页码。
    :param start_page_number: 起始页码。
    :param number_format: 页码格式，默认 decimal。
    :param prefix: 页码前缀。
    :param suffix: 页码后缀。

    AI 调用说明：
    适用：正文填写完成后，统一目录、标题层级样式和分节页码。
    不要用于：正文尚未完成或不需要目录页码的普通文档。
    调用前后：先 describe_format_profile/list_blocks，再生成 entries 和 format_plan；完成后刷新域并保存。
    失败处理：锚点、目录项或节下标越界会失败；重新读取文档结构和格式后修正计划。
    """
    with _op(doc_id) as session:
        anchor = _paragraph(session, anchor_paragraph_index)
        result = session.formatter.finalize_document(
            anchor,
            entries,
            format_plan,
            page_start_section_index=page_start_section_index,
            start_page_number=start_page_number,
            number_format=number_format,
            prefix=prefix,
            suffix=suffix,
        )
        return _result(
            session,
            anchor_paragraph_index=anchor_paragraph_index,
            toc_paragraph_index=anchor_paragraph_index,
            entries=result["entries"],
            format_plan=result["format_plan"],
            page_start_section_index=result["page_start_section_index"],
            start_page_number=result["start_page_number"],
        )


def word_set_toc_level_style(
    doc_id: str,
    level: int,
    font_name: Optional[str] = None,
    font_size: Optional[FontSize] = None,
    bold: Optional[bool] = None,
    line_spacing: Optional[float] = None,
    space_before: Optional[float] = None,
    space_after: Optional[float] = None,
    first_line_indent_chars: Optional[float] = None,
    left_indent: Optional[float] = None,
    alignment: Optional[str] = None,
) -> Dict[str, Any]:
    """设置某一级 TOC 目录样式。

        :param doc_id: 会话 id。
        :param level: 目录层级。
        :param font_name: 字体名，可选。
        :param font_size: 字号，可选。
        :param bold: 是否加粗，可选。
        :param line_spacing: 行距倍数，可选。
        :param space_before: 段前空白（pt），可选。
        :param space_after: 段后空白（pt），可选。
        :param first_line_indent_chars: 首行缩进字符数，可选。
        :param left_indent: 左缩进（pt），可选。
        :param alignment: 对齐方式，可选。

    AI 调用说明：
    适用：调整指定 TOC 级别样式。
    不要用于：改变正文标题本身样式；用 paragraph style/add_heading。
    调用前后：先添加目录或目录样式，再设置级别格式。
    失败处理：级别越界会失败；修正 level 后重试。
    """
    with _op(doc_id) as session:
        session.formatter.set_toc_level_style(
            level, font_name=font_name, font_size=font_size, bold=bold,
            line_spacing=line_spacing, space_before=space_before,
            space_after=space_after,
            first_line_indent_chars=first_line_indent_chars,
            left_indent=left_indent, alignment=alignment,
        )
        return _result(session, level=level)


def word_set_paragraph_style(
    doc_id: str,
    style_name: str,
    base_style_name: str = "Normal",
    font_name: Optional[str] = None,
    font_size: Optional[FontSize] = None,
    bold: Optional[bool] = None,
    line_spacing: Optional[float] = None,
    space_before: Optional[float] = None,
    space_after: Optional[float] = None,
    first_line_indent_chars: Optional[float] = None,
    left_indent: Optional[float] = None,
    alignment: Optional[str] = None,
) -> Dict[str, Any]:
    """创建或更新一个自定义段落样式（配合 word_add_custom_heading 使用）。

        :param doc_id: 会话 id。
        :param style_name: 样式名称。
        :param base_style_name: 基础样式名，默认 Normal。
        其余参数均为可选，含义同 word_set_toc_level_style。

    AI 调用说明：
    适用：创建或更新可复用段落样式。
    不要用于：一次性的正文内容写入。
    调用前后：先定义样式，再用 word_add_custom_heading 或 word_call 应用。
    失败处理：基础样式不存在或参数非法会失败；修正 base_style_name 后重试。
    """
    with _op(doc_id) as session:
        session.formatter.set_paragraph_style(
            style_name, base_style_name=base_style_name, font_name=font_name,
            font_size=font_size, bold=bold, line_spacing=line_spacing,
            space_before=space_before, space_after=space_after,
            first_line_indent_chars=first_line_indent_chars,
            left_indent=left_indent, alignment=alignment,
        )
        return _result(session, style_name=style_name)


def word_add_custom_heading(
    doc_id: str,
    text: str,
    style_name: str,
    level: Optional[int] = None,
    font_name: Optional[str] = None,
    font_size: Optional[FontSize] = None,
) -> Dict[str, Any]:
    """使用自定义样式添加标题。

        :param doc_id: 会话 id。
        :param text: 标题文本。
        :param style_name: 之前用 word_set_paragraph_style 创建（或文档自带）的样式名。
        :param level: 大纲级别，传入后该标题可被目录收集。
        :param font_name: 覆盖字体名，可选。
        :param font_size: 覆盖字号，可选。

    AI 调用说明：
    适用：使用预定义自定义样式添加标题。
    不要用于：样式尚未创建时；先调用 word_set_paragraph_style。
    调用前后：先创建样式，再按文档顺序添加标题。
    失败处理：未知 style_name 会失败；先确认样式存在。
    """
    with _op(doc_id) as session:
        session.formatter.add_custom_heading(
            text, style_name=style_name, level=level,
            font_name=font_name, font_size=font_size,
        )
        return _result(session, style_name=style_name)


# ====================================================================
# Layer 2 · 模板化组合
# ====================================================================


def word_add_cover_page(
    doc_id: str,
    project_name: str = "默认工程名称",
    entrusting_unit: str = "默认委托单位",
    compiling_unit: str = "默认编制单位",
    report_title: str = "结算审核报告",
    date: Optional[str] = None,
    logo_width_cm: float = 4,
    info_blank_lines: int = 11,
    image_path: Optional[str] = None,
    image_base64: Optional[str] = None,
) -> Dict[str, Any]:
    """插入结算审核报告封面（Logo + 工程名 + 报告标题 + 落款表格 + 中文年月）。

        :param doc_id: 会话 id。
        :param project_name: 工程名称。
        :param entrusting_unit: 委托单位名称。
        :param compiling_unit: 编制单位名称。
        :param report_title: 报告标题。
        :param date: 封面日期，省略时用当天。
        :param logo_width_cm: Logo 宽度，单位 cm。
        :param info_blank_lines: 报告标题与落款表格之间的空行数。
        :param image_path: Logo 路径；省略时回退到环境变量 WC_REPORT_MCP_DEFAULT_LOGO。
        :param image_base64: Logo 的 base64（与 image_path 二选一）。

    AI 调用说明：
    适用：生成结算审核报告的标准封面组合。
    不要用于：填写已有封面模板；这种情况用 word_replace_text/fill_underline。
    调用前后：通常作为空白报告的第一个组合工具，之后添加目录和正文。
    失败处理：Logo 或参数非法会失败；修正图片和字段后重试。
    """
    with _op(doc_id) as session:
        logo = _resolve_logo(session, image_path, image_base64)
        session.formatter.fengmian_jiesuan(
            logo, project_name, entrusting_unit, compiling_unit=compiling_unit,
            report_title=report_title, date=date, logo_width=logo_width_cm,
            info_blank_lines=info_blank_lines,
        )
        return _result(session)


def word_add_signature_page(
    doc_id: str,
    project_name: str = "默认工程名称",
    entrusting_unit: str = "默认委托单位",
    personnel: Optional[Dict[str, Any]] = None,
    participants: Optional[List[Dict[str, Any]]] = None,
    compiling_unit: str = "默认编制单位",
    report_title: str = "结算审核报告",
    footer_text: str = "公司从业方针：客观、公正、严谨、专业",
    logo_width_cm: float = 2,
    logo_y_offset_pt: float = 0,
    image_path: Optional[str] = None,
    image_base64: Optional[str] = None,
) -> Dict[str, Any]:
    """新建一节并插入结算审核报告签发页。

        :param doc_id: 会话 id。
        :param project_name: 工程名称。
        :param entrusting_unit: 委托单位名称。
        :param personnel: 各签发岗位信息，键为 公司签发/部门核准/部门审核/小组初审/项目负责人，
                          值可含 姓名、部门、职务、职称，项目负责人还可含 联系电话。
        :param participants: 项目参与者列表，每项可含 姓名、职称。
        :param compiling_unit: 报告编制单位。
        :param report_title: 页眉中的报告标题。
        :param footer_text: 页脚文字。
        :param logo_width_cm: 页眉 Logo 宽度，单位 cm。
        :param logo_y_offset_pt: 页眉 Logo 垂直偏移，单位 pt，正数上移、负数下移。
        :param image_path: 页眉 Logo 路径；省略时回退到环境变量 WC_REPORT_MCP_DEFAULT_LOGO。
        :param image_base64: 页眉 Logo 的 base64。

    AI 调用说明：
    适用：生成结算审核报告标准签发页。
    不要用于：填写已有签发页模板。
    调用前后：通常放在正文之后，新节插入，之后可继续设置页码。
    失败处理：人员结构、Logo 或节参数非法会失败；修正后重试。
    """
    with _op(doc_id) as session:
        logo = _resolve_logo(session, image_path, image_base64)
        session.formatter.qianfaye(
            logo, project_name, entrusting_unit, personnel=personnel,
            participants=participants, compiling_unit=compiling_unit,
            report_title=report_title, footer_text=footer_text,
            logo_width=logo_width_cm, logo_y_offset_pt=logo_y_offset_pt,
        )
        return _result(session, section_index=len(session.doc.sections) - 1)


# ====================================================================
# Layer 2 · 无状态工具
# ====================================================================


def word_rmb_upper(value: Union[float, int, str]) -> Dict[str, Any]:
    """把数字转换为人民币金额中文大写（不需要 doc_id）。

        :param value: 数字或数字字符串，支持千分位逗号，例如 123456.78 或 "1,234.5"。

    AI 调用说明：
    适用：把金额数字转成人民币中文大写。
    不要用于：把普通整数转成文档页码或日期。
    调用前后：无状态，不需要 doc_id；结果可传给正文/表格工具。
    失败处理：无法解析的数字会失败；先规范为数值或数值字符串。
    """
    try:
        chinese_upper = number_to_chinese_upper(value)
    except CONVERTIBLE_ERRORS as exc:
        raise to_tool_error(exc) from exc
    return {"value": str(value), "chinese_upper": chinese_upper}


def word_convert_font_size(size: FontSize, ties: str = "larger") -> Dict[str, Any]:
    """在 pt 与中文字号之间互相换算（不需要 doc_id）。

        传 `11` 得到**吸附**过去的中文字号（偏向五号），传 `"二号"` 得到它对应的 pt。
        返回的 `exact=False` 表示该 pt 并不等于任何中文字号，是被吸附的，
        `delta_pt` 是吸附前后的差值。

        :param size: 字号，pt 数字 / 数字字符串 / 中文字号（"二号"、"小四"…）。
        :param ties: 正好落在两档正中间时的取向，`"larger"`（默认，取更大的号）
                     或 `"smaller"`（取更小的号），例如 13pt 在 小四(12) 与 四号(14) 正中。

    AI 调用说明：
    适用：在 pt 与中文字号之间换算。
    不要用于：直接修改文档字体；设置字体用 add/format/body 工具。
    调用前后：无状态，不需要 doc_id；换算结果再传给字号参数。
    失败处理：未知字号会失败；使用受支持 pt 或中文字号。
    """
    try:
        return describe_font_size(size, ties=ties)
    except CONVERTIBLE_ERRORS as exc:
        raise to_tool_error(exc) from exc


# ====================================================================
# Layer 3 · 兜底工具（保证对 WordFormatter 的全量覆盖）
# ====================================================================

_INDEXABLE_PARAMS = {"section": "sections", "table": "tables"}


def _resolve_indexed_item(items, value: int, *, label: str):
    index = value + len(items) if value < 0 else value
    if index < 0 or index >= len(items):
        upper = len(items) - 1
        raise ReportSessionError(
            f"{label}索引 {value} 超出范围（共 {len(items)} 个，0~{upper}）。"
        )
    return items[index]


def _resolve_cell_ref(session: Session, ref: Dict[str, Any]):
    missing = {"table_index", "row", "col"} - set(ref)
    if missing:
        raise ReportSessionError(
            f"cell 引用缺少字段：{', '.join(sorted(missing))}。"
            "格式应为 {table_index, row, col}。"
        )
    table = _table(session, int(ref["table_index"]))
    row, col = int(ref["row"]), int(ref["col"])
    row_count, col_count = len(table.rows), len(table.columns)
    if row < 0 or row >= row_count or col < 0 or col >= col_count:
        raise ReportSessionError(
            f"单元格 ({row}, {col}) 超出范围（该表共 {row_count} 行 × {col_count} 列）。"
        )
    return table.cell(row, col)


def _resolve_paragraph_ref(session: Session, ref: Union[int, Dict[str, Any]]):
    if isinstance(ref, int) and not isinstance(ref, bool):
        return _paragraph(session, ref)
    if not isinstance(ref, dict):
        raise ReportSessionError(
            "paragraph 引用必须是正文段落下标，或包含 "
            "table_index/row/col/paragraph_index 的对象。"
        )
    if "index" in ref and "table_index" not in ref:
        return _paragraph(session, int(ref["index"]))
    if "table_index" not in ref:
        raise ReportSessionError(
            "单元格 paragraph 引用必须包含 table_index/row/col/paragraph_index。"
        )
    cell = _resolve_cell_ref(session, ref)
    paragraph_index = int(ref.get("paragraph_index", 0))
    paragraphs = cell.paragraphs
    if not paragraphs:
        raise ReportSessionError("目标单元格中没有段落。")
    return _resolve_indexed_item(
        paragraphs, paragraph_index, label="单元格段落 ",
    )


def _resolve_run_ref(session: Session, ref: Dict[str, Any]):
    if not isinstance(ref, dict):
        raise ReportSessionError(
            "run 引用必须是对象，格式为 "
            "{paragraph:<段落下标或单元格段落引用>, run_index:<run 下标>}。"
        )
    if "paragraph" not in ref:
        raise ReportSessionError("run 引用缺少 paragraph 字段。")
    paragraph = _resolve_paragraph_ref(session, ref["paragraph"])
    runs = paragraph.runs
    if not runs:
        raise ReportSessionError("目标段落中没有 run。")
    run_index = int(ref.get("run_index", 0))
    return _resolve_indexed_item(runs, run_index, label="run ")


def _resolve_object_params(session: Session, params: Dict[str, Any]) -> Dict[str, Any]:
    """把 JSON 中的 section/table/paragraph/cell/run 引用换成真实对象。"""
    resolved = dict(params)
    for key, collection in _INDEXABLE_PARAMS.items():
        value = resolved.get(key)
        if isinstance(value, bool) or not isinstance(value, int):
            continue
        items = getattr(session.doc, collection)
        if not items:
            raise ReportSessionError(f"文档中还没有{collection}，无法解析 {key}={value}。")
        resolved[key] = _resolve_indexed_item(items, value, label=f"{key} ")

    paragraph_ref = resolved.get("paragraph")
    if isinstance(paragraph_ref, (int, dict)) and not isinstance(paragraph_ref, bool):
        resolved["paragraph"] = _resolve_paragraph_ref(session, paragraph_ref)
    cell_ref = resolved.get("cell")
    if isinstance(cell_ref, dict):
        resolved["cell"] = _resolve_cell_ref(session, cell_ref)
    run_ref = resolved.get("run")
    if isinstance(run_ref, dict):
        resolved["run"] = _resolve_run_ref(session, run_ref)
    return resolved


def _describe_result(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return type(value).__name__


def word_describe_api() -> str:
    """列出 WordFormatter 的全部公开方法、签名与说明（配合 word_call 使用）。

    AI 调用说明：
    适用：独立工具未覆盖某个 WordFormatter 方法时，查询完整方法签名和说明。
    不要用于：普通高层工具已能满足的场景。
    调用前后：先 describe_api 找到方法，再用 word_call 执行。
    失败处理：无状态工具；若方法说明仍不足，改用高层专用工具避免猜测参数。
    """
    return formatter_api_text()


def word_call(
    doc_id: str,
    method: str,
    params: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """直接调用 WordFormatter 上的任意公开方法（兜底能力，覆盖所有长尾参数）。

        先用 word_describe_api 查看可用方法与签名，再把关键字参数放进 params。
        section / table 可直接传整数索引；paragraph 传正文下标或单元格段落引用；
        cell 传 ``{table_index,row,col}``；
        run 传 ``{paragraph:<paragraph-ref>,run_index}``。负数索引均从末尾计算。

        :param doc_id: 会话 id。
        :param method: WordFormatter 的方法名，例如 "heading"、"set_toc_level_style"。
        :param params: 该方法的关字参数字典。

    AI 调用说明：
    适用：专用 MCP 工具未覆盖的长尾 WordFormatter 方法。
    不要用于：已有专用工具能完成的常见操作；优先用专用工具获得校验和错误信息。
    调用前后：先 word_describe_api 查签名，再按 params 调用；调用后检查 dirty 和返回结果。
    失败处理：方法名、参数或对象引用错误会失败；根据异常修正，不要重复盲调。
    """
    with _op(doc_id) as session:
        if method.startswith("_"):
            raise ReportSessionError(f"不允许调用私有方法: {method!r}。")
        target = getattr(session.formatter, method, None)
        if target is None or not callable(target):
            raise ReportSessionError(
                f"WordFormatter 没有公开方法 {method!r}。请先调用 word_describe_api 查看可用方法。"
            )
        resolved = _resolve_object_params(session, params or {})
        result = target(**resolved)
        return _result(session, method=method, returned=_describe_result(result))


TOOLS = (
    # Layer 1
    word_create_report,
    word_open_report,
    word_save_report,
    word_close_report,
    word_describe_report,
    # Layer 2 · 文档级设置
    word_set_default_font,
    word_set_page_margins,
    word_set_document_language,
    # Layer 2 · 段落与标题
    word_add_heading,
    word_add_body,
    word_add_body_list,
    word_add_body_segments,
    word_add_blank_lines,
    word_add_page_break,
    word_add_cover_text,
    word_add_right_text,
    word_add_date,
    word_insert_image,
    # Layer 2 · 表格
    word_add_table,
    word_merge_cells,
    word_format_cell,
    word_set_table_borders,
    # Layer 2 · 模板填充
    word_list_blocks,
    word_describe_format_profile,
    word_replace_text,
    word_fill_underline,
    word_replace_underline_text,
    word_insert_cell_image,
    # Layer 2 · 页眉页脚
    word_set_header,
    word_set_footer,
    word_set_header_parts,
    word_set_footer_parts,
    word_set_different_first_page,
    word_set_different_odd_even,
    word_clear_header_footer,
    word_set_header_image,
    # Layer 2 · 节与页码
    word_insert_section,
    word_set_section_page,
    word_set_page_numbers,
    word_restart_page_numbering,
    word_set_page_number_start,
    word_set_page_number_format,
    # Layer 2 · 目录与样式
    word_add_toc,
    word_insert_toc,
    word_finalize_document,
    word_set_toc_level_style,
    word_set_paragraph_style,
    word_add_custom_heading,
    # Layer 2 · 模板化组合
    word_add_cover_page,
    word_add_signature_page,
    # Layer 2 · 无状态
    word_rmb_upper,
    word_convert_font_size,
    # Layer 3 · 兜底
    word_describe_api,
    word_call,
)


def register(mcp) -> None:
    """把所有工具注册到 MCPServer 实例。"""
    for tool in TOOLS:
        mcp.tool()(tool)
