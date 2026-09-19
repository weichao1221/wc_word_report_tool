"""测试 wc_word_report_tool 的新 API 与向后兼容性。"""
import base64
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

from wc_word_report_tool import WordFormatter
from wc_word_report_tool.word import FONT_SIZE_MAP, _to_pt


def _read_zip_xml(docx_path: Path, inner_path: str) -> str:
    with ZipFile(docx_path) as zf:
        return zf.read(inner_path).decode("utf-8", errors="ignore")


def _write_test_png(path: Path) -> Path:
    path.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
        "/x8AAusB9Y9Z4k8AAAAASUVORK5CYII="
    ))
    return path


# ====================================================================
# 字号工具
# ====================================================================

def test_to_pt_supports_int_and_float():
    assert _to_pt(14) == 14.0
    assert _to_pt(14.5) == 14.5


def test_to_pt_supports_chinese_font_size():
    assert _to_pt("三号") == 16
    assert _to_pt("小五") == 9
    assert _to_pt("初号") == 42


def test_to_pt_supports_numeric_string():
    assert _to_pt("16") == 16.0


def test_to_pt_raises_on_unknown_string():
    with pytest.raises(ValueError):
        _to_pt("unknown")


def test_font_size_map_contains_common_sizes():
    assert FONT_SIZE_MAP["三号"] == 16
    assert FONT_SIZE_MAP["小五"] == 9
    assert FONT_SIZE_MAP["四号"] == 14


# ====================================================================
# 文档级设置
# ====================================================================

def test_set_default_font_sets_normal_style(tmp_path):
    doc = Document()
    WordFormatter(doc).set_default_font(font_size="三号",
                                        cn_font="仿宋_GB2312", en_font="Times New Roman")
    out = tmp_path / "default_font.docx"
    doc.save(out)

    styles_xml = _read_zip_xml(out, "word/styles.xml")
    assert 'w:eastAsia="仿宋_GB2312"' in styles_xml
    assert 'w:ascii="Times New Roman"' in styles_xml
    assert 'w:sz w:val="32"' in styles_xml  # 16pt = 32 half-points


def test_set_document_language_sets_zh_cn(tmp_path):
    doc = Document()
    WordFormatter(doc).set_document_language("zh-CN")
    out = tmp_path / "lang.docx"
    doc.save(out)

    settings_xml = _read_zip_xml(out, "word/settings.xml")
    assert 'w:val="zh-CN"' in settings_xml
    assert 'w:eastAsia="zh-CN"' in settings_xml


def test_set_page_margins(tmp_path):
    doc = Document()
    WordFormatter(doc).set_page_margins(top=2.5, left=2.6)
    out = tmp_path / "margins.docx"
    doc.save(out)
    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:top="1417"' in document_xml  # 2.5cm
    assert 'w:left="1474"' in document_xml  # 2.6cm


def test_setup_defaults_only_sets_page_margins():
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.setup_defaults(top=3.7, bottom=3.5, left=2.8, right=2.6)

    section = doc.sections[0]
    assert round(section.top_margin.cm, 1) == 3.7
    assert round(section.bottom_margin.cm, 1) == 3.5
    assert round(section.left_margin.cm, 1) == 2.8
    assert round(section.right_margin.cm, 1) == 2.6

    with pytest.raises(TypeError):
        formatter.setup_defaults(body_font_name="宋体")


# ====================================================================
# 段落与标题（新 API）
# ====================================================================

def test_body_creates_paragraph_with_indent(tmp_path):
    doc = Document()
    par = WordFormatter(doc).body("正文内容")
    assert par.alignment == WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    assert par.paragraph_format.first_line_indent is not None
    assert par.paragraph_format.line_spacing == 1.5
    out = tmp_path / "body.docx"
    doc.save(out)
    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:eastAsia="宋体"' in document_xml
    assert 'w:firstLineChars="200"' in document_xml
    assert 'w:firstLine="560"' in document_xml


def test_body_with_no_indent(tmp_path):
    doc = Document()
    par = WordFormatter(doc).body("无缩进", indent=False)
    assert par.paragraph_format.first_line_indent is None
    out = tmp_path / "body-no-indent.docx"
    doc.save(out)
    document_xml = _read_zip_xml(out, "word/document.xml")
    assert "w:firstLineChars" not in document_xml
    assert "w:firstLine=" not in document_xml


def test_set_paragraph_format_uses_chars_with_pt_fallback(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    par = doc.add_paragraph("自定义字符缩进")
    formatter.set_paragraph_format(
        par,
        first_line_indent_chars=1.5,
        first_line_indent_pt=21,
    )

    ind = par._p.pPr.ind
    assert ind.get(qn("w:firstLineChars")) == "150"
    assert ind.get(qn("w:firstLine")) == "420"

    out = tmp_path / "paragraph-indent.docx"
    doc.save(out)
    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:firstLineChars="150"' in document_xml
    assert 'w:firstLine="420"' in document_xml


def test_pt_only_indent_removes_existing_character_priority():
    doc = Document()
    par = doc.add_paragraph("切换缩进单位")
    WordFormatter.set_paragraph_format(
        par,
        first_line_indent_chars=2,
        first_line_indent_pt=28,
    )
    WordFormatter.set_paragraph_format(par, first_line_indent_pt=14)

    ind = par._p.pPr.ind
    assert ind.get(qn("w:firstLineChars")) is None
    assert ind.get(qn("w:firstLine")) == "280"


def test_indent_enabled_related_apis_write_chars_and_fallback(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    cover = formatter.cover_text("封面", font_size=16, indent=True)
    right = formatter.right_text("日期", font_size=16, indent=True)
    formatter.heading("标题", font_size=16, indent=True)

    for par in (cover, right):
        ind = par._p.pPr.ind
        assert ind.get(qn("w:firstLineChars")) == "200"
        assert ind.get(qn("w:firstLine")) == "640"

    out = tmp_path / "related-indent-apis.docx"
    doc.save(out)
    styles_xml = _read_zip_xml(out, "word/styles.xml")
    assert 'w:firstLineChars="200"' in styles_xml
    assert 'w:firstLine="640"' in styles_xml


def test_paragraph_style_supports_chars_with_pt_fallback(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.set_paragraph_style(
        "字符缩进样式",
        first_line_indent_chars=2,
        first_line_indent=28,
    )

    out = tmp_path / "style-indent.docx"
    doc.save(out)
    styles_xml = _read_zip_xml(out, "word/styles.xml")
    assert 'w:firstLineChars="200"' in styles_xml
    assert 'w:firstLine="560"' in styles_xml


def test_toc_style_supports_chars_with_pt_fallback(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.set_toc_level_style(
        1,
        first_line_indent_chars=2,
        first_line_indent=28,
    )

    out = tmp_path / "toc-style-indent.docx"
    doc.save(out)
    styles_xml = _read_zip_xml(out, "word/styles.xml")
    assert 'w:firstLineChars="200"' in styles_xml
    assert 'w:firstLine="560"' in styles_xml


@pytest.mark.parametrize("value, error_type", [("two", TypeError), (-1, ValueError)])
def test_first_line_indent_chars_rejects_invalid_values(value, error_type):
    doc = Document()
    par = doc.add_paragraph("无效缩进")
    with pytest.raises(error_type):
        WordFormatter.set_paragraph_format(par, first_line_indent_chars=value)


def test_heading_uses_builtin_heading_style(tmp_path):
    doc = Document()
    paragraph = WordFormatter(doc).heading("一级标题", level=1)
    out = tmp_path / "h1.docx"
    doc.save(out)

    assert paragraph.style.name == "Heading 1"
    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:pStyle w:val="Heading1"' in document_xml


def test_heading_supports_bold(tmp_path):
    doc = Document()
    par = WordFormatter(doc).heading("二级", level=2, bold=True)
    out = tmp_path / "h2.docx"
    doc.save(out)
    assert par.runs[0].font.bold is True


def test_heading_is_not_bold_by_default():
    doc = Document()
    par = WordFormatter(doc).heading("三级", level=3)
    assert par.runs[0].font.bold is False


def test_cover_text_is_centered_and_bold():
    doc = Document()
    par = WordFormatter(doc).cover_text("项目名", font_size=22, bold=True)
    assert par.alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
    assert par.runs[0].font.bold is True
    assert par.runs[0].font.size.pt == 22


def test_blank_lines_creates_n_paragraphs():
    doc = Document()
    initial = len(doc.paragraphs)
    WordFormatter(doc).blank_lines(3)
    assert len(doc.paragraphs) == initial + 3


def test_fengmian_jiesuan_creates_cover(tmp_path):
    doc = Document()
    logo = _write_test_png(tmp_path / "logo.png")

    table = WordFormatter(doc).fengmian_jiesuan(
        logo, "测试工程", "测试委托单位", date="2026-07-21",
        info_blank_lines=1,
    )

    assert len(table.rows) == 2
    assert "测试工程" in "".join(paragraph.text for paragraph in doc.paragraphs)
    assert table.cell(0, 2).text == "测试委托单位"
    assert table.cell(1, 2).text == "默认编制单位"


def test_qianfaye_creates_section_and_personnel_table(tmp_path):
    doc = Document()
    logo = _write_test_png(tmp_path / "logo.png")
    personnel = {
        "公司签发": {"姓名": "张三", "职务": "总经理", "职称": "正高级工程师"},
        "项目负责人": {
            "姓名": "李四", "部门": "造价部", "职务": "项目经理",
            "职称": "高级工程师", "联系电话": "13800000000",
        },
    }

    table = WordFormatter(doc).qianfaye(
        logo, "测试工程", "测试委托单位", personnel,
        participants=[{"姓名": "王五", "职称": "工程师"}],
    )

    assert len(doc.sections) == 2
    assert len(table.rows) == 11
    assert table.cell(3, 1).text == "张三"
    assert table.cell(9, 2).text == "联系电话：13800000000"
    assert table.cell(10, 0).text == "项目参与者："
    assert table.cell(10, 1).text == "王五"


def test_right_text_is_right_aligned():
    doc = Document()
    par = WordFormatter(doc).right_text("公司名")
    assert par.alignment == WD_PARAGRAPH_ALIGNMENT.RIGHT


def test_created_time_uses_today():
    import datetime
    doc = Document()
    par = WordFormatter(doc).created_time()
    today = datetime.datetime.now().strftime("%Y年%m月%d日")
    assert today in par.text


# ====================================================================
# 表格（新增能力）
# ====================================================================

def test_set_table_borders(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    WordFormatter.set_table_borders(table, color="000000", size=4)
    out = tmp_path / "borders.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert "tblBorders" in document_xml
    assert 'w:color="000000"' in document_xml


def test_set_cell_supports_bold_and_alignment(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    WordFormatter.set_cell(
        table.rows[0].cells[0], "表头",
        bold=True, alignment="居中",
    )
    out = tmp_path / "cell.docx"
    doc.save(out)

    cell = table.rows[0].cells[0]
    assert cell.paragraphs[0].alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
    assert cell.paragraphs[0].runs[0].font.bold is True


def test_add_table_creates_table_with_headers(tmp_path):
    doc = Document()
    headers = ["序号", "项目", "金额"]
    rows = [["1", "建安费", "1200"], ["2", "设备费", "850"]]
    table = WordFormatter(doc).add_table(headers, rows, col_widths=[2, 5, 3])

    out = tmp_path / "table.docx"
    doc.save(out)

    assert len(table.rows) == 3  # 1 header + 2 data
    assert len(table.columns) == 3

    # 表头应该加粗
    header_cell = table.rows[0].cells[0]
    assert header_cell.paragraphs[0].runs[0].font.bold is True

    # 数据不应加粗
    data_cell = table.rows[1].cells[0]
    assert data_cell.paragraphs[0].runs[0].font.bold is False

    # 边框已添加
    document_xml = _read_zip_xml(out, "word/document.xml")
    assert "tblBorders" in document_xml


def test_set_cell_clears_existing_runs():
    """重复调用 set_cell 不应叠加 runs。"""
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    WordFormatter.set_cell(cell, "第一次")
    WordFormatter.set_cell(cell, "第二次")
    assert len(cell.paragraphs[0].runs) == 1
    assert cell.paragraphs[0].text == "第二次"


# ====================================================================
# 页眉页脚（新增能力）
# ====================================================================

def test_set_header(tmp_path):
    doc = Document()
    doc.add_paragraph("正文")
    from docx.enum.section import WD_SECTION_START
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph("第二页")

    sec = doc.sections[1]
    WordFormatter.set_header(sec, "项目名称 报告名称", alignment="居中")
    out = tmp_path / "header.docx"
    doc.save(out)

    # header xml 文件名可能为 header1.xml / header2.xml 等，遍历查找
    with ZipFile(out) as zf:
        header_names = [n for n in zf.namelist() if n.startswith("word/header")]
        header_xml = ""
        for name in header_names:
            content = zf.read(name).decode("utf-8", errors="ignore")
            if "项目名称" in content:
                header_xml = content
                break
    assert "项目名称 报告名称" in header_xml
    assert "w:jc" in header_xml  # 居中


def test_set_footer_uses_kaiti_font(tmp_path):
    doc = Document()
    doc.add_paragraph("正文")
    from docx.enum.section import WD_SECTION_START
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph("第二页")

    sec = doc.sections[1]
    WordFormatter.set_footer(sec, "公司名称", alignment="居右")
    out = tmp_path / "footer.docx"
    doc.save(out)

    # footer xml 文件名可能为 footer1.xml / footer2.xml 等
    with ZipFile(out) as zf:
        footer_names = [n for n in zf.namelist() if n.startswith("word/footer")]
        footer_xml = ""
        for name in footer_names:
            content = zf.read(name).decode("utf-8", errors="ignore")
            if "公司名称" in content:
                footer_xml = content
                break
    assert "公司名称" in footer_xml
    assert 'w:eastAsia="楷体"' in footer_xml
    assert 'w:sz w:val="18"' in footer_xml  # 9pt = 18 half-points


def test_set_header_does_not_stack_runs():
    """重复调用不应叠加。"""
    doc = Document()
    sec = doc.sections[0]
    WordFormatter.set_header(sec, "第一次")
    WordFormatter.set_header(sec, "第二次")
    assert len(sec.header.paragraphs[0].runs) == 1


def test_clear_footer_removes_content():
    doc = Document()
    sec = doc.sections[0]
    WordFormatter.set_footer(sec, "原有页脚")
    WordFormatter.clear_footer(sec)
    assert all(not p.runs for p in sec.footer.paragraphs)


def test_clear_header_removes_content():
    doc = Document()
    sec = doc.sections[0]
    WordFormatter.set_header(sec, "原有页眉")
    WordFormatter.clear_header(sec)
    assert all(not p.runs for p in sec.header.paragraphs)


# ====================================================================
# 跨节页码（新增能力）
# ====================================================================

def test_set_page_number_from_section_clears_previous_sections(tmp_path):
    """前置节无页码，正文节重启页码，后续节链接。"""
    doc = Document()
    doc.add_paragraph("前置内容")

    from docx.enum.section import WD_SECTION_START
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph("正文第一页")

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph("正文第二页")

    WordFormatter(doc).set_page_number_from_section(
        start_section_idx=1, start=1, prefix="第 ", suffix=" 页",
    )
    out = tmp_path / "page_number_from_section.docx"
    doc.save(out)

    # 节 0 页脚应为空（无 runs）
    sec0_footer_runs = sum(
        len(p.runs) for p in doc.sections[0].footer.paragraphs
    )
    assert sec0_footer_runs == 0

    # 节 1 页脚应有 PAGE 域
    with ZipFile(out) as zf:
        footer_names = [n for n in zf.namelist() if n.startswith("word/footer")]
        footer_xml = ""
        for name in footer_names:
            content = zf.read(name).decode("utf-8", errors="ignore")
            if "PAGE" in content:
                footer_xml = content
                break
    assert "PAGE" in footer_xml
    assert "第 " in footer_xml

    # 节 2 应链接到前一节
    assert doc.sections[2].footer.is_linked_to_previous is True


def test_set_page_number_from_section_invalid_index_raises():
    doc = Document()
    with pytest.raises(ValueError, match="超出范围"):
        WordFormatter(doc).set_page_number_from_section(start_section_idx=99)


# ====================================================================
# 严格对齐校验
# ====================================================================

def test_resolve_alignment_strict_raises_on_unknown():
    with pytest.raises(ValueError, match="未知对齐方式"):
        WordFormatter.resolve_alignment("invalid_align", strict=True)


def test_resolve_alignment_non_strict_falls_back_to_left():
    result = WordFormatter.resolve_alignment("invalid_align", strict=False)
    assert result == WD_PARAGRAPH_ALIGNMENT.LEFT


def test_resolve_alignment_supports_chinese():
    assert WordFormatter.resolve_alignment("居中") == WD_PARAGRAPH_ALIGNMENT.CENTER
    assert WordFormatter.resolve_alignment("居左") == WD_PARAGRAPH_ALIGNMENT.LEFT
    assert WordFormatter.resolve_alignment("居右") == WD_PARAGRAPH_ALIGNMENT.RIGHT
    assert WordFormatter.resolve_alignment("左对齐") == WD_PARAGRAPH_ALIGNMENT.LEFT
    assert WordFormatter.resolve_alignment("右对齐") == WD_PARAGRAPH_ALIGNMENT.RIGHT
    assert WordFormatter.resolve_alignment("两端对齐") == WD_PARAGRAPH_ALIGNMENT.JUSTIFY


def test_resolve_alignment_supports_english():
    assert WordFormatter.resolve_alignment("center") == WD_PARAGRAPH_ALIGNMENT.CENTER
    assert WordFormatter.resolve_alignment("LEFT") == WD_PARAGRAPH_ALIGNMENT.LEFT


def test_resolve_alignment_supports_enum():
    enum_val = WD_PARAGRAPH_ALIGNMENT.CENTER
    assert WordFormatter.resolve_alignment(enum_val) == enum_val


# ====================================================================
# 页码与目录
# ====================================================================

def test_restart_page_numbering_unlinks_footer_and_sets_start(tmp_path):
    doc = Document()
    section = WordFormatter(doc).insert_section_with_page_numbering(
        start_page_number=1, prefix="第", suffix="页",
    )
    out = tmp_path / "page.docx"
    doc.save(out)

    assert section.footer.is_linked_to_previous is False
    assert section.header.is_linked_to_previous is False

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:pgNumType w:start="1"' in document_xml

    footer_xml = _read_zip_xml(out, "word/footer1.xml")
    assert "PAGE" in footer_xml


def test_add_toc_with_custom_style_mapping(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.set_paragraph_style("MyHeading1", font_name="黑体",
                                  font_size=14, bold=True)
    formatter.add_toc(
        levels=(1, 3), use_outline_levels=True,
        custom_style_levels={"MyHeading1": 1},
        toc_level_styles={1: {"font_name": "黑体", "font_size": 14, "bold": True}},
    )
    formatter.add_custom_heading("自定义标题", style_name="MyHeading1", level=1)

    out = tmp_path / "toc.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    settings_xml = _read_zip_xml(out, "word/settings.xml")
    styles_xml = _read_zip_xml(out, "word/styles.xml")

    assert 'TOC \\o "1-3" \\h \\z \\u \\t "MyHeading1,1"' in document_xml
    assert "updateFields" in settings_xml
    assert 'styleId="TOC1"' in styles_xml or 'styleId="TOC 1"' in styles_xml


def test_body_supports_custom_font_name_and_size(tmp_path):
    doc = Document()
    WordFormatter(doc).body("正文", font_name="黑体", font_size=18)
    out = tmp_path / "font.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="黑体"' in document_xml
    assert 'w:sz w:val="36"' in document_xml


def test_set_page_margins_supports_numeric_alignment(tmp_path):
    doc = Document()
    WordFormatter(doc).set_page_margins(left=2.5, right=2.6, horizontal_alignment=2)
    out = tmp_path / "layout.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:left="1417"' in document_xml
    assert 'w:right="1474"' in document_xml
    assert 'w:jc w:val="center"' in document_xml


# ====================================================================
# A1 · 页眉页脚：三段分列 / 首页不同 / 奇偶页不同
# ====================================================================

def test_set_header_parts_writes_tabs_and_tab_stops(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    sec = doc.sections[0]

    formatter.set_header_parts(sec, left="公司名", center="报告名", right="第 1 页")
    out = tmp_path / "header-parts.docx"
    doc.save(out)

    with ZipFile(out) as zf:
        header_xml = "".join(
            zf.read(n).decode("utf-8", errors="ignore")
            for n in zf.namelist() if n.startswith("word/header")
        )

    # 制表符必须落成 <w:tab/>，否则 Word 不会按制表位分列
    assert header_xml.count("<w:tab/>") == 2
    assert "公司名" in header_xml and "报告名" in header_xml and "第 1 页" in header_xml
    # 居中 + 居右两个制表位
    assert header_xml.count("<w:tab ") == 2
    assert 'w:val="center"' in header_xml
    assert 'w:val="right"' in header_xml


def test_set_footer_parts_writes_three_segments(tmp_path):
    doc = Document()
    WordFormatter.set_footer_parts(
        doc.sections[0], left="左", center="中", right="右",
    )
    out = tmp_path / "footer-parts.docx"
    doc.save(out)

    with ZipFile(out) as zf:
        footer_xml = "".join(
            zf.read(n).decode("utf-8", errors="ignore")
            for n in zf.namelist() if n.startswith("word/footer")
        )

    assert "左" in footer_xml and "中" in footer_xml and "右" in footer_xml
    assert footer_xml.count("<w:tab/>") == 2


def test_different_first_page_and_odd_even_flags(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    sec = doc.sections[0]

    formatter.set_different_first_page(sec, True)
    formatter.set_different_odd_even(True)
    out = tmp_path / "variants.docx"
    doc.save(out)

    assert "<w:titlePg/>" in _read_zip_xml(out, "word/document.xml")
    assert "evenAndOddHeaders" in _read_zip_xml(out, "word/settings.xml")


def test_different_odd_even_can_be_turned_off(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.set_different_odd_even(True)
    formatter.set_different_odd_even(False)

    out = tmp_path / "odd-even-off.docx"
    doc.save(out)
    assert "evenAndOddHeaders" not in _read_zip_xml(out, "word/settings.xml")


def test_header_variants_write_to_separate_parts(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    sec = doc.sections[0]
    formatter.set_different_first_page(sec, True)
    formatter.set_different_odd_even(True)

    formatter.set_header(sec, "默认页眉")
    formatter.set_header(sec, "首页页眉", variant="first")
    formatter.set_header(sec, "偶数页页眉", variant="even")
    out = tmp_path / "variants-parts.docx"
    doc.save(out)

    from docx import Document as _Read
    reopened = _Read(str(out))
    section = reopened.sections[0]

    assert section.header.paragraphs[0].text == "默认页眉"
    assert section.first_page_header.paragraphs[0].text == "首页页眉"
    assert section.even_page_header.paragraphs[0].text == "偶数页页眉"


def test_header_footer_part_rejects_unknown_variant():
    doc = Document()

    with pytest.raises(ValueError, match="未知的页眉页脚变体"):
        WordFormatter.header_footer_part(doc.sections[0], variant="sidebar")


# ====================================================================
# A2 · 页码编号格式
# ====================================================================

def test_restart_page_numbering_writes_number_format(tmp_path):
    doc = Document()
    WordFormatter.restart_page_numbering(
        doc.sections[0], start=3, number_format="upperRoman",
    )
    out = tmp_path / "roman.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:fmt="upperRoman"' in document_xml
    assert 'w:start="3"' in document_xml


def test_page_number_format_accepts_chinese_alias(tmp_path):
    doc = Document()
    WordFormatter.set_page_number_format(doc.sections[0], "大写罗马")
    out = tmp_path / "roman-alias.docx"
    doc.save(out)

    assert 'w:fmt="upperRoman"' in _read_zip_xml(out, "word/document.xml")


def test_set_page_number_format_keeps_existing_start(tmp_path):
    doc = Document()
    sec = doc.sections[0]
    WordFormatter.set_page_number_start(sec, 5)
    WordFormatter.set_page_number_format(sec, "lowerRoman")
    out = tmp_path / "fmt-only.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:fmt="lowerRoman"' in document_xml
    assert 'w:start="5"' in document_xml


def test_unknown_page_number_format_raises():
    with pytest.raises(ValueError, match="未知的页码格式"):
        WordFormatter.resolve_page_number_format("romanNumerals99")


# ====================================================================
# A3 · 分节纸张方向 / 尺寸 / 边距
# ====================================================================

def test_set_section_page_swaps_dimensions_for_landscape(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    sec = doc.sections[0]
    portrait = (sec.page_width.cm, sec.page_height.cm)

    formatter.set_section_page(sec, orientation="横向")
    out = tmp_path / "landscape.docx"
    doc.save(out)

    assert sec.page_width.cm == pytest.approx(portrait[1], abs=0.1)
    assert sec.page_height.cm == pytest.approx(portrait[0], abs=0.1)
    assert 'w:orient="landscape"' in _read_zip_xml(out, "word/document.xml")


def test_set_section_page_applies_only_given_margins():
    doc = Document()
    formatter = WordFormatter(doc)
    sec = doc.sections[0]
    original_bottom = sec.bottom_margin.cm

    formatter.set_section_page(sec, left=2.54, right=2.54)

    assert sec.left_margin.cm == pytest.approx(2.54, abs=0.01)
    assert sec.right_margin.cm == pytest.approx(2.54, abs=0.01)
    assert sec.bottom_margin.cm == pytest.approx(original_bottom, abs=0.01)


def test_insert_section_sets_per_section_page_setup():
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.body("第一节")

    sec = formatter.insert_section(
        inherit_header=False, inherit_footer=False,
        orientation="横向", top=3.17, bottom=3.17, left=2.54, right=2.54,
    )

    assert sec.page_width.cm > sec.page_height.cm
    assert sec.left_margin.cm == pytest.approx(2.54, abs=0.01)
    # 第一节仍保持纵向，说明是「按节」生效而不是全文
    assert doc.sections[0].page_width.cm < doc.sections[0].page_height.cm


def test_resolve_orientation_rejects_unknown_value():
    with pytest.raises(ValueError, match="未知的纸张方向"):
        WordFormatter.resolve_orientation("斜向")


# ====================================================================
# B1 · 标题对齐
# ====================================================================

def test_heading_supports_alignment(tmp_path):
    doc = Document()
    WordFormatter(doc).heading("居中大标题", level=1, font_size="二号", alignment="居中")
    out = tmp_path / "heading-align.docx"
    doc.save(out)

    assert 'w:jc w:val="center"' in _read_zip_xml(out, "word/document.xml")


def test_heading_alignment_is_per_paragraph_not_per_style():
    """同一级标题对齐方式不同时不能互相串味。"""
    doc = Document()
    formatter = WordFormatter(doc)
    centered = formatter.heading("居中标题", level=1, alignment="居中")
    left = formatter.heading("居左标题", level=1)

    assert centered.alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
    assert left.alignment is None


def test_heading_rejects_unknown_alignment():
    doc = Document()

    with pytest.raises(ValueError, match="未知对齐方式"):
        WordFormatter(doc).heading("标题", alignment="斜着放")


def test_heading_accepts_space_before_and_after(tmp_path):
    doc = Document()
    para = WordFormatter(doc).heading("标题", space_before=18, space_after=6)
    out = tmp_path / "heading-space.docx"
    doc.save(out)

    # w:before / w:after 以 1/20 pt 存储，18pt -> 360
    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:before="360"' in document_xml
    assert 'w:after="120"' in document_xml
    assert para.paragraph_format.space_before.pt == pytest.approx(18)
    assert para.paragraph_format.space_after.pt == pytest.approx(6)


def test_heading_spacing_is_per_paragraph_not_per_style():
    """段间距写进样式会让同级标题互相覆盖，必须落在段落上。"""
    doc = Document()
    formatter = WordFormatter(doc)
    first = formatter.heading("第一个", level=1, space_before=12)
    second = formatter.heading("第二个", level=1, space_before=30)

    assert first.paragraph_format.space_before.pt == pytest.approx(12)
    assert second.paragraph_format.space_before.pt == pytest.approx(30)


# ====================================================================
# B2 · 表格合并单元格
# ====================================================================

def test_add_table_merges_cells(tmp_path):
    doc = Document()
    table = WordFormatter(doc).add_table(
        headers=["序号", "项目", "金额", "备注"],
        rows=[["1", "合并行", "100", "a"], ["2", "x", "200", "b"]],
        merges=[{"row": 1, "col": 1, "rowspan": 1, "colspan": 2}],
    )
    out = tmp_path / "merged.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:gridSpan w:val="2"' in document_xml
    assert table.cell(1, 1).text == "合并行"


def test_add_table_merge_uses_top_left_value_and_drops_extra_paragraphs(tmp_path):
    doc = Document()
    table = WordFormatter(doc).add_table(
        headers=["A", "B", "C"],
        rows=[["保留", "被忽略", "被忽略2"]],
        merges=[(1, 0, 1, 3)],
    )

    assert table.cell(1, 0).text == "保留"
    assert len(table.cell(1, 0).paragraphs) == 1


def test_add_table_vertical_merge(tmp_path):
    doc = Document()
    table = WordFormatter(doc).add_table(
        headers=["A", "B"],
        rows=[["跨行", "1"], ["", "2"]],
        merges=[{"row": 1, "col": 0, "rowspan": 2, "colspan": 1}],
    )
    out = tmp_path / "vmerge.docx"
    doc.save(out)

    assert 'w:vMerge' in _read_zip_xml(out, "word/document.xml")
    assert table.cell(1, 0).text == "跨行"


def test_merge_region_out_of_range_raises():
    doc = Document()

    with pytest.raises(ValueError, match="超出表格范围"):
        WordFormatter(doc).add_table(
            headers=["A", "B"], rows=[["1", "2"]],
            merges=[{"row": 0, "col": 0, "colspan": 9}],
        )


def test_overlapping_merges_raise():
    doc = Document()

    with pytest.raises(ValueError, match="覆盖了同一个单元格"):
        WordFormatter(doc).add_table(
            headers=["A", "B", "C"], rows=[["1", "2", "3"]],
            merges=[{"row": 0, "col": 0, "colspan": 2},
                    {"row": 0, "col": 1, "colspan": 2}],
        )


def test_merge_cells_on_existing_table():
    doc = Document()
    formatter = WordFormatter(doc)
    table = formatter.add_table(headers=["A", "B", "C"], rows=[["1", "2", "3"]])

    WordFormatter.merge_cells(table, [(0, 0, 1, 3)])

    assert table.cell(0, 0).text == "A"


# ====================================================================
# B5 · 表格总宽 / 行高 / 逐格字体
# ====================================================================

def test_add_table_supports_total_width_and_row_height(tmp_path):
    doc = Document()
    WordFormatter(doc).add_table(
        headers=["A", "B"], rows=[["1", "2"], ["3", "4"]],
        table_width_cm=12, row_heights=[1.0, 1.5, 2.0],
    )
    out = tmp_path / "table-size.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:w="6804"' in document_xml          # 12cm -> 6804 twips
    assert 'w:hRule="atLeast"' in document_xml


def test_add_table_supports_per_cell_font_names(tmp_path):
    doc = Document()
    WordFormatter(doc).add_table(
        headers=["A", "B"], rows=[["1", "2"]],
        font_name="宋体",
        cell_font_names=[["黑体", "黑体"], ["楷体_GB2312", None]],
    )
    out = tmp_path / "table-fonts.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:eastAsia="黑体"' in document_xml
    assert 'w:eastAsia="楷体_GB2312"' in document_xml
    assert 'w:eastAsia="宋体"' in document_xml   # 未覆盖的位置回落到全局字体


def test_set_table_width_rejects_non_positive():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)

    with pytest.raises(ValueError, match="表格宽度必须大于 0"):
        WordFormatter.set_table_width(table, 0)


# ====================================================================
# B3 · 图片：自动尺寸与二进制输入
# ====================================================================

def test_insert_img_accepts_bytes_without_width(tmp_path):
    doc = Document()
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
        "/x8AAusB9Y9Z4k8AAAAASUVORK5CYII="
    )

    WordFormatter(doc).insert_img(png)
    out = tmp_path / "img-bytes.docx"
    doc.save(out)

    assert any(n.startswith("word/media/") for n in ZipFile(out).namelist())


def test_insert_img_derives_width_from_image_dpi(tmp_path):
    doc = Document()
    png = tmp_path / "pic.png"
    _write_test_png(png)

    WordFormatter(doc).insert_img(str(png))
    out = tmp_path / "img-auto.docx"
    doc.save(out)

    assert any(n.startswith("word/media/") for n in ZipFile(out).namelist())


def test_insert_img_rejects_non_positive_width(tmp_path):
    doc = Document()

    with pytest.raises(ValueError, match="图片宽度必须大于 0"):
        WordFormatter(doc).insert_img(str(_write_test_png(tmp_path / "p.png")), 0)


def test_insert_img_missing_file_raises(tmp_path):
    doc = Document()

    with pytest.raises(FileNotFoundError, match="图片不存在"):
        WordFormatter(doc).insert_img(str(tmp_path / "nope.png"), 2)


def test_insert_img_supports_floating(tmp_path):
    """印章需要浮在落款文字之上，不能被排版推来推去。"""
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.body("落款：某某公司")
    png = _write_test_png(tmp_path / "seal.png")

    formatter.insert_img(str(png), 3.0, alignment="居右", floating=True, y_offset_pt=6)
    out = tmp_path / "seal.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert "w:pict" in document_xml
    assert "v:shape" in document_xml


def test_insert_img_is_inline_by_default(tmp_path):
    doc = Document()
    WordFormatter(doc).insert_img(_write_test_png(tmp_path / "p.png"), 3.0)
    out = tmp_path / "inline.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert "w:pict" not in document_xml
    assert "wp:inline" in document_xml


# ====================================================================
# 表格参数自洽性校验：越界要给人看得懂的错误，而不是 IndexError
# ====================================================================

def test_add_table_rejects_col_widths_longer_than_headers():
    """OCR 曾产出「表头 1 列、列宽 4 个」，旧代码抛的是 tuple index out of range。"""
    doc = Document()

    with pytest.raises(ValueError, match="col_widths 有 4 项，超过表头的 1 列"):
        WordFormatter(doc).add_table(
            headers=["合并表头"],
            rows=[["a", "b", "c", "d"]],
            col_widths=[1, 2, 3, 4],
        )


def test_add_table_rejects_row_wider_than_headers():
    doc = Document()

    with pytest.raises(ValueError) as excinfo:
        WordFormatter(doc).add_table(headers=["A", "B"], rows=[["1", "2"], ["3", "4", "5"]])

    message = str(excinfo.value)
    assert "第 2 行数据有 3 个单元格" in message
    assert "超过表头的 2 列" in message


def test_add_table_rejects_empty_headers():
    doc = Document()

    with pytest.raises(ValueError, match="headers 不能为空"):
        WordFormatter(doc).add_table(headers=[], rows=[])


def test_add_table_rejects_extra_row_heights():
    doc = Document()

    with pytest.raises(ValueError, match="row_heights 有 5 项，超过表格的 3 行"):
        WordFormatter(doc).add_table(
            headers=["A"], rows=[["1"], ["2"]], row_heights=[1, 1, 1, 1, 1],
        )


def test_add_table_still_allows_ragged_short_rows():
    """短行是合法输入：缺的格子留空即可，不该报错。"""
    doc = Document()

    table = WordFormatter(doc).add_table(
        headers=["A", "B", "C"], rows=[["1"], ["2", "3"]],
    )

    assert len(table.rows) == 3
    assert len(table.columns) == 3


def test_add_table_error_does_not_leave_a_stray_table():
    """校验要在建表之前完成，失败时不应往文档里塞一张空表。"""
    doc = Document()

    with pytest.raises(ValueError):
        WordFormatter(doc).add_table(
            headers=["A", "B"], rows=[["1", "2", "3"]],
        )

    assert len(doc.tables) == 0


# ====================================================================
# 表格列宽/行高：光写 tcW 不够，Word 会按自动布局把列宽冲掉
# ====================================================================

def _tbl_xml(path, index=0):
    from xml.etree import ElementTree as ET
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    return list(root.iter(f"{W}tbl"))[index], W


def test_add_table_locks_fixed_layout_and_writes_tblw(tmp_path):
    """只写 tcW 时 Word 会按 auto 布局重排，列宽白设；必须同时锁 fixed + 写 tblW。"""
    doc = Document()
    WordFormatter(doc).add_table(
        headers=["A", "B", "C", "D"], rows=[["1", "2", "3", "4"]],
        col_widths=[2.5, 4.0, 3.0, 5.0],
    )
    out = tmp_path / "layout.docx"
    doc.save(out)

    tbl, W = _tbl_xml(out)
    tbl_w = tbl.find(f"{W}tblPr/{W}tblW")
    layout = tbl.find(f"{W}tblPr/{W}tblLayout")

    assert tbl_w.get(f"{W}type") == "dxa"
    # 2.5 + 4.0 + 3.0 + 5.0 = 14.5cm
    assert int(tbl_w.get(f"{W}w")) == pytest.approx(14.5 * 566.929, abs=2)
    assert layout.get(f"{W}type") == "fixed"
    grid = [int(g.get(f"{W}w")) for g in tbl.iter(f"{W}gridCol")]
    assert [round(v / 566.929, 1) for v in grid] == [2.5, 4.0, 3.0, 5.0]


def test_add_table_writes_per_cell_width(tmp_path):
    doc = Document()
    WordFormatter(doc).add_table(
        headers=["A", "B", "C"], rows=[["1", "2", "3"]],
        col_widths=[2.0, 6.0, 4.0],
    )
    out = tmp_path / "cells.docx"
    doc.save(out)

    tbl, W = _tbl_xml(out)
    first = next(tbl.iter(f"{W}tr"))
    widths = [round(int(tc.find(f"{W}tcPr/{W}tcW").get(f"{W}w")) / 566.929, 1)
              for tc in first.findall(f"{W}tc")]
    assert widths == [2.0, 6.0, 4.0]


def test_merged_cell_gets_summed_width(tmp_path):
    """合并单元格要写它跨越列的合计宽度，否则 Word 会当成单列宽。"""
    doc = Document()
    WordFormatter(doc).add_table(
        headers=["A", "B", "C", "D"], rows=[["1", "2", "3", "4"]],
        col_widths=[1.0, 2.0, 3.0, 4.0],
        merges=[{"row": 1, "col": 1, "colspan": 3}],
    )
    out = tmp_path / "merged-width.docx"
    doc.save(out)

    tbl, W = _tbl_xml(out)
    second = list(tbl.iter(f"{W}tr"))[1]
    merged = second.findall(f"{W}tc")[1]
    assert int(merged.find(f"{W}tcPr/{W}tcW").get(f"{W}w")) == pytest.approx(
        9.0 * 566.929, abs=2)          # 2+3+4
    assert merged.find(f"{W}tcPr/{W}gridSpan").get(f"{W}val") == "3"


def test_add_table_writes_row_heights(tmp_path):
    doc = Document()
    WordFormatter(doc).add_table(
        headers=["A"], rows=[["1"], ["2"]], row_heights=[1.5, 0.8, 0.8],
    )
    out = tmp_path / "row-heights.docx"
    doc.save(out)

    tbl, W = _tbl_xml(out)
    heights = []
    for tr in tbl.iter(f"{W}tr"):
        h = tr.find(f"{W}trPr/{W}trHeight")
        heights.append(round(int(h.get(f"{W}val")) / 566.929, 2))
    assert heights == [1.5, 0.8, 0.8]


def test_set_table_columns_rejects_bad_widths():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)

    with pytest.raises(ValueError, match="col_widths 不能为空"):
        WordFormatter.set_table_columns(table, [])
    with pytest.raises(ValueError, match="列宽必须全部大于 0"):
        WordFormatter.set_table_columns(table, [2.0, 0])


def test_add_page_break_writes_page_break(tmp_path):
    """模板类文档「一个表单起一页」，漏掉分页符页序就全错。"""
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.body("第一页内容")
    formatter.add_page_break()
    formatter.body("第二页内容")
    out = tmp_path / "page-break.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:type="page"' in document_xml
    assert document_xml.index("第一页内容") < document_xml.index('w:type="page"')
    assert document_xml.index('w:type="page"') < document_xml.index("第二页内容")


def test_add_page_break_can_carry_text(tmp_path):
    doc = Document()
    WordFormatter(doc).add_page_break("新页首行")
    out = tmp_path / "page-break-text.docx"
    doc.save(out)

    assert 'w:type="page"' in _read_zip_xml(out, "word/document.xml")
    assert "新页首行" in _read_zip_xml(out, "word/document.xml")


def test_heading_and_body_support_page_break_before(tmp_path):
    """pageBreakBefore 在段落已处于页首时自动失效，比插独立分页段落安全。"""
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.body("第一页")
    formatter.heading("新页标题", level=1, page_break_before=True)
    formatter.body("接在标题后", page_break_before=True)
    out = tmp_path / "pbb.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert document_xml.count("<w:pageBreakBefore/>") == 2


def test_page_break_before_defaults_off(tmp_path):
    doc = Document()
    WordFormatter(doc).body("普通段落")
    out = tmp_path / "no-pbb.docx"
    doc.save(out)

    assert "pageBreakBefore" not in _read_zip_xml(out, "word/document.xml")


# ====================================================================
# 黄色高亮：标出由 AI 填写的临时数据
# ====================================================================

def test_body_supports_yellow_highlight(tmp_path):
    doc = Document()
    WordFormatter(doc).body("AI 填的临时数据", highlight=True)
    out = tmp_path / "hl.docx"
    doc.save(out)

    assert 'w:highlight w:val="yellow"' in _read_zip_xml(out, "word/document.xml")


def test_set_cell_supports_highlight(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    WordFormatter.set_cell(table.cell(0, 0), "填值", highlight=True)
    out = tmp_path / "hl-cell.docx"
    doc.save(out)

    assert 'w:highlight w:val="yellow"' in _read_zip_xml(out, "word/document.xml")


def test_add_table_supports_per_cell_highlights(tmp_path):
    doc = Document()
    WordFormatter(doc).add_table(
        headers=["姓名", "填值"],
        rows=[["张三", ""], ["李四", ""]],
        cell_highlights=[[False, False], [False, True], [False, True]],
    )
    out = tmp_path / "hl-table.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    # 表头与第一行的第二格不高亮，后两行第二格高亮 => 2 处
    assert document_xml.count('w:highlight w:val="yellow"') == 2


def test_body_segments_highlights_only_filled_value(tmp_path):
    """只该高亮 AI 填的值，标签保持原样。"""
    doc = Document()
    WordFormatter(doc).body_segments(
        [("项目名称：", False), ("雄安新区示范工程", True)],
    )
    out = tmp_path / "segments.docx"
    doc.save(out)

    document_xml = _read_zip_xml(out, "word/document.xml")
    assert document_xml.count('w:highlight w:val="yellow"') == 1
    # 高亮 run 只包住填值，标签在它之前且不带高亮
    label_pos = document_xml.index("项目名称：")
    hl_pos = document_xml.index('w:highlight w:val="yellow"')
    value_pos = document_xml.index("雄安新区示范工程")
    assert label_pos < hl_pos < value_pos


# ====================================================================
# 下划线：还原招投标模板里的「填空线」
# ====================================================================

def test_body_underline_writes_w_u(tmp_path):
    doc = Document()
    WordFormatter(doc).body("整段划线", underline=True)
    out = tmp_path / "u.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert '<w:u ' in xml and 'w:val="single"' in xml


def test_body_without_underline_leaves_runs_clean(tmp_path):
    """默认不能顺手写下划线：Word 会继承样式里的设置，多余一个 w:u 就整篇带线。"""
    doc = Document()
    WordFormatter(doc).body("不划线")
    out = tmp_path / "nou.docx"
    doc.save(out)

    assert '<w:u ' not in _read_zip_xml(out, "word/document.xml")


def test_body_segments_underline_only_marked_segment(tmp_path):
    """填空线只画在填进去的值下面，占位标签不划线——下划线必须是逐片段的。"""
    doc = Document()
    WordFormatter(doc).body_segments([
        {"text": "中国雄安集团生态建设投资有限公司", "underline": True},
        {"text": "（采购人名称）："},
    ])
    out = tmp_path / "seg_u.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert xml.count("<w:u ") == 1
    # 下划线只在第一个 run 上，标签 run 不带
    underline_pos = xml.index("<w:u ")
    label_pos = xml.index("（采购人名称）")
    assert underline_pos < label_pos
    assert xml.index("中国雄安集团生态建设投资有限公司") < label_pos


def test_body_segments_tuple_form_still_works(tmp_path):
    """老的两元组写法（文本, 高亮）不能因为加了 underline 就坏掉。"""
    doc = Document()
    WordFormatter(doc).body_segments([("项目名称：", False), ("雄安", True)])
    out = tmp_path / "seg_tuple.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert xml.count('w:highlight w:val="yellow"') == 1
    assert "<w:u " not in xml


# ====================================================================
# 行距固定值：PDF 里只能量出「相邻行基线差多少 pt」，换算成倍数要依赖字体自身
# 行高，直接写 pt 才精确
# ====================================================================

def test_line_spacing_pt_writes_exact_rule(tmp_path):
    doc = Document()
    WordFormatter(doc).body("固定行距", line_spacing_pt=21.9)
    out = tmp_path / "lspt.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:lineRule="exact"' in xml
    assert 'w:line="438"' in xml          # 21.9pt = 438 twips


def test_line_spacing_pt_overrides_multiple(tmp_path):
    """给了固定值就不能再写倍数，否则两个属性并存、Word 取谁不确定。"""
    doc = Document()
    WordFormatter(doc).body("固定行距", line_spacing=1.5, line_spacing_pt=20)
    out = tmp_path / "lspt2.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:lineRule="exact"' in xml
    assert 'w:lineRule="auto"' not in xml
    assert 'w:line="400"' in xml          # 20pt = 400 twips


def test_multiple_line_spacing_still_available(tmp_path):
    doc = Document()
    WordFormatter(doc).body("倍数行距", line_spacing=1.5)
    out = tmp_path / "lsmulti.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:lineRule="auto"' in xml and 'w:line="360"' in xml


# ====================================================================
# 整段左缩进：还原「整段右移」的表单行（既不是居中，也不是首行缩进）
# ====================================================================

def test_left_indent_pt_writes_w_ind(tmp_path):
    doc = Document()
    # indent=False：这里只验证「整段左缩进」这一个维度，不带首行缩进
    WordFormatter(doc).body("联合体牵头人名称：", left_indent_pt=120.5, indent=False)
    out = tmp_path / "li.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert '<w:ind w:left="2410"' in xml      # 120.5pt = 2410 twips
    assert "w:firstLine" not in xml


def test_left_indent_absent_leaves_indent_unset(tmp_path):
    doc = Document()
    WordFormatter(doc).body("普通正文", indent=False)
    out = tmp_path / "li2.docx"
    doc.save(out)

    # 只在 w:ind 里找：sectPr 的页边距也有 w:left，直接搜 w:left 会误判
    assert "<w:ind" not in _read_zip_xml(out, "word/document.xml")


def test_left_indent_and_first_line_indent_are_independent(tmp_path):
    """整段左缩进和首行缩进是两个维度，同时给要能共存。"""
    doc = Document()
    WordFormatter(doc).body("两段缩进", left_indent_pt=28, indent=True, font_size=14)
    out = tmp_path / "li3.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert '<w:ind w:left="560"' in xml      # 28pt = 560 twips
    assert "w:firstLineChars" in xml


# ====================================================================
# 模板填充：在已有内容上原位改写
# ====================================================================

def test_replace_text_keeps_original_run_format(tmp_path):
    """替换后新文本必须沿用原 run 的字体字号——这是「填完和原文一致」的核心。"""
    doc = Document()
    par = doc.add_paragraph()
    WordFormatter.set_run_font(par.add_run("标段名称："), cn_font="仿宋_GB2312", size="三号")

    assert WordFormatter.replace_text(par, "标段名称：", "标段名称：一标段") == 1
    assert par.text == "标段名称：一标段"

    out = tmp_path / "replace.docx"
    doc.save(out)
    xml = _read_zip_xml(out, "word/document.xml")
    assert "标段名称：一标段" in xml
    assert 'w:eastAsia="仿宋_GB2312"' in xml
    assert 'w:sz w:val="32"' in xml          # 三号 = 16pt = 32 半磅


def test_replace_text_spans_multiple_runs():
    """Word 常把一句话拆成多个 run，跨 run 的匹配同样要能替换。"""
    doc = Document()
    par = doc.add_paragraph()
    WordFormatter.set_run_font(par.add_run("项目名称："), cn_font="宋体", size=14)
    WordFormatter.set_run_font(par.add_run("测试项目"), cn_font="宋体", size=14)

    assert WordFormatter.replace_text(par, "名称：测试", "名称：雄安") == 1
    assert par.text == "项目名称：雄安项目"


def test_replace_text_with_empty_old_appends_to_paragraph():
    """old 传空串 = 追加到段尾，用于「标签：____」这类只在后面补值的填空。"""
    doc = Document()
    par = doc.add_paragraph()
    WordFormatter.set_run_font(par.add_run("委托代理人身份证号码："), cn_font="宋体", size=12)

    assert WordFormatter.replace_text(par, "", "110101199001011234") == 1
    assert par.text == "委托代理人身份证号码：110101199001011234"


def test_replace_text_fills_empty_cell_with_style_font(tmp_path):
    """模板里的空格子（一个 run 都没有）也要能填，字体按样式解析而不是留空。"""
    doc = Document()
    WordFormatter(doc).set_paragraph_style("Normal", font_name="楷体", font_size=15)
    table = doc.add_table(rows=1, cols=1)

    assert WordFormatter.replace_text(table.cell(0, 0).paragraphs[0], "", "赵六") == 1
    assert table.cell(0, 0).text == "赵六"

    out = tmp_path / "empty-cell.docx"
    doc.save(out)
    xml = _read_zip_xml(out, "word/document.xml")
    assert 'w:eastAsia="楷体"' in xml
    assert 'w:sz w:val="30"' in xml          # 15pt = 30 半磅


def test_replace_text_on_empty_paragraph_rejects_non_empty_old():
    doc = Document()
    par = doc.add_paragraph()

    with pytest.raises(ValueError) as excinfo:
        WordFormatter.replace_text(par, "占位", "x")

    assert "old 传空串" in str(excinfo.value)


def test_replace_text_count_zero_replaces_every_occurrence():
    doc = Document()
    par = doc.add_paragraph()
    par.add_run("占位/占位/占位")

    assert WordFormatter.replace_text(par, "占位", "值", count=0) == 3
    assert par.text == "值/值/值"


def test_replace_text_count_limits_replacements():
    doc = Document()
    par = doc.add_paragraph()
    par.add_run("占位/占位/占位")

    assert WordFormatter.replace_text(par, "占位", "值", count=2) == 2
    assert par.text == "值/值/占位"


def test_replace_text_raises_when_source_text_missing():
    doc = Document()
    par = doc.add_paragraph()
    par.add_run("没有这段文字")

    with pytest.raises(ValueError):
        WordFormatter.replace_text(par, "不存在的占位", "x")


def test_replace_text_raises_when_paragraph_has_no_run():
    doc = Document()
    par = doc.add_paragraph()

    with pytest.raises(ValueError):
        WordFormatter.replace_text(par, "任意", "x")


def test_replace_text_highlight_marks_filled_run(tmp_path):
    doc = Document()
    par = doc.add_paragraph()
    par.add_run("项目名称：")

    WordFormatter.replace_text(par, "项目名称：", "项目名称：雄安", highlight=True)
    out = tmp_path / "replace-hl.docx"
    doc.save(out)

    assert 'w:highlight w:val="yellow"' in _read_zip_xml(out, "word/document.xml")


def test_replace_text_without_highlight_leaves_format_clean(tmp_path):
    doc = Document()
    par = doc.add_paragraph()
    par.add_run("项目名称：")

    WordFormatter.replace_text(par, "项目名称：", "项目名称：雄安")
    out = tmp_path / "replace-nohl.docx"
    doc.save(out)

    assert "w:highlight" not in _read_zip_xml(out, "word/document.xml")


def test_insert_cell_image_lands_inside_the_cell(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    WordFormatter.set_cell(table.cell(0, 0), "此格附身份证正面")
    png = _write_test_png(tmp_path / "id.png")

    returned = WordFormatter.insert_cell_image(table.cell(0, 0), png, width=5)

    assert returned.text == ""
    out = tmp_path / "cell-image.docx"
    doc.save(out)
    xml = _read_zip_xml(out, "word/document.xml")
    assert "<w:drawing>" in xml
    assert "此格附身份证正面" not in xml      # 默认清掉格内的占位文字
    with ZipFile(out) as zf:
        assert any(name.startswith("word/media/") for name in zf.namelist())


def test_insert_cell_image_keep_text_preserves_caption(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    WordFormatter.set_cell(table.cell(0, 0), "身份证正面")
    png = _write_test_png(tmp_path / "id2.png")

    WordFormatter.insert_cell_image(table.cell(0, 0), png, width=5, keep_text=True)
    out = tmp_path / "cell-image-keep.docx"
    doc.save(out)

    xml = _read_zip_xml(out, "word/document.xml")
    assert "身份证正面" in xml
    assert "<w:drawing>" in xml


def test_insert_cell_image_rejects_non_positive_width(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    png = _write_test_png(tmp_path / "id3.png")

    with pytest.raises(ValueError):
        WordFormatter.insert_cell_image(table.cell(0, 0), png, width=0)


def test_describe_blocks_lists_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("第二段")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("费率")
    table.cell(0, 1).paragraphs[0].add_run("%")

    blocks = WordFormatter(doc).describe_blocks()

    assert [block["type"] for block in blocks] == ["paragraph", "paragraph", "table"]
    assert [block["index"] for block in blocks[:2]] == [0, 1]
    assert blocks[0]["text"] == "第一段"
    assert blocks[2]["index"] == 0
    assert blocks[2]["grid"] == [["费率", "%"]]
    assert blocks[2]["rows"] == 1 and blocks[2]["columns"] == 2


def test_describe_blocks_skips_empty_paragraphs_by_default():
    doc = Document()
    doc.add_paragraph("")
    doc.add_paragraph("有内容")

    assert [b["index"] for b in WordFormatter(doc).describe_blocks()] == [1]
    assert [b["index"] for b in
            WordFormatter(doc).describe_blocks(include_empty=True)] == [0, 1]


def test_describe_blocks_resolves_font_inherited_from_style():
    """模板里的 run 常常不写字体，全靠样式继承；这里必须回落到实际生效的字体。"""
    doc = Document()
    WordFormatter(doc).set_paragraph_style("Normal", font_name="楷体", font_size=15)
    doc.add_paragraph("项目名称：")

    run = WordFormatter(doc).describe_blocks()[0]["runs"][0]

    assert run["font_name"] == "楷体"
    assert run["font_size_pt"] == 15.0


def test_describe_blocks_truncates_long_text():
    doc = Document()
    doc.add_paragraph("很长的内容" * 20)

    block = WordFormatter(doc).describe_blocks(max_text_chars=10)[0]

    assert block["text"] == "很长的内容很长的内容很"[:10] + "…"
