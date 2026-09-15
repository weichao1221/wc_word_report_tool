"""MCP 会话层：把有状态的 ``WordFormatter`` 包装成以 ``doc_id`` 引用的会话。

MCP 工具是无状态 RPC，而 :class:`~wc_word_report_tool.word.WordFormatter` 绑定在
具体的 ``Document`` 上。本模块用 ``doc_id -> Session`` 的内存映射抹平这个差异。

本模块不依赖 ``mcp`` SDK，可独立测试。
"""

from __future__ import annotations

import base64
import binascii
import os
import re
import shutil
import tempfile
import threading
import time
import uuid
from pathlib import Path

from docx import Document

from ..word import WordFormatter

DEFAULT_OUTPUT_DIRNAME = "wc-reports"
DEFAULT_TTL_SECONDS = 3600
DEFAULT_MAX_SESSIONS = 32

_DATA_URI_RE = re.compile(
    r"^data:image/(?P<ext>[a-zA-Z0-9.+-]+);base64,(?P<data>.*)$",
    re.DOTALL,
)
_EXT_ALIASES = {"jpg": "jpeg", "svg+xml": "svg", "x-png": "png"}


class ReportSessionError(Exception):
    """会话层错误，消息面向使用 MCP 的 AI，尽量给出可执行的下一步。"""


def output_root() -> Path:
    """MCP 保存文档的默认根目录，可用 ``WC_REPORT_MCP_OUTPUT_DIR`` 覆盖。"""
    configured = os.environ.get("WC_REPORT_MCP_OUTPUT_DIR")
    if configured:
        return Path(configured).expanduser()
    return Path.home() / DEFAULT_OUTPUT_DIRNAME


def resolve_output_path(name: str | None) -> Path:
    """解析保存路径：相对路径挂到输出根目录，自动补 ``.docx`` 后缀并建父目录。"""
    if name:
        candidate = Path(name).expanduser()
        path = candidate if candidate.is_absolute() else output_root() / candidate
    else:
        stamp = time.strftime("%Y%m%d-%H%M%S")
        path = output_root() / f"report-{stamp}.docx"
    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def resolve_input_path(name: str, *, what: str = "文件") -> Path:
    """解析已存在的输入路径：绝对路径直接用，相对路径依次尝试 cwd 与输出根目录。"""
    path = Path(name).expanduser()
    if not path.is_absolute():
        cwd_candidate = Path.cwd() / path
        path = cwd_candidate if cwd_candidate.is_file() else output_root() / path
    if not path.is_file():
        raise ReportSessionError(f"{what}不存在: {path}")
    return path


def default_logo_path() -> Path | None:
    """``WC_REPORT_MCP_DEFAULT_LOGO`` 指定的默认 Logo，未配置时返回 None。"""
    configured = os.environ.get("WC_REPORT_MCP_DEFAULT_LOGO")
    if not configured:
        return None
    path = Path(configured).expanduser()
    return path if path.is_file() else None


class Session:
    """一份正在编辑中的文档。"""

    def __init__(self, doc_id: str, doc, source_path: Path | None = None):
        self.doc_id = doc_id
        self.doc = doc
        self.source_path = source_path
        self.saved_path: Path | None = None
        self.created_at = time.time()
        self.last_used = self.created_at
        self._tmp_dir: Path | None = None

    @property
    def formatter(self) -> WordFormatter:
        return WordFormatter(self.doc)

    @property
    def tmp_dir(self) -> Path:
        """存放 base64 图片落盘结果的会话级临时目录。"""
        if self._tmp_dir is None:
            self._tmp_dir = Path(tempfile.mkdtemp(prefix=f"wc-report-{self.doc_id}-"))
        return self._tmp_dir

    def cleanup(self) -> None:
        if self._tmp_dir is not None:
            shutil.rmtree(self._tmp_dir, ignore_errors=True)
            self._tmp_dir = None

    def describe(self) -> dict:
        sections = self.doc.sections
        return {
            "doc_id": self.doc_id,
            "paragraphs": len(self.doc.paragraphs),
            "sections": len(sections),
            "tables": len(self.doc.tables),
            "saved_path": str(self.saved_path) if self.saved_path else None,
            "source_path": str(self.source_path) if self.source_path else None,
            "age_seconds": round(time.time() - self.created_at, 1),
        }


class SessionStore:
    """``doc_id -> Session`` 的内存表，带 TTL 回收和串行锁。

    锁的意义：FastMCP 把同步工具函数放在线程池执行，并发调用会同时改写同一个
    ``Document``。文档构建本身就该串行，因此用一把全局可重入锁即可。
    """

    def __init__(
        self,
        *,
        ttl_seconds: int = DEFAULT_TTL_SECONDS,
        max_sessions: int = DEFAULT_MAX_SESSIONS,
    ):
        self._sessions: dict[str, Session] = {}
        self._ttl_seconds = ttl_seconds
        self._max_sessions = max_sessions
        self.lock = threading.RLock()

    def create(self, *, source_path: Path | None = None) -> Session:
        with self.lock:
            self._reap_locked()
            if len(self._sessions) >= self._max_sessions:
                oldest = min(self._sessions.values(), key=lambda s: s.last_used)
                self._discard_locked(oldest.doc_id)
            doc_id = uuid.uuid4().hex[:12]
            doc = Document(str(source_path)) if source_path else Document()
            session = Session(doc_id, doc, source_path=source_path)
            self._sessions[doc_id] = session
            return session

    def get(self, doc_id: str) -> Session:
        with self.lock:
            self._reap_locked()
            session = self._sessions.get(doc_id)
            if session is None:
                known = ", ".join(sorted(self._sessions)) or "（无）"
                raise ReportSessionError(
                    f"未知的 doc_id: {doc_id!r}。当前可用: {known}。"
                    "请先用 word_create_report 或 word_open_report 创建会话。"
                )
            session.last_used = time.time()
            return session

    def close(self, doc_id: str) -> None:
        with self.lock:
            if doc_id not in self._sessions:
                raise ReportSessionError(f"未知的 doc_id: {doc_id!r}，无法关闭。")
            self._discard_locked(doc_id)

    def close_all(self) -> None:
        with self.lock:
            for doc_id in list(self._sessions):
                self._discard_locked(doc_id)

    def _discard_locked(self, doc_id: str) -> None:
        session = self._sessions.pop(doc_id, None)
        if session is not None:
            session.cleanup()

    def _reap_locked(self) -> None:
        if self._ttl_seconds <= 0:
            return
        deadline = time.time() - self._ttl_seconds
        for doc_id, session in list(self._sessions.items()):
            if session.last_used < deadline:
                self._discard_locked(doc_id)


def materialize_image(
    session: Session,
    *,
    image_path: str | None = None,
    image_base64: str | None = None,
    label: str = "图片",
) -> Path:
    """把 ``image_path`` 或 ``image_base64`` 统一解析成本地文件路径。

    base64 支持裸串和 ``data:image/png;base64,...`` 两种写法，落盘到会话临时目录。
    """
    if image_path and image_base64:
        raise ReportSessionError(f"{label}的 image_path 与 image_base64 只能传一个。")
    if image_path:
        return resolve_input_path(image_path, what=label)
    if not image_base64:
        raise ReportSessionError(f"必须提供 {label} 的 image_path 或 image_base64。")

    payload = image_base64.strip()
    extension = "png"
    match = _DATA_URI_RE.match(payload)
    if match:
        raw_ext = match.group("ext").lower()
        extension = _EXT_ALIASES.get(raw_ext, raw_ext)
        payload = match.group("data")

    try:
        raw = base64.b64decode(payload, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ReportSessionError(f"{label} base64 解码失败: {exc}") from exc
    if not raw:
        raise ReportSessionError(f"{label}内容为空。")

    target = session.tmp_dir / f"{uuid.uuid4().hex[:8]}.{extension}"
    target.write_bytes(raw)
    return target
