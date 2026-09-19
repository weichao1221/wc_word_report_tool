"""MCP 工具层测试：直接调用工具函数，不经过 MCP 协议。

工具函数是普通 Python 函数（不依赖 mcp SDK），所以可以像测普通库一样测试它们。
协议层的冒烟测试见 tests/test_mcp_server.py。
"""

import base64
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from wc_word_report_tool import WordFormatter
from wc_word_report_tool.mcp import tools
from wc_word_report_tool.mcp.errors import ToolError
from wc_word_report_tool.mcp.session import ReportSessionError, SessionStore

TEST_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "/x8AAusB9Y9Z4k8AAAAASUVORK5CYII="
)


@pytest.fixture(autouse=True)
def clean_store(monkeypatch, tmp_path):
    """每个测试用独立的输出目录和空会话表。"""
    monkeypatch.setenv("WC_REPORT_MCP_OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.delenv("WC_REPORT_MCP_DEFAULT_LOGO", raising=False)
    tools.reset_store()
    yield
    tools.reset_store()


def _read_zip_xml(docx_path: Path, inner_path: str) -> str:
    with ZipFile(docx_path) as zf:
        return zf.read(inner_path).decode("utf-8", errors="ignore")


def _zip_names(docx_path: Path) -> list:
    with ZipFile(docx_path) as zf:
        return zf.namelist()


# ====================================================================
# 会话生命周期
# ====================================================================


def test_create_report_returns_usable_doc_id():
    created = tools.word_create_report(title="测试报告", author="willcha")

    assert created["doc_id"]
    assert created["sections"] == 1
    assert created["paragraphs"] == 0

    described = tools.word_describe_report(created["doc_id"])
    assert described["doc_id"] == created["doc_id"]
    assert described["paragraphs_per_section"] == [0]


def test_unknown_doc_id_gives_actionable_error():
    with pytest.raises(ToolError) as excinfo:
        tools.word_add_body("does-not-exist", "正文")

    assert "未知的 doc_id" in str(excinfo.value)
    assert "word_create_report" in str(excinfo.value)


def test_close_report_invalidates_doc_id():
    doc_id = tools.word_create_report()["doc_id"]

    assert tools.word_close_report(doc_id)["closed"] == doc_id

    with pytest.raises(ToolError):
        tools.word_describe_report(doc_id)


def test_save_without_path_reuses_previous_target(tmp_path):
    doc_id = tools.word_create_report()["doc_id"]

    first = tools.word_save_report(doc_id)
    second = tools.word_save_report(doc_id)

    assert first["saved_to"] == second["saved_to"]
    assert Path(first["saved_to"]).parent == tmp_path / "out"
    assert first["size_bytes"] > 0


def test_store_reaps_expired_sessions():
    store = SessionStore(ttl_seconds=60)
    session = store.create()
    session.last_used -= 120

    with pytest.raises(ReportSessionError):
        store.get(session.doc_id)


def test_store_evicts_oldest_when_full():
    store = SessionStore(max_sessions=2)
    oldest = store.create()
    newest = store.create()
    newest.last_used = oldest.last_used + 10

    survivor = store.create()

    assert store.get(newest.doc_id).doc_id == newest.doc_id
    assert survivor.doc_id != oldest.doc_id
    with pytest.raises(ReportSessionError):
        store.get(oldest.doc_id)


# ====================================================================
# 无状态工具
# ====================================================================


def test_rmb_upper_needs_no_doc_id():
    assert tools.word_rmb_upper(123456.78)["chinese_upper"] == (
        "壹拾贰万叁仟肆佰伍拾陆元柒角捌分"
    )
    assert tools.word_rmb_upper("1,234.5")["chinese_upper"] == "壹仟贰佰叁拾肆元伍角"


def test_describe_api_lists_formatter_methods():
    text = tools.word_describe_api()

    for method in (
        "heading",
        "body",
        "add_table",
        "add_toc",
        "fengmian_jiesuan",
        "qianfaye",
        "set_page_number_from_section",
    ):
        assert method in text


# ====================================================================
# 兜底工具
# ====================================================================


def test_word_call_invokes_arbitrary_method():
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_call(doc_id, "body", {"text": "来自兜底工具", "indent": False})

    assert result["method"] == "body"
    assert result["paragraphs"] == 1


def test_word_call_resolves_section_index_to_object():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_call(doc_id, "set_header", {"section": -1, "text": "兜底页眉"})

    saved = tools.word_save_report(doc_id, "call-header.docx")
    head_path = "word/header1.xml"
    assert head_path in _zip_names(Path(saved["saved_to"]))
    assert "兜底页眉" in _read_zip_xml(Path(saved["saved_to"]), head_path)


def test_word_call_rejects_private_and_unknown_methods():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError):
        tools.word_call(doc_id, "_ensure_update_fields_on_open", {})
    with pytest.raises(ToolError):
        tools.word_call(doc_id, "not_a_real_method", {})


# ====================================================================
# 内容构建
# ====================================================================


def test_add_body_defaults_use_character_indent(tmp_path):
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_body(doc_id, "正文内容")

    saved = tools.word_save_report(doc_id, "body.docx")
    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")

    assert 'w:eastAsia="宋体"' in document_xml
    assert 'w:firstLineChars="200"' in document_xml


def test_add_body_list_writes_every_paragraph():
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_add_body_list(doc_id, ["第一段。", "第二段。", "第三段。"])

    assert result["added"] == 3
    assert result["paragraphs"] == 3


def test_add_table_reports_shape():
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_add_table(
        doc_id,
        headers=["序号", "项目", "金额（万元）"],
        rows=[["1", "建筑工程", "1200.50"], ["2", "安装工程", "850.00"]],
        col_widths=[2, 6, 4],
    )

    assert result["rows"] == 3
    assert result["columns"] == 3
    assert result["table_index"] == 0


def test_format_cell_out_of_range_explains_bounds():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_table(doc_id, ["标题"], [["内容"]])

    with pytest.raises(ToolError) as excinfo:
        tools.word_format_cell(doc_id, row=5, col=0, text="越界")

    assert "超出范围" in str(excinfo.value)
    assert "2 行" in str(excinfo.value)


def test_add_date_chinese_style():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_date(doc_id, value="2026-07-16", style="chinese")

    saved = tools.word_save_report(doc_id, "date.docx")

    assert "二零二六年七月十六日" in _read_zip_xml(
        Path(saved["saved_to"]), "word/document.xml"
    )


def test_add_date_rejects_unknown_style():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_add_date(doc_id, style="lunar")

    assert "chinese" in str(excinfo.value)


def test_target_validation_is_converted_to_tool_error():
    """参数校验发生在会话上下文内，才能被转换成模型可见的 ToolError。"""
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_clear_header_footer(doc_id, target="sidebar")
    assert "header/footer/both" in str(excinfo.value)

    with pytest.raises(ToolError) as excinfo:
        tools.word_set_header_image(doc_id, target="sidebar", width_cm=2)
    assert "header/footer" in str(excinfo.value)


def test_section_index_out_of_range():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_set_header(doc_id, "页眉", section_index=9)

    assert "section_index 9 超出范围" in str(excinfo.value)


# ====================================================================
# 图片双通道
# ====================================================================


def test_insert_image_accepts_base64():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_insert_image(doc_id, width_cm=3, image_base64=TEST_PNG_B64)
    saved = tools.word_save_report(doc_id, "image.docx")

    assert any(
        name.startswith("word/media/") for name in _zip_names(Path(saved["saved_to"]))
    )


def test_insert_image_accepts_data_uri(tmp_path):
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_insert_image(
        doc_id, width_cm=3, image_base64=f"data:image/png;base64,{TEST_PNG_B64}"
    )
    saved = tools.word_save_report(doc_id, "image-datauri.docx")

    assert any(
        name.startswith("word/media/") for name in _zip_names(Path(saved["saved_to"]))
    )


def test_insert_image_requires_a_source():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_insert_image(doc_id, width_cm=3)

    assert "image_path" in str(excinfo.value)


def test_insert_image_rejects_both_sources():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_insert_image(
            doc_id, width_cm=3, image_path="a.png", image_base64=TEST_PNG_B64
        )

    assert "只能传一个" in str(excinfo.value)


def test_cover_page_without_logo_explains_the_fix():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_add_cover_page(doc_id)

    assert "Logo" in str(excinfo.value)
    assert "WC_REPORT_MCP_DEFAULT_LOGO" in str(excinfo.value)


def test_cover_page_falls_back_to_default_logo_env(monkeypatch, tmp_path):
    logo = tmp_path / "logo.png"
    import base64

    logo.write_bytes(base64.b64decode(TEST_PNG_B64))
    monkeypatch.setenv("WC_REPORT_MCP_DEFAULT_LOGO", str(logo))

    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_cover_page(doc_id, project_name="某某工程")
    saved = tools.word_save_report(doc_id, "cover.docx")

    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")
    assert "某某工程" in document_xml


# ====================================================================
# 端到端
# ====================================================================


def test_full_report_end_to_end():
    doc_id = tools.word_create_report(title="测试项目结算审核报告")["doc_id"]

    tools.word_add_cover_page(
        doc_id,
        project_name="某某工程",
        entrusting_unit="某委托单位",
        compiling_unit="某编制单位",
        date="2026-07-16",
        image_base64=TEST_PNG_B64,
    )
    tools.word_set_default_font(doc_id, font_size="三号", cn_font="仿宋_GB2312")
    tools.word_insert_section(doc_id, inherit_header=False, inherit_footer=False)
    tools.word_add_toc(doc_id, title="目  录", levels_min=1, levels_max=3)
    tools.word_add_heading(doc_id, "一、项目概况", level=1)
    tools.word_add_body_list(doc_id, ["本项目位于……。", "项目背景说明……。"])
    tools.word_add_table(
        doc_id,
        headers=["序号", "项目", "金额（万元）"],
        rows=[["1", "建筑工程", "1200.50"], ["2", "安装工程", "850.00"]],
    )
    tools.word_set_header(doc_id, "某某工程 结算审核报告", section_index=-1)
    tools.word_set_page_numbers(
        doc_id, start_section_index=1, start=1, prefix="第 ", suffix=" 页",
    )

    saved = tools.word_save_report(doc_id, "demo.docx")
    path = Path(saved["saved_to"])

    assert path.is_file()
    assert path.suffix == ".docx"

    document_xml = _read_zip_xml(path, "word/document.xml")
    assert "一、项目概况" in document_xml
    assert 'w:firstLineChars="200"' in document_xml
    assert "TOC" in document_xml

    settings_xml = _read_zip_xml(path, "word/settings.xml")
    assert "w:updateFields" in settings_xml
    assert "zh-CN" in settings_xml

    footers = [name for name in _zip_names(path) if name.startswith("word/footer")]
    assert footers
    assert any("PAGE" in _read_zip_xml(path, name) for name in footers)


def test_open_report_round_trips_an_existing_file():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_body(doc_id, "原始内容")
    saved = tools.word_save_report(doc_id, "source.docx")

    reopened = tools.word_open_report(saved["saved_to"])
    assert reopened["source_path"] == saved["saved_to"]

    tools.word_add_body(reopened["doc_id"], "追加内容")

    again = tools.word_save_report(reopened["doc_id"], "appended.docx")
    document_xml = _read_zip_xml(Path(again["saved_to"]), "word/document.xml")

    assert "原始内容" in document_xml
    assert "追加内容" in document_xml


# ====================================================================
# A1 · 页眉页脚：三段分列 / 首页 / 奇偶页
# ====================================================================


def test_set_header_parts_writes_tabs(tmp_path):
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_set_header_parts(
        doc_id, left="公司名", center="报告名", right="第 1 页",
    )
    saved = tools.word_save_report(doc_id, "header-parts.docx")

    assert result["variant"] == "primary"
    header_xml = "".join(
        _read_zip_xml(Path(saved["saved_to"]), name)
        for name in _zip_names(Path(saved["saved_to"]))
        if name.startswith("word/header")
    )
    assert header_xml.count("<w:tab/>") == 2
    assert "公司名" in header_xml and "报告名" in header_xml


def test_set_footer_parts_writes_three_segments():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_set_footer_parts(doc_id, left="左", center="中", right="右")
    saved = tools.word_save_report(doc_id, "footer-parts.docx")

    footer_xml = "".join(
        _read_zip_xml(Path(saved["saved_to"]), name)
        for name in _zip_names(Path(saved["saved_to"]))
        if name.startswith("word/footer")
    )
    assert "左" in footer_xml and "中" in footer_xml and "右" in footer_xml


def test_different_first_page_and_odd_even_flags():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_set_different_first_page(doc_id, True)
    tools.word_set_different_odd_even(doc_id, True)
    saved = tools.word_save_report(doc_id, "flags.docx")
    path = Path(saved["saved_to"])

    assert "<w:titlePg/>" in _read_zip_xml(path, "word/document.xml")
    assert "evenAndOddHeaders" in _read_zip_xml(path, "word/settings.xml")


def test_header_variant_first_and_even_are_separate():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_set_different_first_page(doc_id, True)
    tools.word_set_different_odd_even(doc_id, True)
    tools.word_set_header(doc_id, "默认页眉")
    tools.word_set_header(doc_id, "首页页眉", variant="first")
    tools.word_set_header(doc_id, "偶数页页眉", variant="even")
    saved = tools.word_save_report(doc_id, "header-variants.docx")

    from docx import Document as _Read

    section = _Read(saved["saved_to"]).sections[0]
    assert section.header.paragraphs[0].text == "默认页眉"
    assert section.first_page_header.paragraphs[0].text == "首页页眉"
    assert section.even_page_header.paragraphs[0].text == "偶数页页眉"


def test_unknown_header_variant_becomes_tool_error():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_set_header(doc_id, "页眉", variant="sidebar")

    assert "未知的页眉页脚变体" in str(excinfo.value)


def test_clear_header_footer_accepts_variant():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_set_different_first_page(doc_id, True)
    tools.word_set_header(doc_id, "首页页眉", variant="first")

    tools.word_clear_header_footer(doc_id, target="header", variant="first")
    saved = tools.word_save_report(doc_id, "cleared-first.docx")

    from docx import Document as _Read

    section = _Read(saved["saved_to"]).sections[0]
    assert section.first_page_header.paragraphs[0].text == ""


# ====================================================================
# A2 · 页码编号格式
# ====================================================================


def test_page_number_format_romans_are_persisted():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_set_page_number_format(doc_id, "upperRoman", start=1)
    saved = tools.word_save_report(doc_id, "roman.docx")

    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")
    assert 'w:fmt="upperRoman"' in document_xml


def test_restart_page_numbering_accepts_number_format():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_restart_page_numbering(
        doc_id, start=2, number_format="小写罗马",
    )
    saved = tools.word_save_report(doc_id, "roman-restart.docx")

    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")
    assert 'w:fmt="lowerRoman"' in document_xml
    assert 'w:start="2"' in document_xml


def test_set_page_numbers_accepts_number_format():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_body(doc_id, "正文")

    tools.word_set_page_numbers(
        doc_id, start_section_index=0, start=1, number_format="chineseCounting",
    )
    saved = tools.word_save_report(doc_id, "chinese-numbering.docx")

    assert 'w:fmt="chineseCounting"' in _read_zip_xml(
        Path(saved["saved_to"]), "word/document.xml"
    )


def test_unknown_number_format_becomes_tool_error():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_set_page_number_format(doc_id, "roman99")

    assert "未知的页码格式" in str(excinfo.value)


# ====================================================================
# A3 · 分节纸张与页边距
# ====================================================================


def test_insert_section_supports_landscape_and_margins():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_body(doc_id, "第一节")

    result = tools.word_insert_section(
        doc_id, orientation="横向",
        margin_top_cm=3.17, margin_bottom_cm=3.17,
        margin_left_cm=2.54, margin_right_cm=2.54,
    )
    saved = tools.word_save_report(doc_id, "landscape-section.docx")
    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")

    assert result["section_index"] == 1
    assert 'w:orient="landscape"' in document_xml


def test_set_section_page_reports_new_size():
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_set_section_page(doc_id, section_index=0, orientation="横向")

    assert result["page_width_cm"] > result["page_height_cm"]


def test_set_section_page_rejects_unknown_orientation():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_set_section_page(doc_id, orientation="斜向")

    assert "未知的纸张方向" in str(excinfo.value)


def test_set_section_page_index_out_of_range():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_set_section_page(doc_id, section_index=7)

    assert "section_index 7 超出范围" in str(excinfo.value)


# ====================================================================
# B1 + B4 · 标题对齐与段间距
# ====================================================================


def test_add_heading_supports_alignment():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_heading(doc_id, "居中大标题", level=1, alignment="居中")
    saved = tools.word_save_report(doc_id, "heading-align.docx")

    assert 'w:jc w:val="center"' in _read_zip_xml(
        Path(saved["saved_to"]), "word/document.xml"
    )


def test_add_heading_and_body_accept_space_params():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_heading(doc_id, "标题", space_before=18, space_after=6)
    tools.word_add_body(doc_id, "正文", space_before=3, space_after=9)
    saved = tools.word_save_report(doc_id, "space.docx")
    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")

    assert 'w:before="360"' in document_xml
    assert 'w:after="120"' in document_xml


def test_add_body_list_accepts_space_params():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_body_list(doc_id, ["第一段", "第二段"], space_after=6)
    saved = tools.word_save_report(doc_id, "space-list.docx")

    assert 'w:after="120"' in _read_zip_xml(
        Path(saved["saved_to"]), "word/document.xml"
    )


# ====================================================================
# B2 + B5 · 表格合并 / 总宽 / 行高 / 逐格字体
# ====================================================================


def test_add_table_supports_merges():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_table(
        doc_id,
        headers=["序号", "项目", "金额", "备注"],
        rows=[["1", "合并行", "100", "a"], ["2", "x", "200", "b"]],
        merges=[{"row": 1, "col": 1, "rowspan": 1, "colspan": 2}],
    )
    saved = tools.word_save_report(doc_id, "merged.docx")

    assert 'w:gridSpan w:val="2"' in _read_zip_xml(
        Path(saved["saved_to"]), "word/document.xml"
    )


def test_add_table_supports_width_height_and_cell_fonts():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_table(
        doc_id,
        headers=["A", "B"],
        rows=[["1", "2"]],
        font_name="宋体",
        table_width_cm=12,
        row_heights=[1.0, 1.5],
        cell_font_names=[["黑体", "黑体"], ["楷体_GB2312", None]],
    )
    saved = tools.word_save_report(doc_id, "table-extras.docx")
    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")

    assert 'w:w="6804"' in document_xml          # 12cm
    assert 'w:hRule="atLeast"' in document_xml
    assert 'w:eastAsia="黑体"' in document_xml
    assert 'w:eastAsia="楷体_GB2312"' in document_xml


def test_add_table_merge_out_of_range_becomes_tool_error():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_add_table(
            doc_id, headers=["A", "B"], rows=[["1", "2"]],
            merges=[{"row": 0, "col": 0, "colspan": 5}],
        )

    assert "超出表格范围" in str(excinfo.value)


def test_merge_cells_tool_keeps_only_anchor_text():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_table(doc_id, headers=["A", "B", "C"], rows=[["1", "2", "3"]])

    tools.word_merge_cells(doc_id, merges=[{"row": 0, "col": 0, "colspan": 3}])
    saved = tools.word_save_report(doc_id, "merge-existing.docx")

    from docx import Document as _Read

    table = _Read(saved["saved_to"]).tables[0]
    assert table.cell(0, 0).text == "A"


def test_merge_cells_requires_an_existing_table():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_merge_cells(doc_id, merges=[{"row": 0, "col": 0, "colspan": 2}])

    assert "还没有表格" in str(excinfo.value)


# ====================================================================
# B3 · 图片自动尺寸
# ====================================================================


def test_insert_image_without_width_is_auto_sized():
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_insert_image(doc_id, image_base64=TEST_PNG_B64)
    saved = tools.word_save_report(doc_id, "auto-image.docx")

    assert result["auto_sized"] is True
    assert any(
        name.startswith("word/media/") for name in _zip_names(Path(saved["saved_to"]))
    )


def test_insert_image_accepts_height_only():
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_insert_image(doc_id, height_cm=2, image_base64=TEST_PNG_B64)
    saved = tools.word_save_report(doc_id, "height-image.docx")

    assert result["auto_sized"] is False
    assert any(
        name.startswith("word/media/") for name in _zip_names(Path(saved["saved_to"]))
    )


def test_insert_image_rejects_non_positive_width():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_insert_image(doc_id, width_cm=0, image_base64=TEST_PNG_B64)

    assert "图片宽度必须大于 0" in str(excinfo.value)


# ====================================================================
# 端到端 · 横向插页 + 罗马数字页码 + 合并表
# ====================================================================


def test_landscape_insert_and_roman_numbering_end_to_end():
    doc_id = tools.word_create_report(title="混排报告")["doc_id"]

    tools.word_set_different_first_page(doc_id, True)
    tools.word_set_header_parts(doc_id, left="公司", center="报告", right="机密")
    tools.word_set_header(doc_id, "首页不显示页码", variant="first")
    tools.word_set_page_numbers(
        doc_id, start_section_index=0, start=1, number_format="upperRoman",
    )
    tools.word_add_heading(doc_id, "一、横向插页", level=1, alignment="居中")
    tools.word_insert_section(
        doc_id, orientation="横向",
        margin_top_cm=3.17, margin_bottom_cm=3.17,
        margin_left_cm=2.54, margin_right_cm=2.54,
        add_page_number=True, restart_page_number=True,
        start_page_number=1, number_format="decimal",
    )
    tools.word_add_table(
        doc_id,
        headers=["序号", "项目", "金额", "备注"],
        rows=[["1", "合并", "100", "x"]],
        merges=[{"row": 1, "col": 1, "colspan": 2}],
        table_width_cm=15,
    )

    saved = tools.word_save_report(doc_id, "mixed-layout.docx")
    path = Path(saved["saved_to"])
    document_xml = _read_zip_xml(path, "word/document.xml")

    assert path.is_file()
    assert 'w:orient="landscape"' in document_xml
    assert 'w:fmt="upperRoman"' in document_xml
    assert 'w:gridSpan w:val="2"' in document_xml
    assert "<w:titlePg/>" in document_xml


def test_insert_image_supports_floating_for_seals():
    """印章/图表需要浮于文字之上时走 floating。"""
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_body(doc_id, "落款：某某公司")

    result = tools.word_insert_image(
        doc_id, width_cm=3, alignment="居右",
        image_base64=TEST_PNG_B64, floating=True, y_offset_pt=6,
    )
    saved = tools.word_save_report(doc_id, "seal.docx")

    assert result["floating"] is True
    assert "w:pict" in _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")


def test_insert_image_is_inline_by_default():
    doc_id = tools.word_create_report()["doc_id"]

    result = tools.word_insert_image(doc_id, width_cm=3, image_base64=TEST_PNG_B64)
    saved = tools.word_save_report(doc_id, "inline-image.docx")

    assert result["floating"] is False
    assert "w:pict" not in _read_zip_xml(
        Path(saved["saved_to"]), "word/document.xml"
    )


def test_add_table_grid_mismatch_becomes_readable_tool_error():
    """列数对不上时必须给 AI 看得懂的错误，而不是 IndexError。"""
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_add_table(
            doc_id,
            headers=["合并表头"],
            rows=[["a", "b", "c", "d"]],
            col_widths=[1, 2, 3, 4],
        )

    message = str(excinfo.value)
    assert "col_widths 有 4 项" in message
    assert "超过表头的 1 列" in message
    assert "tuple index out of range" not in message


def test_add_table_row_width_mismatch_reports_row_number():
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_add_table(doc_id, headers=["A", "B"], rows=[["1", "2"], ["3", "4", "5"]])

    assert "第 2 行数据有 3 个单元格" in str(excinfo.value)


def test_failed_add_table_does_not_corrupt_the_session():
    """工具报错后会话仍可用，且没有留下半张表。"""
    doc_id = tools.word_create_report()["doc_id"]

    with pytest.raises(ToolError):
        tools.word_add_table(doc_id, headers=["A", "B"], rows=[["1", "2", "3"]])

    described = tools.word_describe_report(doc_id)
    assert described["tables"] == 0

    tools.word_add_table(doc_id, headers=["A", "B"], rows=[["1", "2"]])
    assert tools.word_describe_report(doc_id)["tables"] == 1


def test_add_page_break_tool():
    doc_id = tools.word_create_report()["doc_id"]
    tools.word_add_body(doc_id, "第一页")
    tools.word_add_page_break(doc_id)
    tools.word_add_body(doc_id, "第二页")
    saved = tools.word_save_report(doc_id, "page-break.docx")

    assert 'w:type="page"' in _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")


def test_heading_page_break_before_only_on_first_body_paragraph():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_heading(doc_id, "新页标题", page_break_before=True)
    tools.word_add_body_list(doc_id, ["第一段", "第二段"], page_break_before=True)
    saved = tools.word_save_report(doc_id, "pbb-list.docx")

    document_xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")
    # 标题 1 处 + 正文首段 1 处，第二段不分页
    assert document_xml.count("<w:pageBreakBefore/>") == 2


def test_add_body_list_highlights_ai_filled_data():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_body_list(doc_id, ["项目名称：雄安新区示范工程"], highlight=True)
    saved = tools.word_save_report(doc_id, "hl-body.docx")

    assert 'w:highlight w:val="yellow"' in _read_zip_xml(
        Path(saved["saved_to"]), "word/document.xml")


def test_add_table_cell_highlights_tool():
    doc_id = tools.word_create_report()["doc_id"]

    tools.word_add_table(
        doc_id, headers=["项目", "填值"], rows=[["单位名称", "某某公司"]],
        cell_highlights=[[False, False], [False, True]],
    )
    saved = tools.word_save_report(doc_id, "hl-table.docx")

    assert _read_zip_xml(Path(saved["saved_to"]), "word/document.xml").count(
        'w:highlight w:val="yellow"') == 1


# ====================================================================
# 模板填充
# ====================================================================


def _make_template(tmp_path: Path) -> Path:
    """造一份「自带字体、留着空」的小模板，模拟招投标文件。"""
    document = Document()
    par = document.add_paragraph()
    WordFormatter.set_run_font(
        par.add_run("标段名称："), cn_font="仿宋_GB2312", size="三号")
    table = document.add_table(rows=1, cols=2)
    WordFormatter.set_cell(table.cell(0, 0), "费率：")
    WordFormatter.set_cell(table.cell(0, 1), "%")
    source = tmp_path / "template.docx"
    document.save(str(source))
    return source


def test_list_blocks_reports_paragraph_indexes_and_table_grid(tmp_path):
    source = _make_template(tmp_path)
    doc_id = tools.word_open_report(str(source))["doc_id"]

    result = tools.word_list_blocks(doc_id)

    paragraphs = [b for b in result["blocks"] if b["type"] == "paragraph"]
    tables = [b for b in result["blocks"] if b["type"] == "table"]
    assert paragraphs[0]["index"] == 0
    assert paragraphs[0]["text"] == "标段名称："
    assert paragraphs[0]["runs"][0]["font_name"] == "仿宋_GB2312"
    assert tables[0]["index"] == 0
    assert tables[0]["grid"] == [["费率：", "%"]]


def test_replace_text_in_paragraph_keeps_original_format(tmp_path):
    source = _make_template(tmp_path)
    doc_id = tools.word_open_report(str(source))["doc_id"]

    result = tools.word_replace_text(
        doc_id, "标段名称：", "标段名称：一标段", paragraph_index=0)

    assert result["replaced"] == 1
    saved = tools.word_save_report(doc_id, "filled.docx")
    xml = _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")
    assert "标段名称：一标段" in xml
    assert 'w:eastAsia="仿宋_GB2312"' in xml


def test_replace_text_works_inside_table_cell(tmp_path):
    source = _make_template(tmp_path)
    doc_id = tools.word_open_report(str(source))["doc_id"]

    result = tools.word_replace_text(doc_id, "%", "89%", table_index=0, row=0, col=1)

    assert (result["row"], result["col"]) == (0, 1)
    saved = tools.word_save_report(doc_id, "cell.docx")
    assert "89%" in _read_zip_xml(Path(saved["saved_to"]), "word/document.xml")


def test_replace_text_can_append_after_label(tmp_path):
    source = _make_template(tmp_path)
    doc_id = tools.word_open_report(str(source))["doc_id"]

    tools.word_replace_text(doc_id, "", "一标段", paragraph_index=0)

    assert "标段名称：一标段" in tools.word_list_blocks(doc_id)["blocks"][0]["text"]


def test_replace_text_requires_a_location(tmp_path):
    source = _make_template(tmp_path)
    doc_id = tools.word_open_report(str(source))["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_replace_text(doc_id, "标段名称：", "x")

    assert "word_list_blocks" in str(excinfo.value)


def test_replace_text_missing_source_text_gives_actionable_error(tmp_path):
    source = _make_template(tmp_path)
    doc_id = tools.word_open_report(str(source))["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_replace_text(doc_id, "不存在的占位", "x", paragraph_index=0)

    assert "word_list_blocks" in str(excinfo.value)


def test_replace_text_rejects_out_of_range_paragraph(tmp_path):
    source = _make_template(tmp_path)
    doc_id = tools.word_open_report(str(source))["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_replace_text(doc_id, "标段名称：", "x", paragraph_index=99)

    assert "paragraph_index" in str(excinfo.value)


def test_insert_cell_image_puts_picture_into_cell(tmp_path):
    source = _make_template(tmp_path)
    png = tmp_path / "id.png"
    png.write_bytes(base64.b64decode(TEST_PNG_B64))
    doc_id = tools.word_open_report(str(source))["doc_id"]

    result = tools.word_insert_cell_image(
        doc_id, row=0, col=0, table_index=0, image_path=str(png), width_cm=5,
    )

    assert (result["table_index"], result["row"], result["col"]) == (0, 0, 0)
    saved = tools.word_save_report(doc_id, "cell-image.docx")
    path = Path(saved["saved_to"])
    assert "<w:drawing>" in _read_zip_xml(path, "word/document.xml")
    assert any(name.startswith("word/media/") for name in _zip_names(path))


def test_insert_cell_image_rejects_out_of_range_cell(tmp_path):
    source = _make_template(tmp_path)
    png = tmp_path / "id2.png"
    png.write_bytes(base64.b64decode(TEST_PNG_B64))
    doc_id = tools.word_open_report(str(source))["doc_id"]

    with pytest.raises(ToolError) as excinfo:
        tools.word_insert_cell_image(
            doc_id, row=9, col=0, table_index=0, image_path=str(png))

    assert "超出范围" in str(excinfo.value)
