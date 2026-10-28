"""格式读取（format profile）的单测与往返验证。"""

import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wc_word_report_tool import WordFormatter, extract_format_profile


def _build_report(path: Path) -> Path:
    """造一份含两级标题、正文、页眉页脚和表格的报告。"""
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.setup_defaults()
    formatter.heading("一、项目概况", level=1, font_name="黑体", font_size="三号", bold=True)
    formatter.heading("1.1 建设内容", level=2, font_name="楷体", font_size="四号", bold=True)
    formatter.body("正文第一段。", font_name="仿宋", font_size="四号")
    formatter.body("正文第二段。", font_name="仿宋", font_size="四号")
    formatter.set_page_margins(top=3.7, bottom=3.5, left=2.8, right=2.6)
    formatter.set_header_parts(doc.sections[0], center="项目名称")
    formatter.add_footer_page_number(doc.sections[0], prefix="第 ", suffix=" 页")
    formatter.add_table(["项目", "金额"], [["甲", "100"], ["乙", "200"]], font_size=12)
    doc.save(str(path))
    return path


def test_extract_reports_heading_and_body_formats(tmp_path):
    profile = extract_format_profile(_build_report(tmp_path / "report.docx"))

    assert profile["schema_version"] == 1
    assert {"Normal", "Heading 1", "Heading 2"} <= set(profile["styles"])

    heading = profile["styles"]["Heading 1"]
    assert heading["font"]["cn"] == "黑体"
    assert heading["font"]["size_pt"] == 16.0
    assert heading["font"]["bold"] is True
    assert heading["paragraph"]["line_spacing"]["rule"] == "multiple"
    assert heading["paragraph"]["line_spacing"]["value"] == 1.5
    assert heading["paragraph"]["space_before_pt"] == 8.0
    assert heading["paragraph"]["space_after_pt"] == 8.0
    assert heading["paragraph"]["alignment"] == "left"
    assert heading["outline_level"] == 0
    assert heading["used_count"] == 1
    assert heading["resolved_from"]["font.cn"] == "Heading 1"


def test_observed_groups_classify_heading_and_body(tmp_path):
    profile = extract_format_profile(_build_report(tmp_path / "report.docx"))
    groups = {group["role"]: group for group in profile["observed"]["groups"]}

    assert groups["heading_1"]["style"] == "Heading 1"
    assert groups["heading_2"]["style"] == "Heading 2"
    body = groups["body"]
    assert body["count"] == 2
    assert body["font"]["cn"] == "仿宋"
    assert body["paragraph"]["alignment"] == "justify"
    assert body["paragraph"]["indent"]["first_line_chars"] == 2.0


def test_page_header_footer_and_table_are_read(tmp_path):
    profile = extract_format_profile(_build_report(tmp_path / "report.docx"))

    page = profile["document"]["page"]["sections"][0]
    assert page["margins_cm"]["top"] == 3.7
    assert page["margins_cm"]["bottom"] == 3.5
    assert page["margins_cm"]["left"] == 2.8
    assert page["margins_cm"]["right"] == 2.6

    primary_header = profile["document"]["headers_footers"][0]["header"]["primary"]
    assert primary_header["linked_to_previous"] is False
    header_paragraph = primary_header["paragraphs"][0]
    assert "项目名称" in header_paragraph["text"]
    assert len(header_paragraph["tabs"]) == 2
    assert header_paragraph["bottom_border"] is True

    footer_paragraph = profile["document"]["headers_footers"][0]["footer"]["primary"]["paragraphs"][0]
    assert "PAGE" in footer_paragraph["fields"]

    table = profile["observed"]["tables"][0]
    assert table["rows"] == 3
    assert table["columns"] == 2
    assert len(table["column_widths_cm"]) == 2
    assert table["borders"]["top"]["size_pt"] == 0.5
    assert table["header_row"]["font"]["cn"] == "仿宋_GB2312"


def test_default_document_resolves_theme_font_and_size():
    profile = extract_format_profile(Document())
    defaults = profile["document"]["defaults"]["font"]

    assert defaults["size_pt"] == 11.0
    assert defaults["ascii"]  # 主题 minorHAnsi 解析出的实际字体名


def test_character_indent_and_at_least_line_spacing():
    doc = Document()
    paragraph = doc.add_paragraph("正文")
    WordFormatter.set_paragraph_format(
        paragraph,
        line_spacing_pt=24,
        line_spacing_rule="atLeast",
        space_before=6,
        space_after=3,
        first_line_indent_chars=2,
        left_indent_pt=12,
        right_indent_pt=6,
    )

    group = extract_format_profile(doc)["observed"]["groups"][0]
    assert group["paragraph"]["line_spacing"] == {
        "rule": "at_least", "pt": 24.0, "raw_twips": 480,
    }
    assert group["paragraph"]["space_before_pt"] == 6.0
    assert group["paragraph"]["space_after_pt"] == 3.0
    assert group["paragraph"]["indent"]["first_line_chars"] == 2.0
    assert group["paragraph"]["indent"]["left_pt"] == 12.0
    assert group["paragraph"]["indent"]["right_pt"] == 6.0


def test_explicit_false_bold_overrides_heading_style():
    doc = Document()
    paragraph = doc.add_paragraph(style="Heading 1")
    run = paragraph.add_run("不加粗的标题")
    run.font.bold = False

    group = extract_format_profile(doc)["observed"]["groups"][0]
    assert group["font"]["bold"] is False
    assert group["resolved_from"]["font.bold"] == "direct"


def test_custom_outline_paragraph_is_classified_as_heading():
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.set_paragraph_style("我的标题", font_name="黑体", font_size=16, bold=True)
    formatter.add_custom_heading("自定义标题", style_name="我的标题", level=1)

    profile = extract_format_profile(doc)
    assert profile["observed"]["groups"][0]["role"] == "heading_1"
    assert "我的标题" in profile["styles"]


def test_numbering_definition_and_usage_are_read():
    doc = Document()
    numbering = doc.part.numbering_part.element
    abstract = numbering.find(qn("w:abstractNum"))
    abstract_id = abstract.get(qn("w:abstractNumId"))

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "99")
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), abstract_id)
    num.append(reference)
    numbering.append(num)

    paragraph = doc.add_paragraph()
    paragraph.add_run("第一章 总则")
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "99")
    num_pr.append(level)
    num_pr.append(num_id)
    p_pr.append(num_pr)

    profile = extract_format_profile(doc)
    definitions = profile["document"]["numbering"]["definitions"]

    assert len(definitions) == 1
    assert definitions[0]["num_id"] == 99
    assert definitions[0]["used_by"]["paragraphs"] == [0]
    assert definitions[0]["levels"][0]["level"] == 0
    assert definitions[0]["levels"][0]["text"]
    assert any("自动编号" in note for note in profile["notes"])


def test_table_merge_is_reported():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    WordFormatter.set_table_borders(table)

    info = extract_format_profile(doc)["observed"]["tables"][0]
    assert any(
        merge["row"] == 0 and merge["col"] == 0 and merge["colspan"] == 2
        for merge in info["merges"]
    )


def test_profile_is_json_serializable_and_flags_work(tmp_path):
    profile = extract_format_profile(
        _build_report(tmp_path / "report.docx"),
        include_observed=False,
        include_code=False,
        include_raw=True,
    )

    assert "observed" not in profile
    assert "python_code" not in profile
    assert "raw_xml" in profile["styles"]["Normal"]
    assert len(json.dumps(profile, ensure_ascii=False)) > 1000


def test_python_code_round_trips_styles_page_and_header(tmp_path):
    source = _build_report(tmp_path / "source.docx")
    first = extract_format_profile(source)
    code = first["python_code"]
    compile(code, "<format-profile>", "exec")

    target = Document()
    exec(code, {"doc": target})  # noqa: S102 - 测试生成片段的真实可执行性
    target_path = tmp_path / "target.docx"
    target.save(str(target_path))
    second = extract_format_profile(target_path)

    for name in ("Heading 1", "Heading 2"):
        before = first["styles"][name]
        after = second["styles"][name]
        assert before["font"]["cn"] == after["font"]["cn"]
        assert before["font"]["size_pt"] == after["font"]["size_pt"]
        assert before["font"]["bold"] == after["font"]["bold"]
        assert before["paragraph"]["line_spacing"]["value"] == \
            after["paragraph"]["line_spacing"]["value"]
        assert before["paragraph"]["space_before_pt"] == \
            after["paragraph"]["space_before_pt"]

    assert first["document"]["page"]["sections"][0]["margins_cm"] == \
        second["document"]["page"]["sections"][0]["margins_cm"]
    assert "项目名称" in second["document"]["headers_footers"][0]["header"]["primary"]["paragraphs"][0]["text"]
    assert any("PAGE" in note for note in first["notes"])


def test_include_sources_false_and_sample_limit(tmp_path):
    profile = extract_format_profile(
        _build_report(tmp_path / "report.docx"),
        include_sources=False,
        max_samples=1,
    )

    assert "resolved_from" not in profile["styles"]["Heading 1"]
    body = next(group for group in profile["observed"]["groups"] if group["role"] == "body")
    assert body["count"] == 2
    assert len(body["samples"]) == 1


def test_code_emits_at_least_line_spacing_rule(tmp_path):
    doc = Document()
    formatter = WordFormatter(doc)
    formatter.set_paragraph_style(
        "固定行距正文", line_spacing_pt=24, line_spacing_rule="atLeast",
    )
    doc.add_paragraph("正文", style="固定行距正文")
    path = tmp_path / "atleast.docx"
    doc.save(str(path))

    code = extract_format_profile(path)["python_code"]

    assert 'line_spacing_rule="atLeast"' in code
    compile(code, "<format-profile>", "exec")
